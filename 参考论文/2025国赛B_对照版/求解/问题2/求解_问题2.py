"""问题2正式求解：双角色散变投影。

输入为数据/附件1.xlsx、附件2.xlsx；不生成合成观测。主计算保留原始波数坐标，
在1200--3800 cm^-1共同窗口抽取480个原坐标点，先按连续块切分，再在训练块中
估计基线、相位和幅值。所有阶段都写入求解/问题2/结果；总时限为18分钟，
每个搜索情景另有独立小预算；严格固定情景若请求不可行则记录未评估，物理边界夹紧只作单列诊断。
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import os
import time
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import minimize


ROOT = Path(__file__).resolve().parents[2]
RESULT_DIR = ROOT / "求解/问题2/结果"
SOFT_SECONDS = 1080.0
SEARCH_BUDGET_SECONDS = 42.0
FIVE_FOLD_SEARCH_BUDGET_SECONDS = 30.0
SENSITIVITY_SEARCH_BUDGET_SECONDS = 28.0
STANDARD_GRID_SIZE = 96
STANDARD_MAX_STARTS = 12
STANDARD_SEARCH_BUDGET_SECONDS = SEARCH_BUDGET_SECONDS
CALIBRATION_QUANTILE = 0.75
FOLD_STABILITY_RELATIVE_RANGE_THRESHOLD = 0.20
ANGLES = (10.0, 15.0)
BOUNDS = ((0.5, 40.0), (1.2, 6.0), (-1.0, 1.0))
MAIN_SCENARIOS = [(n, c) for n in (2.0, 3.0, 4.0) for c in (-0.5, 0.0, 0.5)]
FOLD_SPECS = [
    {"折号": 1, "训练块": [1, 2, 4, 5, 7, 8, 10, 11], "校准块": [3, 9], "测试块": [6, 12]},
    {"折号": 2, "训练块": [2, 3, 5, 6, 8, 9, 11, 12], "校准块": [4, 10], "测试块": [1, 7]},
]
FIVE_TEST_BLOCKS = [[1, 2], [3, 4], [5, 6], [7, 8], [9, 10]]

# 台账中的检验结果统一从这些实际输出键读取；这里只登记键路径，不把诊断
# 误写成新的数值结论。运行结束时同时写入验证.json和厚度结果.json，便于审计。
HYPOTHESIS_RESULT_KEYS = {
    "A1": ["验证.json/双向留角度冻结诊断", "厚度结果.json/两折条件厚度_微米"],
    "A2": ["验证.json/五折连续留段", "验证.json/五折跨切分稳定性"],
    "A3": ["验证.json/窗口与模型灵敏度", "厚度结果.json/条件区间组成"],
    "A4": ["验证.json/窗口与模型灵敏度", "厚度结果.json/灵敏度情景搜索完整性"],
    "A5": ["区间覆盖.json/校准经验分位点", "区间覆盖.json/测试经验覆盖率",
           "区间覆盖.json/覆盖不足"],
    "A6": ["验证.json/窗口与模型灵敏度", "验证.json/末端新留段"],
    "A7": ["验证.json/五折跨切分稳定性", "厚度结果.json/正式结论句"],
    "A8": ["厚度结果.json/同片最佳厚度_微米", "厚度结果.json/条件区间_微米",
           "区间覆盖.json/测试经验覆盖率"],
}


def clean(value):
    """把numpy标量/数组转为可严格序列化的Python对象。"""
    if isinstance(value, dict):
        return {str(k): clean(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [clean(v) for v in value]
    if isinstance(value, np.ndarray):
        return clean(value.tolist())
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return float(value)
    return value


def write_json(path: Path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + f".{os.getpid()}.tmp")
    with temporary.open("w", encoding="utf-8") as stream:
        json.dump(clean(payload), stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write("\n")
        stream.flush()
        os.fsync(stream.fileno())
    os.replace(temporary, path)


def write_csv(path: Path, rows, fieldnames):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + f".{os.getpid()}.tmp")
    with temporary.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({key: clean(row.get(key, "")) for key in fieldnames})
        stream.flush()
        os.fsync(stream.fileno())
    os.replace(temporary, path)


def elapsed(t0):
    return time.monotonic() - t0


def percentile(values, p):
    return float(np.quantile(np.asarray(values, dtype=float), p))


def read_inputs():
    archive = json.loads((ROOT / "交接/数据档案.json").read_text(encoding="utf-8"))
    metadata = {entry["文件名"]: entry for entry in archive["文件档案"]}
    datasets = []
    audit = {"问题": 2, "文件": [], "共同波数": True, "主窗口_cm^-1": [1200, 3800],
             "抽样点数_每角度": 480, "质量处置": "保留超界值和首点；主窗口与其不相交。"}
    for filename, expected_angle in zip(("附件1.xlsx", "附件2.xlsx"), ANGLES):
        path = ROOT / "数据" / filename
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        expected = metadata[filename]
        if digest != expected["文件哈希"]:
            raise ValueError(f"{filename} SHA-256与数据档案不一致")
        if expected["材料"] != "碳化硅" or expected["入射角度"] != expected_angle:
            raise ValueError(f"{filename}的材料或角度不符合问题2契约")
        workbook = load_workbook(path, read_only=True, data_only=True)
        if workbook.sheetnames != ["Sheet1"]:
            raise ValueError(f"{filename}工作表不是唯一Sheet1")
        rows = list(workbook["Sheet1"].iter_rows(values_only=True))
        workbook.close()
        if len(rows) != 7470 or rows[0][0] != "波数 (cm-1)" or rows[0][1] != "反射率 (%)":
            raise ValueError(f"{filename}表头或数据行数不符合7469点契约")
        data = []
        for source_row, row in enumerate(rows[1:], 2):
            sigma, reflectance = row[:2]
            if not all(isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)
                       for v in (sigma, reflectance)):
                raise ValueError(f"{filename}第{source_row}行含非有限原始值")
            data.append((float(sigma), float(reflectance), source_row))
        data.sort(key=lambda item: item[0])
        if len(data) != 7469 or any(a[0] >= b[0] for a, b in zip(data, data[1:])):
            raise ValueError(f"{filename}波数不是7469个严格递增点")
        datasets.append(data)
        audit["文件"].append({"文件名": filename, "入射角度_度": expected_angle,
                              "SHA256": digest, "原始点数": len(data),
                              "原始反射率单位": "%", "计算反射率单位": "比例",
                              "超百分之百点数": sum(v > 100 for _, v, _ in data),
                              "首点": {"波数_cm^-1": data[0][0], "反射率_百分比": data[0][1],
                                      "原始行号": data[0][2]}})
    if [x[0] for x in datasets[0]] != [x[0] for x in datasets[1]]:
        audit["共同波数"] = False
        raise ValueError("附件1、2波数坐标未逐行对齐")
    source_window = [i for i, item in enumerate(datasets[0]) if 1200 <= item[0] <= 3800]
    selected = [source_window[j * (len(source_window) - 1) // 479] for j in range(480)]
    sigma = np.array([datasets[0][i][0] for i in selected], dtype=float)
    values = np.array([[datasets[a][i][1] / 100.0 for i in selected] for a in range(2)], dtype=float)
    source_rows = [[datasets[a][i][2] for i in selected] for a in range(2)]
    audit["共同窗口原始点数_每角度"] = len(source_window)
    audit["抽样索引"] = [{"样本序号": j + 1, "块号": j // 40 + 1,
                         "波数_cm^-1": float(sigma[j]),
                         "原始行号_附件一": source_rows[0][j],
                         "原始行号_附件二": source_rows[1][j],
                         "质量标记": "共同原坐标；未插值、未截断"} for j in range(480)]
    return sigma, values, source_rows, datasets, audit


def block_of(index):
    return index // 40 + 1


def make_fold(sigma, values, spec, angles=ANGLES):
    train_blocks, calibration_blocks, test_blocks = (set(spec[k]) for k in ("训练块", "校准块", "测试块"))
    if train_blocks & calibration_blocks or train_blocks & test_blocks or calibration_blocks & test_blocks:
        raise ValueError("折分块重叠")
    if train_blocks | calibration_blocks | test_blocks != set(range(1, 13)):
        raise ValueError("折分块未覆盖1至12")
    edges = np.array([(sigma[k - 1] + sigma[k]) / 2 for k in range(40, 480, 40)])
    guard = [edge for block, edge in enumerate(edges, 1)
             if (block in train_blocks) != (block + 1 in train_blocks)]
    train_indices = np.array([i for i in range(len(sigma))
                              if block_of(i) in train_blocks and all(abs(sigma[i] - edge) >= 20 for edge in guard)], dtype=int)
    if len(train_indices) < 30:
        raise ValueError("保护带后训练点过少")
    center = float((sigma[train_indices].min() + sigma[train_indices].max()) / 2)
    half_span = float((sigma[train_indices].max() - sigma[train_indices].min()) / 2)
    x = (sigma[train_indices] - center) / half_span
    angle_data = []
    for ai, angle in enumerate(angles):
        y = values[ai, train_indices]
        base = np.column_stack((np.ones(len(x)), x, x * x))
        coefficient = np.linalg.lstsq(base, y, rcond=None)[0]
        residual = y - base @ coefficient
        scale = max(percentile(residual, 0.75) - percentile(residual, 0.25), 1e-4)
        angle_data.append({"角度_度": float(angle), "sigma": sigma[train_indices], "x": x,
                           "y": y, "scale": float(scale), "基线系数": coefficient})
    return {"spec": spec, "sigma": sigma, "values": values, "angles": angle_data,
            "train_indices": train_indices, "center": center, "half_span": half_span,
            "guard_edges": guard, "折分_训练点数": int(len(train_indices))}


def full_fold(sigma, values, angles=ANGLES):
    spec = {"折号": 0, "训练块": list(range(1, 13)), "校准块": [], "测试块": []}
    return make_fold(sigma, values, spec, angles)


def valid_parameters(params):
    d, n_ref, dispersion = params
    if not all(lo <= value <= hi for value, (lo, hi) in zip(params, BOUNDS)):
        return False
    n_low = n_ref + dispersion * ((2000 / 1200) ** 2 - 1)
    n_high = n_ref + dispersion * ((2000 / 3800) ** 2 - 1)
    return min(n_low, n_high) > math.sin(math.radians(15))


def phase_rate(sigma, angle, n_ref, dispersion):
    n = n_ref + dispersion * ((2000.0 / sigma) ** 2 - 1.0)
    q = np.sqrt(np.maximum(n * n - math.sin(math.radians(angle)) ** 2, 0.0))
    return 4.0 * math.pi / 10000.0 * sigma * q


def fit_angle(params, data, base_degree=2, amp_degree=1, fixed_phase=None, angle_override=None):
    angle = data["角度_度"] if angle_override is None else float(angle_override)
    rate = phase_rate(data["sigma"], angle, params[1], params[2]) * params[0]
    phase_candidates = [float(fixed_phase)] if fixed_phase is not None else [j * math.pi / 12 for j in range(12)]

    def at_phase(phase):
        phase = float(phase) % math.pi
        columns = [data["x"] ** k for k in range(base_degree + 1)]
        columns.extend([(data["x"] ** k) * np.cos(rate + phase) for k in range(amp_degree + 1)])
        matrix = np.column_stack(columns)
        coefficient, _, rank, _ = np.linalg.lstsq(matrix, data["y"], rcond=None)
        residual = data["y"] - matrix @ coefficient
        return {"相位_rad": phase, "系数": coefficient, "残差": residual,
                "损失": float(np.mean(residual * residual) / data["scale"] ** 2),
                "秩": int(rank), "条件数": float(np.linalg.cond(matrix)),
                "rate": rate, "角度_度": angle, "基线阶数": base_degree, "幅值阶数": amp_degree}

    best = min((at_phase(phase) for phase in phase_candidates), key=lambda item: item["损失"])
    if fixed_phase is None:
        step = math.pi / 12
        for _ in range(8):
            candidates = [best, at_phase(best["相位_rad"] - step), at_phase(best["相位_rad"] + step)]
            best = min(candidates, key=lambda item: item["损失"])
            step /= 2
    return best


def fit_candidate(params, fold, base_degree=2, amp_degree=1, angles=ANGLES):
    if not valid_parameters(params):
        return None
    models = [fit_angle(params, data, base_degree, amp_degree)
              for data in fold["angles"]]
    return {"参数": tuple(float(x) for x in params), "模型": models,
            "损失": float(np.mean([model["损失"] for model in models]))}


def search_fold(fold, t0, base_degree=2, amp_degree=1, dispersion_mode="经验",
                grid_size=STANDARD_GRID_SIZE, max_starts=STANDARD_MAX_STARTS, fixed_n=None,
                fixed_n_policy="严格", budget_seconds=STANDARD_SEARCH_BUDGET_SECONDS):
    """在独立小预算内完成搜索；严格情景不把不可行请求投影成±3%。"""
    local_t0 = time.monotonic()
    deadline = min(local_t0 + budget_seconds, t0 + SOFT_SECONDS - 30.0)
    pool = []
    c_values = [0.0] if dispersion_mode == "常数" else [-0.5, 0.0, 0.5]
    requested_fixed_n = None if fixed_n is None else float(fixed_n)
    effective_fixed_n = requested_fixed_n
    fixed_n_clamped = False
    if effective_fixed_n is not None:
        feasible_request = (BOUNDS[1][0] <= effective_fixed_n <= BOUNDS[1][1] and
                            any(valid_parameters((1.0, effective_fixed_n, c)) for c in c_values))
        if not feasible_request:
            if fixed_n_policy == "严格":
                return None, {"粗网格次数": 0, "粗网格计划次数": 0, "局部起点数": 0,
                              "搜索预算截断": False, "搜索未评估": True,
                              "搜索状态": "未评估_请求固定参考折射率不可行",
                              "固定参数一致": False, "搜索用时秒": elapsed(local_t0),
                              "搜索预算秒": float(budget_seconds), "基线阶数": base_degree,
                              "幅值阶数": amp_degree, "色散形式": dispersion_mode,
                              "固定参考折射率": None,
                              "请求固定参考折射率": requested_fixed_n,
                              "固定参考折射率是否夹紧": False,
                              "固定参数策略": fixed_n_policy,
                              "未评估原因": "请求值超出搜索盒或不存在满足传播约束的色散候选"}
            if fixed_n_policy != "夹紧":
                raise ValueError(f"未知固定参考折射率策略: {fixed_n_policy}")
            effective_fixed_n = min(max(effective_fixed_n, BOUNDS[1][0]), BOUNDS[1][1])
            feasible_after_clamp = any(valid_parameters((1.0, effective_fixed_n, c))
                                       for c in c_values)
            if not feasible_after_clamp:
                return None, {"粗网格次数": 0, "粗网格计划次数": 0, "局部起点数": 0,
                              "搜索预算截断": False, "搜索未评估": True,
                              "搜索状态": "未评估_夹紧后仍不可行",
                              "固定参数一致": False, "搜索用时秒": elapsed(local_t0),
                              "搜索预算秒": float(budget_seconds), "基线阶数": base_degree,
                              "幅值阶数": amp_degree, "色散形式": dispersion_mode,
                              "固定参考折射率": None,
                              "请求固定参考折射率": requested_fixed_n,
                              "固定参考折射率是否夹紧": True,
                              "固定参数策略": fixed_n_policy,
                              "未评估原因": "投影到搜索盒边界后仍不存在满足传播约束的色散候选"}
            fixed_n_clamped = True
    fixed_n = effective_fixed_n
    n_values = [float(fixed_n)] if fixed_n is not None else [2.0, 3.0, 4.0]
    thicknesses = np.linspace(0.5, 40.0, grid_size)
    coarse_count = 0

    def retain(candidate):
        if candidate is None:
            return
        if any(abs(candidate["参数"][0] - item["参数"][0]) < 1e-12 and
               np.max(np.abs(np.asarray(candidate["参数"]) - np.asarray(item["参数"]))) < 1e-12
               for item in pool):
            return
        pool.append(candidate)
        pool.sort(key=lambda item: item["损失"])
        del pool[8:]

    stopped = False
    for n_ref in n_values:
        for dispersion in c_values:
            for thickness in thicknesses:
                if time.monotonic() >= deadline:
                    stopped = True
                    break
                candidate = fit_candidate((float(thickness), n_ref, dispersion), fold,
                                          base_degree, amp_degree)
                coarse_count += 1
                retain(candidate)
            if stopped:
                break
        if stopped:
            break
    if not pool:
        fallback_n = float(fixed_n) if fixed_n is not None else 3.0
        fallback_c = 0.0 if dispersion_mode == "常数" else 0.0
        fallback = (5.1, fallback_n, fallback_c)
        fallback_candidate = fit_candidate(fallback, fold, base_degree, amp_degree)
        if fallback_candidate is None:
            raise RuntimeError("没有找到满足固定参数和物理约束的合法候选")
        pool = [fallback_candidate]
        stopped = True
    starts = [item["参数"] for item in pool[:max_starts]]
    starts.extend([(5.171033357319079, 3.0, -0.5), (5.073684210526316, 3.0, -0.5),
                   (8.0, 2.0, 0.0), (20.0, 4.0, 0.5)])
    unique = []
    for start in starts:
        if fixed_n is not None:
            start = (start[0], fixed_n, start[2])
        if dispersion_mode == "常数":
            start = (start[0], start[1], 0.0)
        if valid_parameters(start) and start not in unique:
            unique.append(start)
    for start in unique[:max_starts]:
        if time.monotonic() >= deadline:
            stopped = True
            break
        if fixed_n is None and dispersion_mode == "经验":
            z0, bounds = np.array(start, dtype=float), [(0.5, 40.0), (1.2, 6.0), (-1.0, 1.0)]
            def objective(z):
                if time.monotonic() >= deadline:
                    return 1e10
                trial = fit_candidate(tuple(z), fold, base_degree, amp_degree)
                return 1e10 if trial is None else trial["损失"]
        elif fixed_n is None:
            z0 = np.array([start[0], start[1]], dtype=float)
            bounds = [(0.5, 40.0), (1.2, 6.0)]
            def objective(z):
                if time.monotonic() >= deadline:
                    return 1e10
                p = (float(z[0]), float(fixed_n if fixed_n is not None else z[1]),
                     0.0 if dispersion_mode == "常数" else float(start[2]))
                trial = fit_candidate(p, fold, base_degree, amp_degree)
                return 1e10 if trial is None else trial["损失"]
        else:
            z0 = np.array([start[0]], dtype=float)
            bounds = [(0.5, 40.0)]
            def objective(z):
                if time.monotonic() >= deadline:
                    return 1e10
                p = (float(z[0]), float(fixed_n),
                     0.0 if dispersion_mode == "常数" else float(start[2]))
                trial = fit_candidate(p, fold, base_degree, amp_degree)
                return 1e10 if trial is None else trial["损失"]

        outcome = minimize(objective, z0, method="L-BFGS-B", bounds=bounds,
                           options={"maxiter": 70, "ftol": 1e-9, "maxls": 20})
        if fixed_n is None and dispersion_mode == "经验":
            params = tuple(float(x) for x in outcome.x)
        elif fixed_n is None:
            params = (float(outcome.x[0]), float(fixed_n if fixed_n is not None else outcome.x[1]),
                      0.0 if dispersion_mode == "常数" else float(start[2]))
        else:
            params = (float(outcome.x[0]), float(fixed_n),
                      0.0 if dispersion_mode == "常数" else float(start[2]))
        retain(fit_candidate(params, fold, base_degree, amp_degree))
    best = min(pool, key=lambda item: item["损失"])
    fixed_consistent = fixed_n is None or abs(best["参数"][1] - float(fixed_n)) <= 1e-12
    if not fixed_consistent:
        raise RuntimeError("搜索结果未满足固定参考折射率情景")
    return best, {"粗网格次数": coarse_count, "粗网格计划次数": int(len(n_values) * len(c_values) * grid_size),
                  "局部起点数": min(len(unique), max_starts), "搜索预算截断": stopped,
                  "搜索未评估": False,
                  "搜索状态": "限时完成但候选合法" if stopped else "完整完成",
                  "固定参数一致": fixed_consistent,
                  "搜索用时秒": elapsed(local_t0), "搜索预算秒": float(budget_seconds),
                  "基线阶数": base_degree, "幅值阶数": amp_degree,
                  "色散形式": dispersion_mode, "固定参考折射率": fixed_n,
                  "请求固定参考折射率": requested_fixed_n,
                  "固定参考折射率是否夹紧": fixed_n_clamped,
                  "固定参数策略": fixed_n_policy}


def predict_model(model, sigma, center, half_span):
    x = (np.asarray(sigma, dtype=float) - center) / half_span
    rate = model["rate"] if len(model["rate"]) == len(x) else None
    if rate is None:
        raise ValueError("预测坐标必须与模型rate同步生成")
    matrix = np.column_stack([np.ones(len(x)), x, x * x,
                              np.cos(rate + model["相位_rad"]),
                              x * np.cos(rate + model["相位_rad"])])
    if len(model["系数"]) != 5:
        powers = [x ** k for k in range(len(model["系数"]) - 2)]
        # 仅供灵敏度结果使用；主模型始终为二次基线、一次幅值。
        matrix = np.column_stack(powers[:3] + [x ** k * np.cos(rate + model["相位_rad"])
                                                 for k in range(len(model["系数"]) - 3)])
    return matrix @ model["系数"]


def predict_params(params, model, sigma, center, half_span, angle):
    rate = phase_rate(np.asarray(sigma, dtype=float), angle, params[1], params[2]) * params[0]
    x = (np.asarray(sigma, dtype=float) - center) / half_span
    coefficient = model["系数"]
    base_degree = int(model.get("基线阶数", 2))
    amp_degree = int(model.get("幅值阶数", len(coefficient) - base_degree - 1))
    columns = [x ** k for k in range(base_degree + 1)]
    columns.extend([(x ** k) * np.cos(rate + model["相位_rad"]) for k in range(amp_degree + 1)])
    return np.column_stack(columns) @ coefficient


def score_blocks(fold, candidate):
    blocks = []
    if not fold["spec"]["测试块"]:
        return blocks
    cal_indices = np.array([i for i in range(len(fold["sigma"]))
                            if block_of(i) in set(fold["spec"]["校准块"])], dtype=int)
    for ai, model in enumerate(candidate["模型"]):
        angle = ANGLES[ai]
        if len(cal_indices):
            cal_pred = predict_params(candidate["参数"], model, fold["sigma"][cal_indices],
                                       fold["center"], fold["half_span"], angle)
            cal_abs = np.sort(np.abs(fold["values"][ai, cal_indices] - cal_pred))
            rank = min(len(cal_abs) - 1,
                       max(0, math.ceil(CALIBRATION_QUANTILE * len(cal_abs)) - 1))
            half_width = float(cal_abs[rank])
        else:
            half_width = float("nan")
        for block in fold["spec"]["测试块"]:
            indices = np.array([i for i in range(len(fold["sigma"])) if block_of(i) == block], dtype=int)
            pred = predict_params(candidate["参数"], model, fold["sigma"][indices],
                                  fold["center"], fold["half_span"], angle)
            observed = fold["values"][ai, indices]
            residual = observed - pred
            coverage = None if not math.isfinite(half_width) else float(np.mean(np.abs(residual) <= half_width))
            blocks.append({"折号": fold["spec"]["折号"], "角度_度": angle, "测试块": block,
                           "点数": int(len(indices)), "样本序号": (indices + 1).tolist(),
                           "波数_cm^-1": fold["sigma"][indices].tolist(),
                           "观测反射率_比例": observed.tolist(), "预测反射率_比例": pred.tolist(),
                           "残差_比例": residual.tolist(), "均方根误差_比例": float(np.sqrt(np.mean(residual ** 2))),
                           "训练尺度_比例": float(fold["angles"][ai]["scale"]),
                           "标准化均方根误差": float(np.sqrt(np.mean(residual ** 2)) / fold["angles"][ai]["scale"]),
                           "校准点数": int(len(cal_indices)), "区间半宽_比例": half_width,
                           "校准经验分位点": CALIBRATION_QUANTILE,
                           "经验覆盖率": coverage,
                           "区间下界_比例": (pred - half_width).tolist() if math.isfinite(half_width) else [],
                           "区间上界_比例": (pred + half_width).tolist() if math.isfinite(half_width) else []})
    return blocks


def baseline_blocks(fold):
    rows = []
    for ai, data in enumerate(fold["angles"]):
        x_train, y_train = data["x"], data["y"]
        mean_value = float(np.mean(y_train))
        quad = np.polyfit(x_train, y_train, 2)
        for block in fold["spec"]["测试块"]:
            indices = np.array([i for i in range(len(fold["sigma"])) if block_of(i) == block], dtype=int)
            x_test = (fold["sigma"][indices] - fold["center"]) / fold["half_span"]
            observed = fold["values"][ai, indices]
            predictions = {"训练均值": np.full(len(indices), mean_value),
                           "训练二次趋势": np.polyval(quad, x_test)}
            for name, pred in predictions.items():
                residual = observed - pred
                rows.append({"折号": fold["spec"]["折号"], "角度_度": ANGLES[ai], "测试块": block,
                             "基线": name, "均方根误差_比例": float(np.sqrt(np.mean(residual ** 2))),
                             "标准化均方根误差": float(np.sqrt(np.mean(residual ** 2)) / data["scale"])})
    return rows


def frozen_angle_checks(folds, candidates):
    checks = []
    for fold, candidate in zip(folds, candidates):
        for source, target in ((0, 1), (1, 0)):
            source_phase = candidate["模型"][source]["相位_rad"]
            target_data = fold["angles"][target]
            frozen_model = fit_angle(candidate["参数"], target_data, fixed_phase=source_phase)
            for block in fold["spec"]["测试块"]:
                indices = np.array([i for i in range(len(fold["sigma"])) if block_of(i) == block], dtype=int)
                pred = predict_params(candidate["参数"], frozen_model, fold["sigma"][indices],
                                      fold["center"], fold["half_span"], ANGLES[target])
                residual = fold["values"][target, indices] - pred
                checks.append({"折号": fold["spec"]["折号"], "来源角度_度": ANGLES[source],
                               "被预测角度_度": ANGLES[target], "测试块": block,
                               "冻结厚度_um": candidate["参数"][0],
                               "冻结参考折射率": candidate["参数"][1], "冻结色散系数": candidate["参数"][2],
                               "来源相位冻结": True, "目标角度仅重估有限基线幅值": True,
                               "标准化均方根误差": float(np.sqrt(np.mean(residual ** 2)) /
                                                       fold["angles"][target]["scale"])})
    return checks


def five_fold_validation(sigma, values, t0):
    rows, fold_info = [], []
    for fold_no, test_blocks in enumerate(FIVE_TEST_BLOCKS, 1):
        calibration = [11, 12]
        training = [b for b in range(1, 13) if b not in set(test_blocks + calibration)]
        spec = {"折号": fold_no + 10, "训练块": training, "校准块": calibration, "测试块": test_blocks}
        fold = make_fold(sigma, values, spec)
        candidate, info = search_fold(fold, t0, grid_size=48, max_starts=4,
                                      budget_seconds=FIVE_FOLD_SEARCH_BUDGET_SECONDS)
        blocks = score_blocks(fold, candidate)
        score = float(np.mean([b["标准化均方根误差"] for b in blocks])) if blocks else None
        fold_info.append({"折号": fold_no, "测试块": test_blocks, "训练块": training,
                          "主指标": score, "搜索": info, "搜索状态": info["搜索状态"],
                          "条件厚度_um": candidate["参数"][0],
                          "搜索未截断": not bool(info["搜索预算截断"]),
                          "条件厚度有效": bool(valid_parameters(candidate["参数"])),
                          "逐折搜索完整且参数有效": (info["搜索状态"] == "完整完成" and
                                                     bool(valid_parameters(candidate["参数"])))})
        rows.extend(blocks)
        write_json(RESULT_DIR / "验证.json", {"状态": "五折连续留段进行中", "五折": fold_info})
    return fold_info, rows


def summarize_five_fold(fold_info):
    """把五折厚度离散度转成可复核的稳定性标签，不缩窄条件范围。"""
    valid_rows = [row for row in fold_info if row.get("条件厚度有效", False)]
    thicknesses = [float(row["条件厚度_um"]) for row in valid_rows]
    all_complete_valid = bool(fold_info) and all(
        row.get("逐折搜索完整且参数有效", False) for row in fold_info)
    if not thicknesses:
        return {"有效折数": 0, "五折条件厚度_微米": [], "跨切分稳定性": "无有效折",
                "逐折搜索完整且参数有效": all_complete_valid,
                "结论口径": "低置信：五折没有形成有效条件厚度，不能给出跨切分稳定性判断。",
                "稳定性判据": "有效折数不足，未计算相对极差。"}
    median = float(np.median(thicknesses))
    minimum, maximum = float(min(thicknesses)), float(max(thicknesses))
    relative_range = float((maximum - minimum) / abs(median)) if median else float("inf")
    unstable = (not all_complete_valid or len(thicknesses) < len(fold_info) or
                relative_range >= FOLD_STABILITY_RELATIVE_RANGE_THRESHOLD)
    status = "不稳定" if unstable else "相对稳定"
    wording = ("低置信、跨切分不稳定；保留最佳估计与完整条件范围，不能写成稳定绝对厚度。"
               if unstable else
               "五折相对稳定；仍只把结果解释为条件厚度诊断，不作为统计置信区间。")
    return {"有效折数": len(thicknesses), "五折总折数": len(fold_info),
            "逐折搜索完整且参数有效": all_complete_valid,
            "五折条件厚度_微米": thicknesses, "最小厚度_微米": minimum,
            "最大厚度_微米": maximum, "中位数厚度_微米": median,
            "标准差_微米": float(np.std(thicknesses, ddof=1)) if len(thicknesses) > 1 else 0.0,
            "相对极差": relative_range,
            "稳定性判据": "五折厚度极差/中位数≥20%即标记不稳定；仅用于结论分级",
            "跨切分稳定性": status, "结论口径": wording}


def assess_publication(full_candidate, condition_range, main_score, baseline_score,
                       coverage, five_summary):
    """把“答案是否形成”与质量诊断分开，避免低置信被误写成无答案。"""
    point_valid = (full_candidate is not None and len(full_candidate.get("参数", ())) == 3 and
                   all(math.isfinite(float(value)) for value in full_candidate["参数"]))
    interval_valid = (len(condition_range) == 2 and
                      all(math.isfinite(float(value)) for value in condition_range) and
                      float(condition_range[0]) <= float(condition_range[1]))
    unstable = five_summary.get("跨切分稳定性") == "不稳定"
    coverage_low = math.isfinite(float(coverage)) and float(coverage) < CALIBRATION_QUANTILE
    low_confidence = unstable or coverage_low
    reasons = []
    if unstable:
        reasons.append("五折跨切分不稳定")
    if coverage_low:
        reasons.append("反射率留出覆盖低于校准经验分位点")
    return {
        "答案发布判定": "PASS（条件发布）" if point_valid and interval_valid else "FAIL（无有效答案）",
        "点估计已形成": point_valid,
        "条件范围已形成": interval_valid,
        "质量诊断": {
            "正式基线方向": "主方法不劣于正式二次趋势基线" if main_score <= baseline_score
            else "主方法劣于正式二次趋势基线",
            "跨切分稳定性": five_summary.get("跨切分稳定性", "未形成"),
            "反射率留出覆盖": float(coverage),
        },
        "置信等级": "低" if low_confidence else "条件中",
        "降级依据": "；".join(reasons) if reasons else "未触发已登记的降级条件",
        "口径": "PASS仅表示真实附件已形成点估计与条件范围；不表示绝对厚度稳定性或统计置信协议全部通过。",
    }


def profile(fold, candidate, t0):
    d, n, c = candidate["参数"]
    d_axis = sorted(set(np.linspace(0.5, 40.0, 17).tolist() + [d]))
    n_axis = sorted(set(np.linspace(1.2, 6.0, 13).tolist() + [n]))
    entries = []
    for n_ref in n_axis:
        for thickness in d_axis:
            if elapsed(t0) >= SOFT_SECONDS - 30:
                break
            params = (float(thickness), float(n_ref), float(c))
            trial = fit_candidate(params, fold)
            if trial is not None:
                entries.append({"厚度_um": thickness, "参考折射率": n_ref,
                                "色散系数": c, "训练标准化均方损失": trial["损失"]})
    if not entries:
        entries = [{"厚度_um": float(d), "参考折射率": float(n),
                    "色散系数": float(c), "训练标准化均方损失": float(candidate["损失"])}]
    best_loss = min(item["训练标准化均方损失"] for item in entries)
    near = [item for item in entries if item["训练标准化均方损失"] <= best_loss * 1.10 + 1e-12]
    return {"构造方法": "固定全量拟合的经验色散系数，网格扫描厚度与参考折射率；每格重估角度常相位和线性系数。",
            "厚度轴_um": d_axis, "参考折射率轴": n_axis, "数值": entries,
            "最低训练标准化损失": best_loss, "近等损失阈值": 0.10,
            "近等损失厚度范围_um": [min(x["厚度_um"] for x in near), max(x["厚度_um"] for x in near)],
            "近等损失组合数": len(near), "是否统计置信区间": False,
            "解释": "这是光学条件剖面范围，不是统计置信区间；未知折射率使厚度存在条件可辨识性。"}


def sensitivity(sigma_full, values_full, main_fold, main_candidate, datasets, t0):
    rows = []
    variants = [
        {"名称": "主模型", "基线阶数": 2, "幅值阶数": 1, "色散形式": "经验",
         "固定参考折射率": None, "固定参数策略": "严格", "厚度网格点数": STANDARD_GRID_SIZE,
         "扰动口径": "正式主方法同口径；用于搜索设置对照", "纳入条件范围": True},
        {"名称": "常数基线", "基线阶数": 0, "幅值阶数": 1, "色散形式": "经验",
         "固定参考折射率": None, "固定参数策略": "严格", "厚度网格点数": STANDARD_GRID_SIZE,
         "扰动口径": "基线阶数从二次改为常数", "纳入条件范围": True},
        {"名称": "一次基线", "基线阶数": 1, "幅值阶数": 1, "色散形式": "经验",
         "固定参考折射率": None, "固定参数策略": "严格", "厚度网格点数": STANDARD_GRID_SIZE,
         "扰动口径": "基线阶数从二次改为一次", "纳入条件范围": True},
        {"名称": "常数幅值", "基线阶数": 2, "幅值阶数": 0, "色散形式": "经验",
         "固定参考折射率": None, "固定参数策略": "严格", "厚度网格点数": STANDARD_GRID_SIZE,
         "扰动口径": "幅值阶数从一次改为常数", "纳入条件范围": True},
        {"名称": "常数折射率", "基线阶数": 2, "幅值阶数": 1, "色散形式": "常数",
         "固定参考折射率": None, "固定参数策略": "严格", "厚度网格点数": STANDARD_GRID_SIZE,
         "扰动口径": "色散形式从经验改为常数", "纳入条件范围": True},
        {"名称": "参考折射率相对减3%_严格请求", "基线阶数": 2, "幅值阶数": 1,
         "色散形式": "经验", "固定参考折射率": main_candidate["参数"][1] * 0.97,
         "固定参数策略": "严格", "厚度网格点数": STANDARD_GRID_SIZE,
         "请求相对变化": -0.03,
         "扰动口径": "相对全量主模型参考折射率的-3%严格请求；不可行请求标记未评估",
         "纳入条件范围": True},
        {"名称": "参考折射率相对加3%_严格请求", "基线阶数": 2, "幅值阶数": 1,
         "色散形式": "经验", "固定参考折射率": main_candidate["参数"][1] * 1.03,
         "固定参数策略": "严格", "厚度网格点数": STANDARD_GRID_SIZE,
         "请求相对变化": 0.03,
         "扰动口径": "相对全量主模型参考折射率的+3%严格请求；不可行请求标记未评估",
         "纳入条件范围": True},
        {"名称": "物理可行域夹紧诊断_下边界外请求", "基线阶数": 2, "幅值阶数": 1,
         "色散形式": "经验", "固定参考折射率": max(0.0, BOUNDS[1][0] - 0.03 * main_candidate["参数"][1]),
         "固定参数策略": "夹紧", "厚度网格点数": STANDARD_GRID_SIZE,
         "请求相对变化": None,
         "扰动口径": "仅检验搜索盒下边界投影；不代表±3%且不纳入条件范围",
         "纳入条件范围": False},
        {"名称": "厚度网格加密2倍", "基线阶数": 2, "幅值阶数": 1, "色散形式": "经验",
         "固定参考折射率": None, "固定参数策略": "严格", "厚度网格点数": 192,
         "扰动口径": "厚度粗网格由96点加密为192点", "纳入条件范围": True},
    ]
    full = full_fold(sigma_full, values_full)

    def flush_rows():
        write_json(RESULT_DIR / "灵敏度_中间.json",
                   {"状态": "灵敏度情景进行中", "窗口与模型灵敏度": rows,
                    "实际用时秒": elapsed(t0)})

    for variant in variants:
        name = variant["名称"]
        candidate, info = search_fold(
            full, t0, variant["基线阶数"], variant["幅值阶数"], variant["色散形式"],
            grid_size=variant["厚度网格点数"], max_starts=STANDARD_MAX_STARTS,
            fixed_n=variant["固定参考折射率"], fixed_n_policy=variant["固定参数策略"],
            budget_seconds=STANDARD_SEARCH_BUDGET_SECONDS)
        admissible = bool(candidate is not None and info["固定参数一致"] and
                          not info.get("搜索未评估", False) and variant["纳入条件范围"])
        rows.append({"灵敏度项": name, "情景名称": name,
                     "厚度_um": None if candidate is None else candidate["参数"][0],
                     "参考折射率": None if candidate is None else candidate["参数"][1],
                     "色散系数": None if candidate is None else candidate["参数"][2],
                     "训练标准化损失": None if candidate is None else candidate["损失"],
                     "搜索": json.dumps(clean(info), ensure_ascii=False),
                     "搜索状态": info["搜索状态"],
                     "搜索未截断": not bool(info["搜索预算截断"]),
                     "搜索未评估": bool(info.get("搜索未评估", False)),
                     "固定参考折射率": info["固定参考折射率"],
                     "请求固定参考折射率": info["请求固定参考折射率"],
                     "实际固定参考折射率": info["固定参考折射率"],
                     "固定参考折射率是否夹紧": info["固定参考折射率是否夹紧"],
                     "固定参数策略": variant["固定参数策略"],
                     "固定参数一致": info["固定参数一致"],
                     "请求相对变化": variant.get("请求相对变化"),
                     "实际相对变化": (None if variant.get("请求相对变化") is None or
                                      info["固定参考折射率"] is None else
                                      info["固定参考折射率"] / main_candidate["参数"][1] - 1.0),
                     "参考折射率扰动口径": variant["扰动口径"],
                     "是否纳入条件范围": admissible,
                     "未评估原因": info.get("未评估原因", ""),
                     "基线阶数": variant["基线阶数"], "幅值阶数": variant["幅值阶数"],
                     "色散形式": variant["色散形式"],
                     "厚度网格点数": variant["厚度网格点数"],
                     "搜索设置与正式主方法一致": (variant["厚度网格点数"] == STANDARD_GRID_SIZE and
                                                   STANDARD_MAX_STARTS == 12 and
                                                   STANDARD_SEARCH_BUDGET_SECONDS == SEARCH_BUDGET_SECONDS)})
        flush_rows()
    for name, low, high in (("窗口下界加100", 1300, 3800), ("窗口上界加100", 1200, 3900)):
        selected = [i for i, item in enumerate(datasets[0]) if low <= item[0] <= high]
        chosen = [selected[j * (len(selected) - 1) // 479] for j in range(480)]
        sigma = np.array([datasets[0][i][0] for i in chosen], dtype=float)
        vals = np.array([[datasets[a][i][1] / 100.0 for i in chosen] for a in range(2)], dtype=float)
        fold = full_fold(sigma, vals)
        candidate, info = search_fold(fold, t0, grid_size=STANDARD_GRID_SIZE,
                                      max_starts=STANDARD_MAX_STARTS,
                                      budget_seconds=STANDARD_SEARCH_BUDGET_SECONDS)
        rows.append({"灵敏度项": name, "情景名称": name,
                     "厚度_um": candidate["参数"][0], "参考折射率": candidate["参数"][1],
                     "色散系数": candidate["参数"][2], "训练标准化损失": candidate["损失"],
                     "搜索": json.dumps(clean(info), ensure_ascii=False),
                     "搜索状态": info["搜索状态"], "搜索未截断": not bool(info["搜索预算截断"]),
                     "搜索未评估": bool(info.get("搜索未评估", False)),
                     "固定参考折射率": None, "请求固定参考折射率": None,
                     "实际固定参考折射率": None, "固定参考折射率是否夹紧": False,
                     "固定参数策略": "严格", "固定参数一致": info["固定参数一致"],
                     "请求相对变化": None, "实际相对变化": None,
                     "参考折射率扰动口径": "窗口移动100 cm^-1；不改变参考折射率定义",
                     "是否纳入条件范围": True, "未评估原因": "",
                     "基线阶数": 2, "幅值阶数": 1, "色散形式": "经验",
                     "厚度网格点数": STANDARD_GRID_SIZE, "波数窗口_cm^-1": f"{low}-{high}",
                     "搜索设置与正式主方法一致": True})
        flush_rows()
    return rows


def end_window_check(datasets, full_candidate, full_fold_obj):
    rows = []
    indices = [i for i, item in enumerate(datasets[0]) if 3800 <= item[0] <= 4000.122]
    protected = sum(3750 <= item[0] < 3800 for item in datasets[0])
    for ai, model in enumerate(full_candidate["模型"]):
        sigma = np.array([datasets[0][i][0] for i in indices], dtype=float)
        pred = predict_params(full_candidate["参数"], model, sigma, full_fold_obj["center"],
                              full_fold_obj["half_span"], ANGLES[ai])
        obs = np.array([datasets[ai][i][1] / 100.0 for i in indices])
        residual = obs - pred
        rows.append({"角度_度": ANGLES[ai], "测试窗口_cm^-1": [3800, 4000.122],
                     "保护窗口_cm^-1": [3750, 3800], "保护点数": protected,
                     "测试点数": len(indices), "均方根误差_比例": float(np.sqrt(np.mean(residual ** 2))),
                     "平均绝对误差_比例": float(np.mean(np.abs(residual))),
                     "预测反射率范围_比例": [float(np.min(pred)), float(np.max(pred))],
                     "限制": "末端段属于选型后外层检验，不能替代真实厚度真值验证。"})
    return rows


def main():
    t0 = time.monotonic()
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    state = {"问题": 2, "运行状态": "开始", "实际用时秒": 0.0, "时间预算秒": SOFT_SECONDS}
    write_json(RESULT_DIR / "执行状态.json", state)
    try:
        sigma_full, values_full, source_rows, datasets, audit = read_inputs()
        audit["实际用时秒"] = elapsed(t0)
        write_json(RESULT_DIR / "数据审查.json", audit)
        folds = [make_fold(sigma_full[0:480] if len(sigma_full) == 480 else sigma_full,
                           values_full, spec) for spec in FOLD_SPECS]
        # 上面使用的sigma就是480点主样本；显式记录训练尺度、保护边界和源行。
        candidates, searches, all_blocks, baselines = [], [], [], []
        for fold in folds:
            candidate, info = search_fold(fold, t0)
            candidates.append(candidate)
            searches.append(info)
            all_blocks.extend(score_blocks(fold, candidate))
            baselines.extend(baseline_blocks(fold))
            state.update({"运行状态": f"已完成第{fold['spec']['折号']}折", "实际用时秒": elapsed(t0),
                          "主指标已形成块数": len(all_blocks)})
            write_json(RESULT_DIR / "执行状态.json", state)
            write_json(RESULT_DIR / "留段预测_中间.json", {"分块": all_blocks, "搜索": searches})
        if len(all_blocks) != 8:
            raise RuntimeError("主方法未覆盖两折八个测试块")
        main_score = float(np.mean([item["标准化均方根误差"] for item in all_blocks]))
        coverage = float(np.mean([item["经验覆盖率"] for item in all_blocks]))
        baseline_summary = {}
        for name in ("训练均值", "训练二次趋势"):
            vals = [item["标准化均方根误差"] for item in baselines if item["基线"] == name]
            baseline_summary[name] = {"连续留段标准化均方根误差": float(np.mean(vals)), "块数": len(vals)}
        full_fold_obj = full_fold(sigma_full, values_full)
        full_candidate, full_search = search_fold(full_fold_obj, t0,
                                                  grid_size=STANDARD_GRID_SIZE,
                                                  max_starts=STANDARD_MAX_STARTS,
                                                  budget_seconds=STANDARD_SEARCH_BUDGET_SECONDS)
        full_profile = profile(full_fold_obj, full_candidate, t0)
        five_info, five_blocks = five_fold_validation(sigma_full, values_full, t0)
        five_summary = summarize_five_fold(five_info)
        sensitivity_rows = sensitivity(sigma_full, values_full, full_fold_obj, full_candidate,
                                       datasets, t0)
        fold_thickness = [float(c["参数"][0]) for c in candidates]
        five_thickness = [float(row["条件厚度_um"]) for row in five_info
                          if row.get("条件厚度有效", False)]
        sensitivity_thickness = [float(row["厚度_um"]) for row in sensitivity_rows
                                 if row.get("是否纳入条件范围", False)
                                 and row.get("固定参数一致", False)
                                 and row.get("厚度_um") is not None
                                 and math.isfinite(float(row["厚度_um"]))]
        condition_components = {
            "两折条件厚度_微米": fold_thickness,
            "五折条件厚度_微米": five_thickness,
            "光学剖面厚度范围_微米": full_profile["近等损失厚度范围_um"],
            "灵敏度情景厚度_微米": sensitivity_thickness,
        }
        condition_values = [value for values in condition_components.values() for value in values]
        if not condition_values:
            raise RuntimeError("没有形成可用于条件区间的合法厚度结果")
        condition_range = [float(min(condition_values)), float(max(condition_values))]
        formal_baseline_score = baseline_summary["训练二次趋势"]["连续留段标准化均方根误差"]
        formal_improvement = (100.0 * (formal_baseline_score - main_score) /
                              formal_baseline_score if formal_baseline_score else None)
        publication = assess_publication(full_candidate, condition_range, main_score,
                                         formal_baseline_score, coverage, five_summary)
        prototype_main_score = 1.1716972867113729
        prototype_baseline_score = 1.4403713129480724
        prototype_improvement = (100.0 * (prototype_baseline_score - prototype_main_score) /
                                 prototype_baseline_score)
        thickness_result = {
            "问题": 2, "材料": "碳化硅", "输入附件": ["附件1.xlsx", "附件2.xlsx"],
            "运行状态": "已完成", "同片最佳厚度_微米": float(full_candidate["参数"][0]),
            "同片最佳厚度的参考折射率": float(full_candidate["参数"][1]),
            "同片最佳厚度的经验色散系数": float(full_candidate["参数"][2]),
            "条件区间_微米": condition_range,
            "条件区间构造": "两折、五折连续留段、固定经验色散下厚度—参考折射率训练损失10%近等值剖面及标记为纳入条件范围的合法灵敏度情景并集；严格±3%请求不可行时不纳入，边界夹紧诊断单列；不是统计置信区间。",
            "条件区间组成": condition_components,
            "两折条件厚度_微米": fold_thickness,
            "五折条件厚度_微米": five_thickness,
            "灵敏度情景搜索完整性": [{"灵敏度项": row["灵敏度项"],
                                   "情景名称": row["情景名称"],
                                   "厚度_um": row["厚度_um"],
                                   "搜索状态": row["搜索状态"],
                                   "搜索未截断": row["搜索未截断"],
                                   "搜索未评估": row["搜索未评估"],
                                   "固定参数一致": row["固定参数一致"],
                                   "请求固定参考折射率": row["请求固定参考折射率"],
                                   "实际固定参考折射率": row["实际固定参考折射率"],
                                   "固定参考折射率是否夹紧": row["固定参考折射率是否夹紧"],
                                   "是否纳入条件范围": row["是否纳入条件范围"],
                                   "参考折射率扰动口径": row["参考折射率扰动口径"]}
                                  for row in sensitivity_rows],
            "全量训练标准化损失": float(full_candidate["损失"]),
            "全量拟合参数": full_search,
            "光学情景": {"模型": "透明两束近似", "折射率": "n(σ)=n参+c[(2000/σ)^2−1]，为经验参数化，不是材料常数",
                         "传播约束": "n(σ)>sin(15°)"},
            "五折跨切分稳定性": five_summary,
            "答案发布判定": publication,
            "假设检验结果键": HYPOTHESIS_RESULT_KEYS,
            "不确定性说明": five_summary["结论口径"] + "实测无折射率曲线、重复测量和真实厚度；条件区间同时保留跨切分、光学剖面和明确物理口径的灵敏度变化。反射率区间覆盖不转写为厚度置信度。",
            "正式结论句": "同片最佳厚度是条件点估计；五折连续留段若标记不稳定，则按低置信、跨切分不稳定表述，并保留完整条件范围。",
            "实际用时秒": elapsed(t0), "原型留段主指标复核值": 1.1716972867113729}
        write_json(RESULT_DIR / "厚度结果.json", thickness_result)
        validation = {"问题": 2, "主方法": "双角色散变投影",
                      "正式主方法分数_连续留段标准化均方根误差": main_score,
                      "正式二次趋势基线分数_连续留段标准化均方根误差": formal_baseline_score,
                      "连续留段标准化均方根误差": main_score,
                      "主指标含义": "两折、两角、每折两个40点测试块共八块的RMSE/本折本角训练残差IQR等权平均；越小越好，不是厚度误差。",
                      "原始反射率RMSE_比例均值": float(np.mean([x["均方根误差_比例"] for x in all_blocks])),
                      "经验覆盖率": {"构造方法": "校准块绝对残差的75%经验分位对称区间；不宣称有限样本名义覆盖保证",
                                  "校准经验分位点": CALIBRATION_QUANTILE, "名义覆盖率": None,
                                  "测试经验覆盖率": coverage,
                                  "测试点数": int(sum(x["点数"] for x in all_blocks)),
                                  "覆盖不足": coverage < CALIBRATION_QUANTILE,
                                  "覆盖解释": "仅反射率点预测覆盖；这是校准分位区间的留出诊断，不是厚度覆盖。"},
                      "基线对比": baseline_summary,
                      "三路线原型对比": {"双角色散变投影": prototype_main_score,
                                      "双角峰序匹配": 1.2354246858438795, "双角相位回归": 1.2201909148889387,
                                      "二次趋势诊断基线": prototype_baseline_score},
                      "正式主方法相对正式二次趋势改善_百分比": formal_improvement,
                      "改善率计算": "100×(正式二次趋势基线分数_连续留段标准化均方根误差−正式主方法分数_连续留段标准化均方根误差)/正式二次趋势基线分数_连续留段标准化均方根误差；两项均取本文件字段",
                      "原型相对二次趋势改善_百分比": prototype_improvement,
                      "原型改善率口径": "仅由三路线原型对比中的双角色散变投影与二次趋势诊断基线计算，不代表正式留段改善",
                      "答案发布判定": publication,
                      "假设检验结果键": HYPOTHESIS_RESULT_KEYS,
                      "分块": all_blocks, "基线分块": baselines, "搜索": searches,
                      "双向留角度冻结诊断": frozen_angle_checks(folds, candidates),
                      "五折连续留段": five_info, "五折跨切分稳定性": five_summary,
                      "五折测试块数": len(five_blocks),
                      "窗口与模型灵敏度": sensitivity_rows,
                      "实际用时秒": elapsed(t0)}
        write_json(RESULT_DIR / "验证.json", validation)
        write_json(RESULT_DIR / "区间覆盖.json", validation["经验覆盖率"])
        profile_rows = []
        for row in full_profile["数值"]:
            profile_rows.append(row)
        write_csv(RESULT_DIR / "参数剖面.csv", profile_rows,
                  ["厚度_um", "参考折射率", "色散系数", "训练标准化损失"])
        write_csv(RESULT_DIR / "灵敏度.csv", sensitivity_rows,
                  ["灵敏度项", "情景名称", "波数窗口_cm^-1", "厚度_um", "参考折射率", "色散系数",
                   "训练标准化损失", "搜索", "搜索状态", "搜索未截断", "搜索未评估",
                   "固定参考折射率", "请求固定参考折射率", "实际固定参考折射率",
                   "固定参考折射率是否夹紧", "固定参数策略", "固定参数一致",
                   "请求相对变化", "实际相对变化",
                   "参考折射率扰动口径", "是否纳入条件范围", "未评估原因",
                   "基线阶数", "幅值阶数", "色散形式", "厚度网格点数",
                   "搜索设置与正式主方法一致"])
        new_end = end_window_check(datasets, full_candidate, full_fold_obj)
        write_csv(RESULT_DIR / "留段预测.csv",
                  [{"类型": "主方法测试块", "折号": row["折号"], "角度_度": row["角度_度"],
                    "测试块": row["测试块"], "波数_cm^-1": sigma, "观测反射率_比例": obs,
                    "预测反射率_比例": pred, "残差_比例": res}
                   for row in all_blocks
                   for sigma, obs, pred, res in zip(row["波数_cm^-1"], row["观测反射率_比例"],
                                                    row["预测反射率_比例"], row["残差_比例"])] +
                  [{"类型": "末端新留段", "角度_度": row["角度_度"], "测试块": "3800-4000.122",
                    "波数_cm^-1": "", "观测反射率_比例": "", "预测反射率_比例": "", "残差_比例": "",
                    "备注": json.dumps(row, ensure_ascii=False)} for row in new_end],
                  ["类型", "折号", "角度_度", "测试块", "波数_cm^-1", "观测反射率_比例",
                   "预测反射率_比例", "残差_比例", "备注"])
        write_json(RESULT_DIR / "基准交接.json", {
            "问题": 2, "用途": "供问题三逐项复算碳化硅条件基准，不将原型数值当作实测真值。",
            "输入哈希": audit["文件"], "共同窗口_cm^-1": [1200, 3800],
            "抽样点数_每角度": 480, "抽样源行": source_rows, "折分": [fold["spec"] for fold in folds],
            "训练保护带_cm^-1": 20, "主模型": "双角色散变投影",
            "全量参数": {"厚度_um": full_candidate["参数"][0], "参考折射率": full_candidate["参数"][1],
                         "色散系数": full_candidate["参数"][2]},
            "两折条件厚度_um": fold_thickness, "条件范围_um": condition_range,
            "分块评分": all_blocks, "五折连续留段": five_info,
            "五折跨切分稳定性": five_summary,
            "灵敏度情景": sensitivity_rows,
            "末端新留段": new_end, "评价单位": "反射率比例；原始附件B列百分数除以100",
            "区间": validation["经验覆盖率"], "真实数据无合成输入": True,
            "实际用时秒": elapsed(t0)})
        validation["末端新留段"] = new_end
        validation["实际用时秒"] = elapsed(t0)
        write_json(RESULT_DIR / "验证.json", validation)
        state = {"问题": 2, "运行状态": "已完成", "实际用时秒": elapsed(t0), "结果文件": [
            "厚度结果.json", "基准交接.json", "验证.json", "参数剖面.csv", "灵敏度.csv", "留段预测.csv", "区间覆盖.json"]}
        write_json(RESULT_DIR / "执行状态.json", state)
        print(json.dumps({"问题": 2, "运行状态": "已完成", "同片最佳厚度_微米": thickness_result["同片最佳厚度_微米"],
                          "条件区间_微米": condition_range, "主指标": main_score,
                          "经验覆盖率": coverage, "实际用时秒": elapsed(t0)}, ensure_ascii=False), flush=True)
    except Exception as error:
        state = {"问题": 2, "运行状态": "异常结束但已保存阶段结果", "实际用时秒": elapsed(t0),
                 "错误类型": type(error).__name__, "错误": str(error)}
        write_json(RESULT_DIR / "执行状态.json", state)
        print(json.dumps(state, ensure_ascii=False), flush=True)
        raise


if __name__ == "__main__":
    main()
