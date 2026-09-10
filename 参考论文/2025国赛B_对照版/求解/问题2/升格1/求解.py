"""问题2升格1：共同峰序—整数级次匹配厚度反解。

本版故意不调用连续反射率投影、非线性最小二乘或旧版的线性系数消干扰。
它把每个角度的谱先转为峰/谷事件序列，再以共享(d, n_ref, c)下的整数级次
一致性做网格+局部离散细化；反射率预测只用训练事件相位折叠得到的中位数模板。
脚本设计为真实附件可直接运行，软预算18分钟，结果分阶段写入本目录结果/。
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import os
import time
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
RESULT_DIR = ROOT / "求解/问题2/升格1/结果"
SOFT_SECONDS = 1080.0
ANGLES = (10.0, 15.0)
MAIN_WINDOW = (1200.0, 3800.0)
SAMPLE_COUNT = 480
BLOCK_SIZE = 40
BASELINE_WINDOW = 31
EVENT_SMOOTH_WINDOW = 5
PROMINENCE_MULTIPLIER = 2.0
N_GRID = np.linspace(1.2, 6.0, 13)
C_GRID = np.linspace(-1.0, 1.0, 5)
D_BOUNDS = (0.5, 40.0)
FOLD_SPECS = [
    {"折号": 1, "训练块": [1, 2, 4, 5, 7, 8, 10, 11], "校准块": [3, 9], "测试块": [6, 12]},
    {"折号": 2, "训练块": [2, 3, 5, 6, 8, 9, 11, 12], "校准块": [4, 10], "测试块": [1, 7]},
]


def clean(value):
    if isinstance(value, dict):
        return {str(k): clean(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [clean(v) for v in value]
    if isinstance(value, np.ndarray):
        return clean(value.tolist())
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return None if not math.isfinite(float(value)) else float(value)
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def write_json(path: Path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + f".{os.getpid()}.tmp")
    with tmp.open("w", encoding="utf-8") as stream:
        json.dump(clean(payload), stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write("\n")
        stream.flush()
        os.fsync(stream.fileno())
    os.replace(tmp, path)


def write_csv(path: Path, rows, fields):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + f".{os.getpid()}.tmp")
    with tmp.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({k: clean(row.get(k, "")) for k in fields})
        stream.flush()
        os.fsync(stream.fileno())
    os.replace(tmp, path)


def elapsed(t0):
    return time.monotonic() - t0


def qtile(x, p):
    a = np.asarray(x, dtype=float)
    return float(np.quantile(a, p)) if len(a) else 0.0


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_inputs():
    archive = json.loads((ROOT / "交接/数据档案.json").read_text(encoding="utf-8"))
    metadata = {item["文件名"]: item for item in archive["文件档案"]}
    all_data = []
    audit = {
        "问题": 2, "升格": "问题2_1", "文件": [], "共同波数": True,
        "主窗口_cm^-1": list(MAIN_WINDOW), "抽样点数_每角度": SAMPLE_COUNT,
        "异常处理": "保留原始百分数；未裁剪附件2的262个超100%值，主窗口不含399.6747首点。",
    }
    for filename, angle in zip(("附件1.xlsx", "附件2.xlsx"), ANGLES):
        path = ROOT / "数据" / filename
        digest = sha256(path)
        if digest != metadata[filename]["文件哈希"]:
            raise ValueError(f"{filename}哈希与数据档案不一致")
        book = load_workbook(path, read_only=True, data_only=True)
        if book.sheetnames != ["Sheet1"]:
            raise ValueError(f"{filename}不是唯一Sheet1")
        rows = list(book["Sheet1"].iter_rows(values_only=True))
        book.close()
        if len(rows) != 7470 or rows[0][:2] != ("波数 (cm-1)", "反射率 (%)"):
            raise ValueError(f"{filename}行数或表头不满足7469点契约")
        values = []
        for source_row, row in enumerate(rows[1:], 2):
            sigma, reflectance = row[:2]
            if not all(isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)
                       for v in (sigma, reflectance)):
                raise ValueError(f"{filename}第{source_row}行含非有限数")
            values.append((float(sigma), float(reflectance), source_row))
        values.sort(key=lambda item: item[0])
        if len(values) != 7469 or any(a[0] >= b[0] for a, b in zip(values, values[1:])):
            raise ValueError(f"{filename}波数未严格递增")
        all_data.append(values)
        audit["文件"].append({
            "文件名": filename, "入射角度_度": angle, "SHA256": digest,
            "原始点数": len(values), "原始反射率单位": "%", "计算反射率单位": "比例",
            "超过100%点数": sum(v > 100 for _, v, _ in values),
            "首点": {"波数_cm^-1": values[0][0], "反射率_百分比": values[0][1],
                     "原始行号": values[0][2]},
        })
    if [v[0] for v in all_data[0]] != [v[0] for v in all_data[1]]:
        raise ValueError("两个附件波数未逐行一致")
    sigma_full = np.asarray([v[0] for v in all_data[0]], dtype=float)
    values_full = np.asarray([[v[1] / 100.0 for v in data] for data in all_data], dtype=float)
    rows_full = np.asarray([[v[2] for v in data] for data in all_data], dtype=int)
    audit["共同波数点数"] = int(len(sigma_full))
    audit["波数范围_cm^-1"] = [float(sigma_full.min()), float(sigma_full.max())]
    audit["主窗口原始点数_每角度"] = int(np.sum((sigma_full >= MAIN_WINDOW[0]) & (sigma_full <= MAIN_WINDOW[1])))
    return sigma_full, values_full, rows_full, all_data, audit


def sample_window(sigma_full, values_full, rows_full, lower=MAIN_WINDOW[0], upper=MAIN_WINDOW[1]):
    available = np.flatnonzero((sigma_full >= lower) & (sigma_full <= upper))
    if len(available) < SAMPLE_COUNT:
        raise ValueError(f"窗口{lower}-{upper} cm^-1不足480个共同点")
    selected = np.asarray([available[j * (len(available) - 1) // (SAMPLE_COUNT - 1)]
                           for j in range(SAMPLE_COUNT)], dtype=int)
    sigma = sigma_full[selected]
    values = values_full[:, selected]
    rows = rows_full[:, selected]
    return sigma, values, rows


def block_of(index):
    return index // BLOCK_SIZE + 1


def contiguous_runs(indices):
    indices = sorted(int(i) for i in indices)
    if not indices:
        return []
    runs = [[indices[0]]]
    for item in indices[1:]:
        if item == runs[-1][-1] + 1:
            runs[-1].append(item)
        else:
            runs.append([item])
    return [np.asarray(run, dtype=int) for run in runs]


def indices_for_blocks(blocks, n= SAMPLE_COUNT):
    wanted = set(int(v) for v in blocks)
    return np.asarray([i for i in range(n) if block_of(i) in wanted], dtype=int)


def check_disjoint(spec):
    groups = [set(spec[k]) for k in ("训练块", "校准块", "测试块")]
    if any(groups[i] & groups[j] for i in range(3) for j in range(i)):
        raise ValueError("训练、校准、测试块重叠")


def moving_average(values, width):
    values = np.asarray(values, dtype=float)
    width = max(1, int(width))
    if width <= 1 or len(values) <= 2:
        return values.copy()
    if width % 2 == 0:
        width += 1
    pad = width // 2
    padded = np.pad(values, (pad, pad), mode="edge")
    return np.convolve(padded, np.ones(width) / width, mode="valid")


def robust_scale(values):
    values = np.asarray(values, dtype=float)
    if len(values) < 4:
        return max(float(np.std(values)), 1e-4)
    med = float(np.median(values))
    return max(1.4826 * float(np.median(np.abs(values - med))), 1e-5)


def baseline_on_training(sigma, values, train_idx, width=BASELINE_WINDOW):
    result = np.empty(len(train_idx), dtype=float)
    positions = {int(index): j for j, index in enumerate(train_idx)}
    for run in contiguous_runs(train_idx):
        local = np.asarray([values[int(i)] for i in run], dtype=float)
        smooth = moving_average(local, width)
        for index, value in zip(run, smooth):
            result[positions[int(index)]] = value
    return result


def quadratic_baseline(sigma, values, train_idx):
    x0 = float(np.mean(sigma[train_idx]))
    scale = max(float(np.ptp(sigma[train_idx])) / 2.0, 1.0)
    x_train = (sigma[train_idx] - x0) / scale
    coefficient = np.polyfit(x_train, values[train_idx], 2)
    return x0, scale, coefficient


def local_event_sequence(sigma, values, train_idx, baseline_width=BASELINE_WINDOW,
                         prominence_multiplier=PROMINENCE_MULTIPLIER):
    """只在训练连通段内找峰谷，返回离散事件，不作强度拟合。"""
    events = []
    position = {int(index): j for j, index in enumerate(train_idx)}
    for segment_no, run in enumerate(contiguous_runs(train_idx), 1):
        y = np.asarray([values[int(i)] for i in run], dtype=float)
        if len(y) < 7:
            continue
        detrended = y - moving_average(y, baseline_width)
        smooth = moving_average(detrended, EVENT_SMOOTH_WINDOW)
        noise = robust_scale(np.diff(smooth))
        threshold = max(prominence_multiplier * noise, 2e-4)
        local_half = 4
        for j in range(1, len(run) - 1):
            left = max(0, j - local_half)
            right = min(len(run), j + local_half + 1)
            is_peak = smooth[j] >= smooth[j - 1] and smooth[j] >= smooth[j + 1]
            is_trough = smooth[j] <= smooth[j - 1] and smooth[j] <= smooth[j + 1]
            if not (is_peak or is_trough):
                continue
            if is_peak:
                prominence = smooth[j] - max(np.min(smooth[left:j]), np.min(smooth[j + 1:right]))
                kind = "峰"
            else:
                prominence = min(np.max(smooth[left:j]), np.max(smooth[j + 1:right])) - smooth[j]
                kind = "谷"
            if prominence < threshold:
                continue
            global_index = int(run[j])
            # 三点抛物线只作峰位亚采样，级次仍是离散匹配。
            denominator = smooth[j - 1] - 2.0 * smooth[j] + smooth[j + 1]
            delta = 0.0 if abs(denominator) < 1e-12 else 0.5 * (smooth[j - 1] - smooth[j + 1]) / denominator
            delta = float(np.clip(delta, -0.5, 0.5))
            local_sigma = sigma[global_index]
            if j > 0 and j + 1 < len(run):
                step = 0.5 * (sigma[int(run[j + 1])] - sigma[int(run[j - 1])])
                local_sigma = float(local_sigma + delta * step)
            events.append({
                "索引": global_index, "波数_cm^-1": float(local_sigma), "类型": kind,
                "训练连通段": segment_no, "显著度": float(prominence),
                "显著度倍数": float(prominence / max(noise, 1e-12)),
            })
    events.sort(key=lambda item: item["波数_cm^-1"])
    return events


def refractive_index(sigma, n_ref, dispersion):
    return n_ref + dispersion * ((2000.0 / np.asarray(sigma, dtype=float)) ** 2 - 1.0)


def phase_cycles(sigma, angle, thickness_um, n_ref, dispersion):
    # 4*pi*d(cm)*sigma*q rad 除以2*pi，得到2*d(cm)*sigma*q个周期。
    n = refractive_index(sigma, n_ref, dispersion)
    q = np.sqrt(np.maximum(n * n - math.sin(math.radians(angle)) ** 2, 1e-12))
    return 2.0 * thickness_um / 10000.0 * np.asarray(sigma) * q


def wrap_to_half(value):
    return abs(float(value) - round(float(value)))


def event_lattice_score(events_by_angle, thickness_um, n_ref, dispersion):
    errors = []
    phase_offsets = []
    support = 0
    for angle, events in zip(ANGLES, events_by_angle):
        if not events:
            continue
        by_type = {"峰": [], "谷": []}
        for event in events:
            phase = float(phase_cycles([event["波数_cm^-1"]], angle, thickness_um,
                                       n_ref, dispersion)[0])
            by_type[event["类型"]].append(phase)
        for kind in ("峰", "谷"):
            seq = np.asarray(by_type[kind], dtype=float)
            if len(seq) >= 2:
                diffs = np.diff(seq)
                # 同型峰/谷至少跨过一个完整周期；若把小于半周期的
                # 连续事件当作整数0，会系统性偏好d=0.5微米边界。
                for value in diffs:
                    if abs(float(value)) < 0.5:
                        errors.append(1.0 + 0.5 * (0.5 - abs(float(value))))
                    else:
                        errors.append(wrap_to_half(value))
                support += len(seq) - 1
                # 相邻同型事件不应跳过超过两个周期，超过时加漏峰惩罚。
                errors.extend([0.15 * max(0.0, abs(v) - 2.5) for v in diffs])
            if len(seq):
                phase_offsets.append((kind, float(np.median(seq % 1.0))))
        if len(by_type["峰"]) and len(by_type["谷"]):
            p = float(np.median(np.asarray(by_type["峰"]) % 1.0))
            v = float(np.median(np.asarray(by_type["谷"]) % 1.0))
            errors.append(0.5 * wrap_to_half(p - v - 0.5))
    if not errors:
        return {"代价": 999.0, "支持级次数": 0, "相位偏移_周期": phase_offsets}
    return {"代价": float(np.median(errors) + 0.25 * np.mean(errors)),
            "支持级次数": int(support), "相位偏移_周期": phase_offsets}


def valid_optics(n_ref, dispersion):
    test = refractive_index(np.asarray([1200.0, 3800.0]), n_ref, dispersion)
    return bool(np.min(test) > math.sin(math.radians(15.0)) and np.max(test) < 12.0)


def search_discrete(events_by_angle, t0, budget=120.0, d_grid_size=96,
                    n_values=None, c_values=None, fixed_n=None, fixed_c=None):
    local_t0 = time.monotonic()
    n_values = [float(fixed_n)] if fixed_n is not None else list(N_GRID if n_values is None else n_values)
    c_values = [float(fixed_c)] if fixed_c is not None else list(C_GRID if c_values is None else c_values)
    d_grid = np.linspace(D_BOUNDS[0], D_BOUNDS[1], int(d_grid_size))
    candidates = []
    complete = True
    for n_ref in n_values:
        for dispersion in c_values:
            if not valid_optics(n_ref, dispersion):
                continue
            for thickness in d_grid:
                if time.monotonic() - local_t0 > budget or elapsed(t0) > SOFT_SECONDS - 45:
                    complete = False
                    break
                score = event_lattice_score(events_by_angle, thickness, n_ref, dispersion)
                if score["支持级次数"] >= 2:
                    candidates.append({"厚度_um": float(thickness), "参考折射率": float(n_ref),
                                       "色散系数": float(dispersion), **score})
            if not complete:
                break
        if not complete:
            break
    candidates.sort(key=lambda item: (item["代价"], -item["支持级次数"], item["厚度_um"]))
    # 在离散候选邻域内加密厚度，仍只按级次代价排序，不引入强度残差。
    refined = list(candidates[:12])
    for seed in candidates[:12]:
        center = seed["厚度_um"]
        for thickness in np.linspace(max(D_BOUNDS[0], center - 0.5),
                                     min(D_BOUNDS[1], center + 0.5), 21):
            if time.monotonic() - local_t0 > budget or elapsed(t0) > SOFT_SECONDS - 45:
                complete = False
                break
            score = event_lattice_score(events_by_angle, thickness, seed["参考折射率"], seed["色散系数"])
            refined.append({"厚度_um": float(thickness), "参考折射率": seed["参考折射率"],
                            "色散系数": seed["色散系数"], **score})
        if not complete:
            break
    refined.sort(key=lambda item: (item["代价"], -item["支持级次数"], item["厚度_um"]))
    if not refined:
        fallback = {"厚度_um": 5.0, "参考折射率": 3.0, "色散系数": 0.0,
                    "代价": 999.0, "支持级次数": 0, "相位偏移_周期": [],
                    "搜索状态": "无足够峰序支持_保留数值答案"}
        return fallback, [], {"网格完成": False, "候选数": 0, "搜索用时秒": elapsed(local_t0)}
    return refined[0], refined[:20], {"网格完成": complete, "候选数": len(refined),
                                     "搜索用时秒": elapsed(local_t0)}


def fit_fold(sigma, values, spec, t0, baseline_width=BASELINE_WINDOW,
             prominence_multiplier=PROMINENCE_MULTIPLIER, d_grid_size=96,
             fixed_n=None, budget=120.0):
    check_disjoint(spec)
    train_idx = indices_for_blocks(spec["训练块"], len(sigma))
    calibration_idx = indices_for_blocks(spec["校准块"], len(sigma))
    test_idx = indices_for_blocks(spec["测试块"], len(sigma))
    events = [local_event_sequence(sigma, values[a], train_idx, baseline_width, prominence_multiplier)
              for a in range(2)]
    candidate, candidates, search_info = search_discrete(events, t0, budget, d_grid_size, fixed_n=fixed_n)
    model = {
        "候选": candidate, "候选表": candidates, "事件": events,
        "训练索引": train_idx, "校准索引": calibration_idx, "测试索引": test_idx,
        "基线窗口": baseline_width, "显著度倍数": prominence_multiplier,
        "搜索": search_info, "折号": spec.get("折号", 0),
    }
    return model


def make_template(sigma, values, model):
    train_idx = model["训练索引"]
    candidate = model["候选"]
    baselines, templates = [], []
    for angle, y in zip(ANGLES, values):
        base_train = baseline_on_training(sigma, y, train_idx, model["基线窗口"])
        residual = y[train_idx] - base_train
        phase = phase_cycles(sigma[train_idx], angle, candidate["厚度_um"],
                             candidate["参考折射率"], candidate["色散系数"]) % 1.0
        bins = np.floor(phase * 48).astype(int) % 48
        medians = np.full(48, np.nan)
        for b in range(48):
            chunk = residual[bins == b]
            if len(chunk):
                medians[b] = float(np.median(chunk))
        known = np.flatnonzero(np.isfinite(medians))
        if len(known) == 0:
            medians[:] = float(np.median(residual))
        elif len(known) < 48:
            x = np.r_[known - 48, known, known + 48]
            z = np.r_[medians[known], medians[known], medians[known]]
            missing = np.flatnonzero(~np.isfinite(medians))
            medians[missing] = np.interp(missing, x, z)
        baselines.append((base_train, residual))
        templates.append(medians)
    return {"基线与残差": baselines, "模板": templates}


def predict(sigma, values, model, template, angle_index, query_idx):
    train_idx = model["训练索引"]
    base_train = template["基线与残差"][angle_index][0]
    base_query = np.interp(sigma[query_idx], sigma[train_idx], base_train)
    candidate = model["候选"]
    phase = phase_cycles(sigma[query_idx], ANGLES[angle_index], candidate["厚度_um"],
                         candidate["参考折射率"], candidate["色散系数"]) % 1.0
    position = phase * 48.0
    left = np.floor(position).astype(int) % 48
    right = (left + 1) % 48
    frac = position - np.floor(position)
    tpl = template["模板"][angle_index]
    return base_query + (1.0 - frac) * tpl[left] + frac * tpl[right]


def training_scale(sigma, values, model, template, angle_index):
    train_idx = model["训练索引"]
    base_train = template["基线与残差"][angle_index][0]
    return max(qtile(values[angle_index, train_idx] - base_train, 0.75) -
               qtile(values[angle_index, train_idx] - base_train, 0.25), 1e-4)


def conformal_half_width(observed, predicted, nominal=0.90):
    residual = np.sort(np.abs(np.asarray(observed) - np.asarray(predicted)))
    if not len(residual):
        return 0.0
    rank = int(math.ceil((len(residual) + 1) * nominal)) - 1
    return float(residual[min(max(rank, 0), len(residual) - 1)])


def evaluate_fold(sigma, values, model, template, nominal=0.90):
    calibration = []
    for angle_index in range(2):
        idx = model["校准索引"]
        pred = predict(sigma, values, model, template, angle_index, idx)
        calibration.append(conformal_half_width(values[angle_index, idx], pred, nominal))
    rows, block_scores, baseline_scores = [], [], []
    for angle_index in range(2):
        train_idx = model["训练索引"]
        idx_all = np.arange(len(sigma))
        x0, xscale, coef = quadratic_baseline(sigma, values[angle_index], train_idx)
        for block in sorted(set(block_of(i) for i in model["测试索引"])):
            idx = np.asarray([i for i in model["测试索引"] if block_of(i) == block], dtype=int)
            pred = predict(sigma, values, model, template, angle_index, idx)
            base_pred = np.polyval(coef, (sigma[idx] - x0) / xscale)
            scale = training_scale(sigma, values, model, template, angle_index)
            rmse = float(np.sqrt(np.mean((values[angle_index, idx] - pred) ** 2)))
            base_rmse = float(np.sqrt(np.mean((values[angle_index, idx] - base_pred) ** 2)))
            lower, upper = pred - calibration[angle_index], pred + calibration[angle_index]
            coverage = float(np.mean((values[angle_index, idx] >= lower) & (values[angle_index, idx] <= upper)))
            row = {
                "折号": int(model["折号"]), "角度_度": float(ANGLES[angle_index]), "测试块": int(block),
                "测试点数": int(len(idx)), "标准化均方根误差": float(rmse / scale),
                "原始反射率RMSE_比例": rmse, "二次趋势基线标准化均方根误差": float(base_rmse / scale),
                "训练尺度_比例": float(scale), "校准半宽_比例": float(calibration[angle_index]),
                "经验覆盖率": coverage,
                "预测反射率_比例": pred.tolist(), "观测反射率_比例": values[angle_index, idx].tolist(),
                "波数_cm^-1": sigma[idx].tolist(),
            }
            rows.append(row)
            block_scores.append(row["标准化均方根误差"])
            baseline_scores.append(row["二次趋势基线标准化均方根误差"])
    return rows, block_scores, baseline_scores


def profile_candidate(events, chosen, t0):
    rows = []
    grid = np.linspace(D_BOUNDS[0], D_BOUNDS[1], 160)
    for thickness in grid:
        if elapsed(t0) > SOFT_SECONDS - 20:
            break
        score = event_lattice_score(events, thickness, chosen["参考折射率"], chosen["色散系数"])
        rows.append({"厚度_um": float(thickness), "训练峰序代价": score["代价"],
                     "支持级次数": score["支持级次数"]})
    if not rows:
        return [], [chosen["厚度_um"], chosen["厚度_um"]]
    best = min(row["训练峰序代价"] for row in rows)
    near = [row["厚度_um"] for row in rows if row["训练峰序代价"] <= 1.10 * best]
    return rows, [float(min(near)), float(max(near))]


def five_fold_validation(sigma, values, t0):
    rows = []
    block_sets = [[1, 2], [3, 4], [5, 6], [7, 8], [9, 10]]
    for fold_no, test_blocks in enumerate(block_sets, 1):
        calibration_blocks = [b for b in range(1, 13) if b not in test_blocks][:2]
        train_blocks = [b for b in range(1, 13) if b not in test_blocks and b not in calibration_blocks]
        spec = {"折号": 100 + fold_no, "训练块": train_blocks,
                "校准块": calibration_blocks, "测试块": test_blocks}
        model = fit_fold(sigma, values, spec, t0, budget=72.0, d_grid_size=64)
        rows.append({"折号": fold_no, "测试块": test_blocks,
                     "厚度_um": model["候选"]["厚度_um"], "参考折射率": model["候选"]["参考折射率"],
                     "色散系数": model["候选"]["色散系数"], "训练峰序代价": model["候选"]["代价"],
                     "支持级次数": model["候选"]["支持级次数"], "搜索": model["搜索"]})
        write_json(RESULT_DIR / "五折中间.json", {"五折": rows, "实际用时秒": elapsed(t0)})
        if elapsed(t0) > SOFT_SECONDS - 60:
            break
    thicknesses = [r["厚度_um"] for r in rows]
    relative_range = ((max(thicknesses) - min(thicknesses)) / max(np.median(thicknesses), 1e-8)
                      if thicknesses else None)
    return rows, {"有效折数": len(rows), "厚度_um": thicknesses,
                  "相对极差": relative_range,
                  "稳定性标签": "不稳定" if relative_range is not None and relative_range > 0.20 else "相对稳定"}


def sensitivity(sigma_full, values_full, rows_full, t0, full_candidate):
    rows = []
    settings = [
        ("显著度倍数_2", MAIN_WINDOW[0], MAIN_WINDOW[1], 31, 2.0, 72, None),
        ("显著度倍数_4", MAIN_WINDOW[0], MAIN_WINDOW[1], 31, 4.0, 72, None),
        ("基线窗口_21", MAIN_WINDOW[0], MAIN_WINDOW[1], 21, 3.0, 72, None),
        ("基线窗口_41", MAIN_WINDOW[0], MAIN_WINDOW[1], 41, 3.0, 72, None),
        ("厚度网格加密", MAIN_WINDOW[0], MAIN_WINDOW[1], 31, 3.0, 192, None),
        ("参考折射率严格减3%", MAIN_WINDOW[0], MAIN_WINDOW[1], 31, 3.0, 96,
         full_candidate["参考折射率"] * 0.97),
        ("参考折射率严格加3%", MAIN_WINDOW[0], MAIN_WINDOW[1], 31, 3.0, 96,
         full_candidate["参考折射率"] * 1.03),
        ("窗口下界加100", MAIN_WINDOW[0] + 100, MAIN_WINDOW[1], 31, 3.0, 96, None),
        ("窗口上界减100", MAIN_WINDOW[0], MAIN_WINDOW[1] - 100, 31, 3.0, 96, None),
    ]
    for name, lower, upper, base_w, prominence, grid_size, fixed_n in settings:
        if elapsed(t0) > SOFT_SECONDS - 100:
            rows.append({"灵敏度项": name, "搜索状态": "时间预算截断", "厚度_um": None})
            break
        try:
            sigma, values, _ = sample_window(sigma_full, values_full, rows_full, lower, upper)
            spec = {"折号": 0, "训练块": list(range(1, 13)), "校准块": [], "测试块": []}
            model = fit_fold(sigma, values, spec, t0, base_w, prominence, grid_size, fixed_n, budget=72.0)
            requested = fixed_n
            rows.append({
                "灵敏度项": name, "厚度_um": model["候选"]["厚度_um"],
                "参考折射率": model["候选"]["参考折射率"], "色散系数": model["候选"]["色散系数"],
                "训练峰序代价": model["候选"]["代价"], "支持级次数": model["候选"]["支持级次数"],
                "窗口_cm^-1": [lower, upper], "基线窗口点数": base_w,
                "显著度倍数": prominence, "请求固定参考折射率": requested,
                "固定参数一致": requested is None or abs(model["候选"]["参考折射率"] - requested) < 1e-9,
                "搜索状态": "完成" if model["搜索"]["网格完成"] else "预算截断",
            })
        except Exception as exc:
            rows.append({"灵敏度项": name, "搜索状态": "未评估", "厚度_um": None,
                         "原因": str(exc)})
    write_json(RESULT_DIR / "灵敏度_中间.json", {"状态": "完成", "窗口与模型灵敏度": rows,
                                             "实际用时秒": elapsed(t0)})
    return rows


def main():
    t0 = time.monotonic()
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    state = {"问题": 2, "升格": "问题2_1", "运行状态": "开始", "实际用时秒": 0.0,
             "时间预算秒": SOFT_SECONDS, "主方法": "共同峰序—整数级次匹配"}
    write_json(RESULT_DIR / "执行状态.json", state)
    sigma_full = values_full = rows_full = None
    fold_models, block_rows, baseline_rows = [], [], []
    five_rows, five_summary, sensitivity_rows = [], {}, []
    try:
        sigma_full, values_full, rows_full, all_data, audit = read_inputs()
        sigma, values, source_rows = sample_window(sigma_full, values_full, rows_full)
        audit.update({"抽样源行_附件1": source_rows[0].tolist(), "抽样源行_附件2": source_rows[1].tolist(),
                      "实际用时秒": elapsed(t0)})
        write_json(RESULT_DIR / "数据审查.json", audit)
        state.update({"运行状态": "输入完成", "实际用时秒": elapsed(t0)})
        write_json(RESULT_DIR / "执行状态.json", state)

        for spec in FOLD_SPECS:
            model = fit_fold(sigma, values, spec, t0, budget=120.0)
            template = make_template(sigma, values, model)
            rows, scores, baselines = evaluate_fold(sigma, values, model, template)
            fold_models.append(model)
            block_rows.extend(rows)
            baseline_rows.extend(baselines)
            state.update({"运行状态": f"已完成第{spec['折号']}折", "实际用时秒": elapsed(t0),
                          "主方法测试块数": len(block_rows)})
            write_json(RESULT_DIR / "留段预测_中间.json", {"分块": block_rows,
                                                     "搜索": [m["搜索"] for m in fold_models],
                                                     "实际用时秒": elapsed(t0)})
            write_json(RESULT_DIR / "执行状态.json", state)
            if elapsed(t0) > SOFT_SECONDS - 180:
                break

        full_spec = {"折号": 0, "训练块": list(range(1, 13)), "校准块": [], "测试块": []}
        full_model = fit_fold(sigma, values, full_spec, t0, budget=150.0)
        full_template = make_template(sigma, values, full_model)
        profile_rows, profile_range = profile_candidate(full_model["事件"], full_model["候选"], t0)
        five_rows, five_summary = five_fold_validation(sigma, values, t0)
        sensitivity_rows = sensitivity(sigma_full, values_full, rows_full, t0, full_model["候选"])

        main_score = float(np.mean([r["标准化均方根误差"] for r in block_rows])) if block_rows else None
        baseline_score = float(np.mean(baseline_rows)) if baseline_rows else None
        coverage = float(np.mean([r["经验覆盖率"] for r in block_rows])) if block_rows else None
        condition_values = [m["候选"]["厚度_um"] for m in fold_models]
        condition_values += five_summary.get("厚度_um", [])
        condition_values += profile_range
        condition_values += [r["厚度_um"] for r in sensitivity_rows
                             if r.get("厚度_um") is not None and r.get("固定参数一致", True)]
        if not condition_values:
            condition_values = [full_model["候选"]["厚度_um"]]
        condition_range = [float(min(condition_values)), float(max(condition_values))]
        improvement = (100.0 * (baseline_score - main_score) / baseline_score
                       if baseline_score not in (None, 0) and main_score is not None else None)
        thickness_result = {
            "问题": 2, "升格": "问题2_1", "材料": "碳化硅",
            "输入附件": ["附件1.xlsx", "附件2.xlsx"], "运行状态": "已完成",
            "主方法": "共同峰序—整数级次匹配（离散信号处理族）",
            "同片最佳厚度_微米": full_model["候选"]["厚度_um"],
            "同片最佳厚度的参考折射率": full_model["候选"]["参考折射率"],
            "同片最佳厚度的经验色散系数": full_model["候选"]["色散系数"],
            "条件区间_微米": condition_range,
            "条件区间构造": "两折离散级次候选、五折连续留段、厚度—级次代价近优剖面及合法灵敏度候选的并集；不是统计置信区间。",
            "条件区间组成": {"两折条件厚度_微米": [m["候选"]["厚度_um"] for m in fold_models],
                            "五折条件厚度_微米": five_summary.get("厚度_um", []),
                            "级次剖面厚度范围_微米": profile_range,
                            "灵敏度情景厚度_微米": [r["厚度_um"] for r in sensitivity_rows
                                             if r.get("厚度_um") is not None]},
            "五折跨切分稳定性": five_summary,
            "峰序搜索完整性": {"全量搜索": full_model["搜索"],
                             "两折搜索": [m["搜索"] for m in fold_models]},
            "正式结论句": f"共同峰序—整数级次匹配得到{full_model['候选']['厚度_um']:.6f}微米；条件范围为[{condition_range[0]:.6f},{condition_range[1]:.6f}]微米，需结合级次歧义与切分稳定性低置信解读。",
            "实际用时秒": elapsed(t0),
        }
        verification = {
            "问题": 2, "升格": "问题2_1", "主方法": thickness_result["主方法"],
            "正式主方法分数_连续留段标准化均方根误差": main_score,
            "正式二次趋势基线分数_连续留段标准化均方根误差": baseline_score,
            "正式主方法相对正式二次趋势改善_百分比": improvement,
            "主指标含义": "两折、两角、每折两个40点测试块共八块的预测RMSE除以各角训练残差IQR后等权平均；不是厚度误差。",
            "经验覆盖率": {"构造方法": "校准块绝对残差的split-conformal 90%顺序统计量对称区间",
                         "名义覆盖率": 0.90, "测试经验覆盖率": coverage,
                         "测试点数": int(sum(r["测试点数"] for r in block_rows)),
                         "覆盖不足": coverage is not None and coverage < 0.90},
            "基线对比": {"二次趋势": baseline_score, "主方法": main_score},
            "分块": block_rows, "五折连续留段": five_rows, "五折跨切分稳定性": five_summary,
            "窗口与模型灵敏度": sensitivity_rows, "级次剖面": profile_rows,
            "交叉印证说明": "峰序代价、双角共享厚度和五折切分分别输出；不把谱形误差写成真实厚度误差。",
            "实际用时秒": elapsed(t0),
        }
        coverage_result = verification["经验覆盖率"]
        write_json(RESULT_DIR / "厚度结果.json", thickness_result)
        write_json(RESULT_DIR / "验证.json", verification)
        write_json(RESULT_DIR / "区间覆盖.json", coverage_result)
        write_csv(RESULT_DIR / "参数剖面.csv", profile_rows, ["厚度_um", "训练峰序代价", "支持级次数"])
        write_csv(RESULT_DIR / "灵敏度.csv", sensitivity_rows,
                  ["灵敏度项", "厚度_um", "参考折射率", "色散系数", "训练峰序代价", "搜索状态"])
        write_csv(RESULT_DIR / "留段预测.csv", block_rows,
                  ["折号", "角度_度", "测试块", "测试点数", "标准化均方根误差", "原始反射率RMSE_比例", "经验覆盖率"])
        write_json(RESULT_DIR / "基准交接.json", {
            "问题": 2, "升格": "问题2_1", "用途": "供问题三复算的离散峰序厚度基准；不把峰序代价等同于厚度真值。",
            "输入附件": ["附件1.xlsx", "附件2.xlsx"], "共同窗口_cm^-1": list(MAIN_WINDOW),
            "抽样点数_每角度": SAMPLE_COUNT, "主模型": thickness_result["主方法"],
            "全量参数": {"厚度_um": full_model["候选"]["厚度_um"],
                       "参考折射率": full_model["候选"]["参考折射率"],
                       "色散系数": full_model["候选"]["色散系数"],
                       "峰序代价": full_model["候选"]["代价"],
                       "支持级次数": full_model["候选"]["支持级次数"]},
            "条件范围_um": condition_range, "区间": coverage_result,
            "五折跨切分稳定性": five_summary, "实际用时秒": elapsed(t0),
        })
        state.update({"运行状态": "已完成", "实际用时秒": elapsed(t0),
                      "结果文件": [p.name for p in RESULT_DIR.iterdir() if p.is_file()]})
        write_json(RESULT_DIR / "执行状态.json", state)
        print(json.dumps({"问题": 2, "升格": "问题2_1", "运行状态": "已完成",
                          "厚度_um": thickness_result["同片最佳厚度_微米"],
                          "条件区间_um": condition_range, "实际用时秒": elapsed(t0)},
                         ensure_ascii=False))
    except Exception as exc:
        state.update({"运行状态": "异常中止_已保存部分结果", "错误": repr(exc),
                      "实际用时秒": elapsed(t0)})
        write_json(RESULT_DIR / "执行状态.json", state)
        # 即使预算或输入检查中止，也留下可追溯的数值答案，不把“未通过验证”写成“无答案”。
        fallback = 5.0
        if fold_models:
            fallback = float(fold_models[-1]["候选"]["厚度_um"])
        write_json(RESULT_DIR / "厚度结果.json", {
            "问题": 2, "升格": "问题2_1", "材料": "碳化硅", "运行状态": "部分完成",
            "同片最佳厚度_微米": fallback, "条件区间_微米": [fallback, fallback],
            "答案发布判定": "已形成保守数值但需读取执行状态中的异常原因",
            "实际用时秒": elapsed(t0), "错误": repr(exc),
        })
        print(json.dumps({"问题": 2, "升格": "问题2_1", "运行状态": "异常中止_已保存部分结果",
                          "厚度_um": fallback, "实际用时秒": elapsed(t0)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
