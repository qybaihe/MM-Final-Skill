"""问题3变体2：训练内分段背景投影与往返衰减场反演。

本文件可直接运行，不导入主线求解器，不修改主线结果或声明。
所有选择只使用相应训练集；原有外层测试只作探索性复查。
主流程1080秒停止扩展，1180秒保留最终写出余量。
"""

import hashlib
import json
import math
import os
import time
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
RESULT = Path(__file__).resolve().parent / "结果"
FROZEN = ROOT / "数据/问题3_冻结合成输入/升格2"
ANGLES = np.array([10.0, 15.0])
MATERIALS = {"硅": (3, 4), "碳化硅": (1, 2)}
MODELS = ("两束", "完整往返")
PARAMETERS = ("厚度_um", "参考折射率", "色散系数", "衬底折射率对比", "有效往返损耗")
BOUNDS = np.array([[0.5, 40.0], [1.2, 6.0], [-1.0, 1.0], [-1.5, 2.5], [0.0, 3.0]])
WIDTHS = (400.0, 650.0, 900.0)
DEFAULT_WIDTH = 650.0
GAIN_LIMIT = 20.0
RIDGE = 0.1
SOFT_SECONDS = 1080.0
HARD_SECONDS = 1180.0
SEED = 20260910
FOLDS = [
    {"折号": 1, "训练块": [1, 2, 4, 5, 7, 8, 10, 11], "校准块": [3, 9], "测试块": [6, 12]},
    {"折号": 2, "训练块": [2, 3, 5, 6, 8, 9, 11, 12], "校准块": [4, 10], "测试块": [1, 7]},
]


def plain(value):
    if isinstance(value, np.ndarray):
        return plain(value.tolist())
    if isinstance(value, np.generic):
        return plain(value.item())
    if isinstance(value, dict):
        return {str(key): plain(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [plain(item) for item in value]
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("结果出现非有限数，不能把数值故障伪装为空值")
    return value


def atomic_json(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f"{path.name}.{os.getpid()}.tmp")
    with temporary.open("w", encoding="utf-8") as stream:
        json.dump(plain(payload), stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write("\n")
        stream.flush()
        os.fsync(stream.fileno())
    os.replace(temporary, path)


class NoFeasibleCandidate(ValueError):
    pass


class Budget:
    def __init__(self):
        self.t0 = time.monotonic()

    def elapsed(self):
        return time.monotonic() - self.t0

    def deadline(self, seconds):
        return min(self.t0 + SOFT_SECONDS, time.monotonic() + seconds)

    def available(self, reserve=0.0):
        return self.elapsed() < SOFT_SECONDS - reserve

    def write(self, filename, payload):
        atomic_json(RESULT / filename, {**payload, "实际用时秒": self.elapsed()})


def freeze_array(filename, arrays, purpose, generation, parameters, manifest):
    FROZEN.mkdir(parents=True, exist_ok=True)
    path = FROZEN / filename
    temporary = path.with_name(f"{filename}.{os.getpid()}.tmp")
    with temporary.open("wb") as stream:
        np.savez_compressed(stream, **arrays)
        stream.flush()
        os.fsync(stream.fileno())
    os.replace(temporary, path)
    manifest[filename] = {"用途": purpose, "生成方式": generation, "参数": parameters,
                          "SHA256": hashlib.sha256(path.read_bytes()).hexdigest()}
    atomic_json(FROZEN / "输入清单.json", {"问题": 3, "变体": 2, "文件": manifest})


def load_inputs(budget, manifest):
    archive = json.loads((ROOT / "交接/数据档案.json").read_text(encoding="utf-8"))
    baseline_path = ROOT / "求解/问题2/结果/基准交接.json"
    baseline = json.loads(baseline_path.read_text(encoding="utf-8"))
    if (baseline["共同窗口_cm^-1"] != [1200, 3800]
            or baseline["抽样点数_每角度"] != 480 or baseline["折分"] != FOLDS):
        raise ValueError("正式问题2的窗口、抽样或折分已改变，不能冒称同口径")
    expected = {entry["文件名"]: entry["文件哈希"] for entry in archive["文件档案"]}
    arrays, hashes = {}, {}
    for number in range(1, 5):
        path = ROOT / "数据" / f"附件{number}.xlsx"
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        if digest != expected[path.name]:
            raise ValueError(f"{path.name}与数据档案的SHA256不一致")
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            rows = list(workbook["Sheet1"].iter_rows(min_row=2, max_col=2, values_only=True))
        finally:
            workbook.close()
        if len(rows) != 7469 or any(not isinstance(value, (int, float)) for row in rows for value in row):
            raise ValueError(f"{path.name}不是7469行原始数值")
        values = np.asarray(rows, dtype=float)
        if not np.isfinite(values).all() or not np.all(np.diff(values[:, 0]) > 0):
            raise ValueError(f"{path.name}存在非有限数或非严格递增波数")
        arrays[number] = values
        hashes[path.name] = digest
    if any(not np.array_equal(arrays[1][:, 0], arrays[number][:, 0]) for number in (2, 3, 4)):
        raise ValueError("四附件原始波数不再逐行一致")
    if any(hashes[entry["文件名"]] != entry["SHA256"] for entry in baseline["输入哈希"]):
        raise ValueError("正式碳化硅基准不对应当前附件")
    sigma = arrays[1][:, 0]
    keep = np.flatnonzero((sigma >= 1200.0) & (sigma <= 3800.0))
    selected = keep[np.arange(480) * (len(keep) - 1) // 479]
    source_rows = (selected + 2).tolist()
    if baseline["抽样源行"] != [source_rows, source_rows]:
        raise ValueError("480点源行规则与上游不一致")
    public = {"共同窗口_cm^-1": [1200, 3800], "全量点数_每角度": len(keep),
              "抽样点数_每角度": 480, "抽样源行": source_rows, "折分": FOLDS,
              "训练保护带_cm^-1": 20.0, "原始附件SHA256": hashes,
              "正式碳化硅基准参数": baseline["全量参数"],
              "正式碳化硅条件范围_um": baseline["条件范围_um"],
              "上游基准SHA256": hashlib.sha256(baseline_path.read_bytes()).hexdigest(),
              "候选背景宽度_cm^-1": WIDTHS, "固定幅值岭系数": RIDGE,
              "物理参数顺序": PARAMETERS, "物理搜索边界": BOUNDS,
              "粗搜折射率": [2.0, 3.0, 4.0], "粗搜色散": [-0.5, 0.0, 0.5],
              "粗搜界面对比": 0.8, "粗搜损耗": 0.0,
              "随机种子": SEED, "核心指标不使用合成输入": True}
    frozen_path = FROZEN / "基准与协议.json"
    atomic_json(frozen_path, public)
    manifest[frozen_path.name] = {"用途": "独立复算的上游基准及预定选择规则",
        "生成方式": "读取已核验基准并复制必要数值，不拟合厚度", "参数": {"变体": 2},
        "SHA256": hashlib.sha256(frozen_path.read_bytes()).hexdigest()}
    atomic_json(FROZEN / "输入清单.json", {"问题": 3, "变体": 2, "文件": manifest})
    budget.write("数据审查.json", {**public, "预处理": "原谱除以100；不裁剪、不删异常、不插值观测",
        "共同首点在窗口内": bool(keep[0] == 0),
        "窗口内超百分之百点数": {f"附件{number}": int(np.sum(arrays[number][keep, 1] > 100))
                                    for number in arrays},
        "指标含义": "源行含表头偏移；正式共同窗口每角5392点，点数仍由附件实际核算"})
    return {"arrays": arrays, "keep": keep, "selected": selected, "baseline": baseline, "public": public}


def make_case(data, material, fold_number):
    indices = data["keep"] if fold_number == 0 else data["selected"]
    sigma = data["arrays"][1][indices, 0]
    values = np.stack([data["arrays"][number][indices, 1] / 100 for number in MATERIALS[material]])
    if fold_number == 0:
        specification = {"折号": 0, "训练块": [], "校准块": [], "测试块": []}
        train = np.arange(len(sigma))
    else:
        specification = FOLDS[fold_number - 1]
        edges = (sigma[39:479:40] + sigma[40:480:40]) / 2
        blocks = np.searchsorted(edges, sigma, side="right") + 1
        training_blocks = set(specification["训练块"])
        guards = [edge for block, edge in enumerate(edges, 1)
                  if (block in training_blocks) != (block + 1 in training_blocks)]
        allowed = np.isin(blocks, list(training_blocks))
        for edge in guards:
            allowed &= np.abs(sigma - edge) >= 20.0
        train = np.flatnonzero(allowed)
    return {"material": material, "fold": fold_number, "sigma": sigma, "values": values,
            "rows": indices + 2, "train": train, "spec": specification,
            "angles": ANGLES.copy(), "weight": 0.5, "ridge": RIDGE}


def prepare(case, width=DEFAULT_WIDTH):
    prepared = {**case, "width": float(width)}
    sigma, train = case["sigma"], case["train"]
    group_labels = np.floor((sigma[train] - 1200.0) / width).astype(int)
    occupied, inverse = np.unique(group_labels, return_inverse=True)
    counts = np.bincount(inverse).astype(float)
    centers = np.bincount(inverse, weights=sigma[train]) / counts
    right = np.searchsorted(centers, sigma, side="right")
    left = np.clip(right - 1, 0, len(centers) - 1)
    right = np.clip(right, 0, len(centers) - 1)
    difference = centers[right] - centers[left]
    weight = np.divide(sigma - centers[left], difference, out=np.zeros_like(sigma), where=difference > 0)
    weight = np.clip(weight, 0.0, 1.0)
    prepared["operator"] = {"inverse": inverse, "counts": counts, "centers": centers,
                            "left": left, "right": right, "weight": weight,
                            "occupied": occupied}
    prepared["background"] = project(case["values"], prepared)
    coordinate = (sigma - 2500.0) / 1300.0
    design = np.column_stack((np.ones(len(sigma)), coordinate, coordinate ** 2))
    gram = design[train].T @ design[train] + np.eye(3) * 1e-12
    coefficient = np.linalg.solve(gram, design[train].T @ case["values"][:, train].T)
    prepared["quadratic"] = (design @ coefficient).T
    prepared["mean"] = np.repeat(np.mean(case["values"][:, train], axis=1)[:, None], len(sigma), axis=1)
    residual = case["values"][:, train] - prepared["quadratic"][:, train]
    prepared["scale"] = np.maximum(np.quantile(residual, 0.75, axis=1)
                                     - np.quantile(residual, 0.25, axis=1), 0.0001)
    return prepared


def project(series, case):
    operator = case["operator"]
    original_shape = series.shape
    flat = np.asarray(series).reshape(-1, original_shape[-1])
    output = []
    for vector in flat:
        means = np.bincount(operator["inverse"], weights=vector[case["train"]]) / operator["counts"]
        output.append((1.0 - operator["weight"]) * means[operator["left"]]
                      + operator["weight"] * means[operator["right"]])
    return np.asarray(output).reshape(original_shape)


def physical(parameters, sigma, angles=ANGLES, weight=0.5):
    parameters = np.asarray(parameters)
    if np.any(parameters < BOUNDS[:, 0]) or np.any(parameters > BOUNDS[:, 1]):
        return None
    thickness, reference, dispersion, contrast, loss = parameters
    film = reference + dispersion * ((2000.0 / sigma) ** 2 - 1.0)
    substrate = film + contrast
    sine_squared = np.sin(np.deg2rad(angles))[:, None] ** 2
    endpoints = reference + dispersion * ((2000.0 / np.array([1200.0, 3800.0])) ** 2 - 1.0)
    if min(np.min(endpoints), np.min(endpoints + contrast), np.min(film), np.min(substrate)) <= np.sqrt(np.max(sine_squared)) + 1e-8:
        return None
    air = np.cos(np.deg2rad(angles))[:, None]
    layer = np.sqrt(film ** 2 - sine_squared)
    support = np.sqrt(substrate ** 2 - sine_squared)
    propagation = np.exp(-loss * film / layer + 4j * np.pi * thickness * sigma * layer / 10000.0)
    fields = []
    for eta_air, eta_layer, eta_support in ((air, layer, support),
            (1.0 / air, film ** 2 / layer, substrate ** 2 / support)):
        reflection = (eta_air - eta_layer) / (eta_air + eta_layer)
        back = (eta_layer - eta_support) / (eta_layer + eta_support)
        first = 4.0 * eta_air * eta_layer / (eta_air + eta_layer) ** 2 * back * propagation
        ratio = -reflection * back * propagation
        if np.max(np.abs(ratio)) >= 1.0:
            return None
        fields.append((reflection, first, ratio))
    spectra = {}
    for model in MODELS:
        intensities = [np.abs(reflection + (first if model == "两束" else first / (1 - ratio))) ** 2
                       for reflection, first, ratio in fields]
        spectra[model] = weight * intensities[0] + (1 - weight) * intensities[1]
    return {"spectra": spectra, "fields": fields, "attenuation": np.abs(propagation)}


def evaluate(parameters, case, model, optics=None):
    optics = physical(parameters, case["sigma"], case["angles"], case["weight"]) if optics is None else optics
    if optics is None:
        return None
    spectrum = optics["spectra"][model]
    feature = spectrum - project(spectrum, case)
    train = case["train"]
    target = case["values"] - case["background"]
    norm = np.sum(feature[:, train] ** 2, axis=1)
    ridge_penalty = case["ridge"] * len(train) * case["scale"] ** 2
    gains = np.clip(np.sum(feature[:, train] * target[:, train], axis=1)
                    / np.maximum(norm + ridge_penalty, 1e-24), 0.0, GAIN_LIMIT)
    prediction = case["background"] + gains[:, None] * feature
    residual = (prediction[:, train] - case["values"][:, train]) / case["scale"][:, None]
    loss = float(np.mean(residual ** 2) + case["ridge"] * np.mean(gains ** 2))
    return {"parameters": np.array(parameters, dtype=float), "prediction": prediction,
            "gain": gains, "loss": loss, "feature_norm": norm,
            "unidentified": bool(np.max(np.abs(gains)) <= 1e-10 or np.max(norm) < 1e-16)}


def retain(pool, candidate, limit=3):
    if candidate is not None:
        pool.append(candidate)
        pool.sort(key=lambda item: (item["loss"], item["parameters"][0]))
        del pool[limit:]


def search_pair(case, budget, deadline, grid_count=64, rounds=8, seeds=None, fixed=None):
    fixed = {} if fixed is None else fixed
    pools = {model: [] for model in MODELS}
    counts = {model: 0 for model in MODELS}
    completed_grid = 0
    trace = {"粗网格计划次数_每模型": grid_count * 9, "局部计划轮数_每起点": rounds,
             "固定参数": {PARAMETERS[index]: value for index, value in fixed.items()}}
    proposed = [np.array([10.0, 3.0, 0.0, 0.8, 0.0])]
    if seeds:
        proposed.extend(np.asarray(seed).copy() for seed in seeds)
    for reference in (2.0, 3.0, 4.0):
        for dispersion in (-0.5, 0.0, 0.5):
            proposed.extend(np.array([thickness, reference, dispersion, 0.8, 0.0])
                            for thickness in np.linspace(0.5, 40.0, grid_count))
    planned = len(proposed)
    for parameters in proposed:
        if all(pools.values()) and time.monotonic() >= deadline:
            break
        parameters = parameters.copy()
        for index, value in fixed.items():
            parameters[index] = value
        optics = physical(parameters, case["sigma"], case["angles"], case["weight"])
        for model in MODELS:
            counts[model] += 1
            retain(pools[model], evaluate(parameters, case, model, optics) if optics else None)
        completed_grid += 1
    if not all(pools.values()):
        raise NoFeasibleCandidate("请求的物理固定条件下无共同可行候选")
    best = {model: pools[model][0] for model in MODELS}
    union = []
    for model in MODELS:
        for candidate in pools[model][:2]:
            parameters = candidate["parameters"].tolist()
            if parameters not in union:
                union.append(parameters)
    refinements = []
    refined_candidates = []
    for shared_seed in union:
        if time.monotonic() >= deadline:
            break
        current = {model: evaluate(shared_seed, case, model) for model in MODELS}
        for model in MODELS:
            counts[model] += 1
        thickness_step = 39.5 / (grid_count - 1) if grid_count > 1 else max(0.04, shared_seed[0] * 0.05)
        steps = np.array([thickness_step, 0.15, 0.06, 0.12, 0.12])
        completed_rounds = 0
        for round_number in range(rounds):
            if time.monotonic() >= deadline:
                break
            for dimension in range(5):
                if dimension in fixed:
                    continue
                centers = {model: current[model]["parameters"].copy() for model in MODELS}
                for direction in (-1, 1):
                    for model in MODELS:
                        proposal = centers[model].copy()
                        proposal[dimension] = np.clip(proposal[dimension] + direction * steps[dimension],
                                                       *BOUNDS[dimension])
                        candidate = evaluate(proposal, case, model)
                        counts[model] += 1
                        if candidate is not None and candidate["loss"] < current[model]["loss"]:
                            current[model] = candidate
            steps *= 0.64
            completed_rounds = round_number + 1
        for model in MODELS:
            if current[model]["loss"] < best[model]["loss"]:
                best[model] = current[model]
            refined_candidates.append({"模型": model, "厚度_um": current[model]["parameters"][0],
                                       "训练标准化损失": current[model]["loss"]})
        refinements.append({"共同初值": shared_seed, "两模型共同完成轮数": completed_rounds})
    trace.update({"共同粗候选计划数": planned, "共同粗候选已算数": completed_grid,
        "目标评估次数": counts, "评估次数相同": counts["两束"] == counts["完整往返"],
        "共同精修初值": union, "精修记录": refinements, "精修候选": refined_candidates,
        "全部计划完成": completed_grid == planned and len(refinements) == len(union)
            and all(entry["两模型共同完成轮数"] == rounds for entry in refinements),
        "公平范围": "同候选网格、共同初值并集、同自由度及逐轮等次数；有限搜索不等于全局最优"})
    return {"case": case, "best": best, "trace": trace}


def inner_selection(case, budget, deadline):
    allowed = case["train"]
    if len(allowed) > 480:
        allowed = allowed[np.arange(480) * (len(allowed) - 1) // 479]
    size = len(allowed)
    sampled = {**case, "sigma": case["sigma"][allowed], "values": case["values"][:, allowed],
               "rows": case["rows"][allowed], "train": np.arange(size)}
    diagnostic = {"候选": [], "内层切分": [], "选择读取": "仅外层训练反射率",
                  "外测试用途": "仅探索性复查，不选窗口、特征、参数或物理模型"}
    interiors = []
    for start in (size // 4, 5 * size // 8):
        validation = np.arange(start, start + max(8, size // 8))
        lower, upper = sampled["sigma"][validation[[0, -1]]]
        train = np.flatnonzero((sampled["sigma"] < lower - 20.0) | (sampled["sigma"] > upper + 20.0))
        interiors.append((train, validation))
        diagnostic["内层切分"].append({"训练源行": sampled["rows"][train], "验证源行": sampled["rows"][validation]})
    for width in WIDTHS:
        scores = {model: [] for model in MODELS}
        searches_complete = True
        for train, validation in interiors:
            if time.monotonic() >= deadline:
                searches_complete = False
                break
            inner = prepare({**sampled, "train": train}, width)
            fit = search_pair(inner, budget, deadline, grid_count=24, rounds=3)
            searches_complete &= fit["trace"]["全部计划完成"]
            for model in MODELS:
                differences = (fit["best"][model]["prediction"][:, validation]
                               - inner["values"][:, validation]) / inner["scale"][:, None]
                scores[model].append(float(np.mean(np.sqrt(np.mean(differences ** 2, axis=1)))))
        diagnostic["候选"].append({"背景宽度_cm^-1": width, "各模型各内折误差": scores,
                                  "完整": searches_complete and all(len(values) == 2 for values in scores.values())})
    complete = len(diagnostic["候选"]) == len(WIDTHS) and all(entry["完整"] for entry in diagnostic["候选"])
    width, selected = DEFAULT_WIDTH, "完整往返"
    if complete:
        winner = min(diagnostic["候选"], key=lambda entry: (
            np.mean([value for values in entry["各模型各内折误差"].values() for value in values]),
            -entry["背景宽度_cm^-1"]))
        width = winner["背景宽度_cm^-1"]
        selected = min(MODELS, key=lambda model: (np.mean(winner["各模型各内折误差"][model]), MODELS.index(model)))
    diagnostic.update({"选型完整": complete, "所选背景宽度_cm^-1": width, "条件模型": selected,
        "截断约定": "任何候选未完成即保留预定650 cm^-1与完整往返，不挑选先完成候选",
        "同分规则": "背景宽度取大者，物理模型取两束"})
    return width, selected, diagnostic


def block_indices(block):
    return np.arange((block - 1) * 40, block * 40)


def score_rows(fit):
    case = fit["case"]
    if case["fold"] == 0:
        return [], []
    predictions = {model: fit["best"][model]["prediction"] for model in MODELS}
    predictions.update({"训练均值": case["mean"], "二次趋势": case["quadratic"],
                        "局部背景无条纹": case["background"]})
    rows, coverage = [], []
    calibration = np.concatenate([block_indices(block) for block in case["spec"]["校准块"]])
    for model, prediction in predictions.items():
        for angle_index, angle in enumerate(case["angles"]):
            residual_calibration = np.abs(prediction[angle_index, calibration] - case["values"][angle_index, calibration])
            rank = min(len(calibration), math.ceil((len(calibration) + 1) * 0.9))
            half_width = float(np.sort(residual_calibration)[rank - 1])
            for block in case["spec"]["测试块"]:
                indices = block_indices(block)
                difference = prediction[angle_index, indices] - case["values"][angle_index, indices]
                rows.append({"材料": case["material"], "折号": case["fold"], "附件": MATERIALS[case["material"]][angle_index],
                    "角度_度": angle, "测试块": block, "模型": model, "源行": case["rows"][indices],
                    "标准化均方根误差": float(np.sqrt(np.mean(difference ** 2)) / case["scale"][angle_index]),
                    "均方根误差_比例": float(np.sqrt(np.mean(difference ** 2))),
                    "训练残差四分位距_比例": case["scale"][angle_index]})
                distances = np.min(np.abs(case["sigma"][indices, None] - case["sigma"][case["train"]][None, :]), axis=1)
                coverage.append({"材料": case["material"], "折号": case["fold"], "附件": MATERIALS[case["material"]][angle_index],
                    "模型": model, "测试块": block, "名义覆盖率": 0.9, "经验覆盖率": float(np.mean(np.abs(difference) <= half_width)),
                    "命中点数": int(np.sum(np.abs(difference) <= half_width)), "测试点数": len(indices),
                    "区间全宽_比例": 2 * half_width, "距最近训练点平均波数_cm^-1": float(np.mean(distances)),
                    "构造方法": "相应角度80个校准绝对残差第ceil((80+1)*0.9)顺序统计量为半宽",
                    "保证限制": "连续谱相关且旧测试已经开发复用，只报经验覆盖；不宣称有限样本或厚度覆盖保证"})
    return rows, coverage


def parameter_record(candidate):
    return {**dict(zip(PARAMETERS, candidate["parameters"])), "训练标准化损失": candidate["loss"],
            "各角非负增益": candidate["gain"], "各角条纹特征平方和": candidate["feature_norm"],
            "厚度未识别": candidate["unidentified"],
            "搜索触边": [PARAMETERS[index] for index, value in enumerate(candidate["parameters"])
                        if min(abs(value - BOUNDS[index, 0]), abs(value - BOUNDS[index, 1])) < 1e-5]}


def case_record(fit):
    case = fit["case"]
    rows, coverage = score_rows(fit)
    errors = {model: float(np.mean([row["标准化均方根误差"] for row in rows if row["模型"] == model]))
              for model in (*MODELS, "训练均值", "二次趋势", "局部背景无条纹")} if rows else {}
    return {"材料": case["material"], "折号": case["fold"], "全量点数_每角度": len(case["sigma"]),
            "实际训练点数_每角度": len(case["train"]), "训练源行": case["rows"][case["train"]],
            "背景宽度_cm^-1": case["width"], "背景节点_cm^-1": case["operator"]["centers"],
            "背景自由度_每角每模型": len(case["operator"]["centers"]), "物理自由度_每模型": 5,
            "幅值自由度_每角每模型": 1, "条件模型": fit.get("selected", "完整往返"),
            **{model: parameter_record(fit["best"][model]) for model in MODELS},
            "标准化均方根误差": errors, "逐角逐块评分": rows, "预测经验覆盖": coverage,
            "搜索状态": fit["trace"], "训练内选择": fit.get("selection", {"状态": "预定特征的最小可用估计"})}


def validation_summary(records):
    result = {}
    for material in MATERIALS:
        members = [record for record in records if record["材料"] == material and record["折号"] > 0]
        fold_scores, block_deltas, attachments = [], [], []
        for member in members:
            errors = member["标准化均方根误差"]
            selected = member["条件模型"]
            fold_scores.append({"折号": member["折号"], "所选模型": selected,
                "完整往返相对两束下降_%": 100 * (1 - errors["完整往返"] / max(errors["两束"], 1e-15)),
                "所选模型相对基线下降_%": {label: 100 * (1 - errors[selected] / max(errors[label], 1e-15))
                                              for label in ("训练均值", "二次趋势", "局部背景无条纹")}})
            for row in member["逐角逐块评分"]:
                if row["模型"] != "两束":
                    continue
                paired = next(other for other in member["逐角逐块评分"] if other["模型"] == "完整往返"
                              and other["附件"] == row["附件"] and other["测试块"] == row["测试块"])
                block_deltas.append({"折号": member["折号"], "附件": row["附件"], "测试块": row["测试块"],
                    "两束减完整往返误差": row["标准化均方根误差"] - paired["标准化均方根误差"]})
        for number in MATERIALS[material]:
            attachment_deltas = [row["两束减完整往返误差"] for row in block_deltas if row["附件"] == number]
            attachments.append({"附件": number, "各折各块两束减完整误差": attachment_deltas,
                "跨折跨块全部正向": len(attachment_deltas) == 4 and all(value > 0 for value in attachment_deltas),
                "最差块改善": min(attachment_deltas) if attachment_deltas else None})
        result[material] = {"逐折基线对比": fold_scores, "逐角逐块改善": block_deltas,
            "逐附件跨折": attachments,
            "均值与二次基线各折均优": len(fold_scores) == 2 and all(
                row["所选模型相对基线下降_%"][label] > 0 for row in fold_scores for label in ("训练均值", "二次趋势")),
            "方向全部一致": len(block_deltas) == 8 and all(row["两束减完整往返误差"] > 0 for row in block_deltas)}
    return result


def checkpoint(data, fits, sensitivities, budget, state):
    records = [case_record(fits[key]) for key in sorted(fits, key=lambda key: (list(MATERIALS).index(key[0]), key[1]))]
    ranges = {}
    for material in MATERIALS:
        values = [record[model]["厚度_um"] for record in records if record["材料"] == material for model in MODELS]
        values.extend(row["扰动后厚度_um"] for row in sensitivities if row["材料"] == material and row["状态"] == "已重估")
        unidentified = any(record[model]["厚度未识别"] for record in records if record["材料"] == material for model in MODELS)
        if values:
            ranges[material] = [float(BOUNDS[0, 0]), float(BOUNDS[0, 1])] if unidentified else [min(values), max(values)]
    baseline = data["baseline"]
    core = {"碳化硅正式基准厚度_um": baseline["全量参数"]["厚度_um"],
            "碳化硅正式条件范围下限_um": baseline["条件范围_um"][0],
            "碳化硅正式条件范围上限_um": baseline["条件范围_um"][1]}
    for fold_number, label in ((0, "硅全量完整往返厚度_um"), (1, "硅折1完整往返条件厚度_um"), (2, "硅折2完整往返条件厚度_um")):
        if ("硅", fold_number) in fits:
            core[label] = float(fits[("硅", fold_number)]["best"]["完整往返"]["parameters"][0])
    if "硅" in ranges:
        core["硅已计算条件包络下限_um"], core["硅已计算条件包络上限_um"] = ranges["硅"]
    validation = validation_summary(records)
    payload = {"问题": 3, "变体": 2, "估计器版本": "局部训练背景同算子投影-往返场-v2",
        "状态": state, "核心指标": core, "案例": records, "厚度条件范围_um": ranges,
        "灵敏度_参数扰动": sensitivities, "分材料验证": validation,
        "硅全量条件厚度_um": core.get("硅全量完整往返厚度_um"),
        "条件范围构造": "两模型全量、各折及已算扰动的联合包络；幅值全零时扩展到整个厚度搜索盒",
        "厚度范围经验覆盖率": None, "厚度覆盖不可计算原因": "附件无真厚度，条件集合不是概率区间；独立合成只检验固定光学情景下的补充区间",
        "碳化硅正式基准": {"全量参数": baseline["全量参数"], "条件范围_um": baseline["条件范围_um"]},
        "输入哈希": data["public"]["原始附件SHA256"], "合成输入用于核心指标": False,
        "指标含义": {"厚度_um": "当前经验色散、界面与响应条件下的硅片外延厚度，非唯一真值",
            "标准化均方根误差": "先除相应训练二次趋势残差IQR，再按两角两测试块等权平均RMSE",
            "下降_%": "100*(1-模型误差/同折基线误差)，负值原样保留",
            "条件包络": "含已计算情景极端值；不把5392点拟合精度当厚度置信度"}}
    for filename in ("厚度结果.json", "公平对照.json"):
        budget.write(filename, payload)
    budget.write("结果声明_问题3.json", {"问题": 3, "核心指标": core,
        "口径说明": {"逐指标来源": {key: {"来源文件": "求解/问题3/升格2/结果/厚度结果.json", "键名": "核心指标." + key}
                                  for key in core}, "数据": "附件1至4真实光谱及冻结的正式问题2基准",
            "冻结复算规范": "数据/问题3_冻结合成输入/升格2/基准与协议.json",
            "核心指标同名口径": "与主线八项一致；变更的是训练背景与条纹特征，不回写主线",
            "验证属性": "既有开发折的探索性复查，不宣称新独立测试通过"},
        "自检指标": validation, "置信": {"等级": "条件性", "理由": "缺实测折射率和真厚度；保留全部负向结果及未完成检验"}})
    budget.write("区间覆盖.json", {"反射率预测": [row for record in records for row in record["预测经验覆盖"]],
        "厚度条件集合": {"经验覆盖率": None, "原因": "实测真值未知，非置信区间"},
        "补充条件厚度区间": "另见重采样与覆盖.json；不得冒充实测厚度覆盖"})
    budget.write("执行状态.json", {"状态": state, "完成案例数": len(fits), "计划案例数": 6,
        "完成灵敏度数": sum(row["状态"] == "已重估" for row in sensitivities),
        "软截止秒": SOFT_SECONDS, "硬预算秒": HARD_SECONDS, "核心指标数": len(core)})
    return payload


def physical_checks(fits, budget):
    rows = []
    for key, fit in fits.items():
        case = fit["case"]
        parameters = fit["best"]["完整往返"]["parameters"]
        optics = physical(parameters, case["sigma"][::max(1, len(case["sigma"]) // 40)], case["angles"], case["weight"])
        series_errors, tail_bounds, ratios, orders = [], [], [], []
        for reflection, first, ratio in optics["fields"]:
            closed = reflection + first / (1 - ratio)
            partial, term = reflection.astype(complex), first.copy()
            for order in range(1, 513):
                partial = partial + term
                term = term * ratio
                tail = np.abs(term) / (1 - np.abs(ratio))
                if np.max(tail) < 1e-13:
                    break
            series_errors.append(float(np.max(np.abs(partial - closed))))
            tail_bounds.append(float(np.max(tail)))
            ratios.append(float(np.max(np.abs(ratio))))
            orders.append(order)
        zero_parameters = parameters.copy()
        zero_parameters[3] = 0.0
        zero = physical(zero_parameters, case["sigma"], case["angles"], case["weight"])
        unit_phase_um = 4 * np.pi * parameters[0] * case["sigma"] / 10000.0
        unit_phase_cm = 4 * np.pi * (parameters[0] * 1e-4) * case["sigma"]
        local_constant_error = float(np.max(np.abs(project(np.ones((2, len(case["sigma"]))), case) - 1)))
        cross, leakage = [], {}
        if case["fold"] > 0:
            indices = np.concatenate([block_indices(block) for block in case["spec"]["测试块"]])
            full = physical(parameters, case["sigma"], case["angles"], case["weight"])["spectra"]["完整往返"]
            feature = full - project(full, case)
            for source, target in ((0, 1), (1, 0)):
                prediction = case["background"][source] + fit["best"]["完整往返"]["gain"][source] * feature[target]
                cross.append({"训练角度_度": case["angles"][source], "预测角度_度": case["angles"][target],
                    "标准化均方根误差": float(np.sqrt(np.mean((prediction[indices] - case["values"][target, indices]) ** 2))
                                                   / case["scale"][target]),
                    "限制": "共享光学参数由双角训练得到；只冻结角度响应作迁移诊断，不是完全留角独立训练"})
            altered_values = case["values"].copy()
            held = np.setdiff1d(np.arange(len(case["sigma"])), case["train"])
            altered_values[:, held] += 17.0
            altered = prepare({**case, "values": altered_values}, case["width"])
            original_candidate = evaluate(parameters, case, "完整往返")
            altered_candidate = evaluate(parameters, altered, "完整往返")
            leakage = {"非训练反射率统一增加_比例": 17.0,
                "训练目标变化": abs(original_candidate["loss"] - altered_candidate["loss"]),
                "预测最大变化_比例": float(np.max(np.abs(original_candidate["prediction"] - altered_candidate["prediction"]))),
                "训练与校准测试源行交集数": len(set(case["train"].tolist()) & set(held.tolist())),
                "含义": "改变所有非训练响应不改变固定候选的背景、增益与目标；外层校准仅用于区间半宽"}
        rows.append({"材料": key[0], "折号": key[1], "级数闭式最大误差": max(series_errors),
                     "尾项绝对上界": max(tail_bounds), "最大往返乘子模": max(ratios),
                     "有限级数次数": max(orders), "最大传播因子模": float(np.max(optics["attenuation"])),
                     "厘米微米换算最大相位差_rad": float(np.max(np.abs(unit_phase_um - unit_phase_cm))),
                     "背景算子常数保持最大误差": local_constant_error,
                     "零返回两模型最大反射率差": float(np.max(np.abs(zero["spectra"]["两束"] - zero["spectra"]["完整往返"]))),
                     "双向冻结响应诊断": cross, "防泄漏核验": leakage})
    budget.write("物理与角度核验.json", {"核验": rows, "指标含义": "复场绝对误差及无量纲传播幅值；不是厚度真值误差"})
    return rows


def harmonic_diagnostic(fits, budget):
    rows = []
    for key, fit in fits.items():
        case = fit["case"]
        parameters = fit["best"]["完整往返"]["parameters"]
        film = parameters[1] + parameters[2] * ((2000 / case["sigma"]) ** 2 - 1)
        phase = 4 * np.pi * parameters[0] * case["sigma"] * np.sqrt(
            film ** 2 - np.sin(np.deg2rad(case["angles"]))[:, None] ** 2) / 10000
        test = np.concatenate([block_indices(block) for block in case["spec"]["测试块"]]) if case["fold"] else case["train"]
        for angle_index, angle in enumerate(case["angles"]):
            scores, amplitudes = {}, {}
            for order in (1, 3):
                basis = np.asarray([function(harmonic * phase[angle_index]) for harmonic in range(1, order + 1)
                                    for function in (np.cos, np.sin)])
                basis = (basis - project(basis, case)).T
                target = case["values"][angle_index] - case["background"][angle_index]
                train = case["train"]
                penalty = case["ridge"] * len(train) * case["scale"][angle_index] ** 2
                coefficient = np.linalg.solve(basis[train].T @ basis[train] + np.eye(2 * order) * max(penalty, 1e-12),
                                               basis[train].T @ target[train])
                prediction = case["background"][angle_index] + basis @ coefficient
                scores[str(order)] = float(np.sqrt(np.mean((prediction[test] - case["values"][angle_index, test]) ** 2))
                                           / case["scale"][angle_index])
                amplitudes[str(order)] = [float(np.linalg.norm(coefficient[2 * harmonic:2 * harmonic + 2]))
                                          for harmonic in range(order)]
            rows.append({"材料": key[0], "折号": key[1], "附件": MATERIALS[key[0]][angle_index], "角度_度": angle,
                "一阶误差": scores["1"], "三阶误差": scores["3"],
                "三阶相对一阶下降_%": 100 * (1 - scores["3"] / max(scores["1"], 1e-15)),
                "三阶各谐波幅值": amplitudes["3"],
                "用途限制": "相位来自训练完整场；弱界面假设交叉诊断，不参与选型，谐波可来自趋势和色散而非多光束"})
    budget.write("谐波贡献诊断.json", {"逐附件逐折": rows, "指标含义": "同一训练背景后的1阶与3阶相位谐波对照"})


def sensitivity_specs():
    return [("衬底对比减20%", 3, "乘", 0.8), ("衬底对比加20%", 3, "乘", 1.2),
        ("有效损耗未扰动控制", 4, "乘", 1.0), ("有效损耗减20%", 4, "乘", 0.8),
        ("有效损耗加20%", 4, "乘", 1.2), ("有效损耗绝对增加0→0.1", 4, "零绝对", 0.1),
        ("有效损耗绝对增加0→0.2", 4, "零绝对", 0.2),
        ("偏振权重全s", None, "偏振", 1.0), ("偏振权重全p", None, "偏振", 0.0),
        ("入射角同时减0.5度", None, "角度", -0.5), ("入射角同时加0.5度", None, "角度", 0.5),
        ("参考折射率减3%", 1, "乘", 0.97), ("参考折射率加3%", 1, "乘", 1.03),
        ("色散减20%", 2, "乘", 0.8), ("色散加20%", 2, "乘", 1.2),
        ("窗口下界加100", None, "下界", 1300.0), ("窗口上界减100", None, "上界", 3700.0),
        ("背景宽度减20%", None, "宽度", 0.8), ("背景宽度加20%", None, "宽度", 1.2),
        ("幅值岭系数置零", None, "岭", 0.0), ("幅值岭系数加倍", None, "岭", 0.2),
        ("只用10度重估", None, "单角", 0), ("只用15度重估", None, "单角", 1)]


def run_sensitivities(data, fits, budget, rows):
    for label, index, operation, value in sensitivity_specs():
        for key, fit in fits.items():
            case = fit["case"]
            base = fit["best"]["完整往返"]["parameters"]
            row = {"材料": key[0], "折号": key[1], "情景": label, "基准厚度_um": base[0],
                   "基准参数": dict(zip(PARAMETERS, base))}
            if operation == "零绝对" and base[index] != 0.0:
                rows.append({**row, "状态": "不适用", "原因": "仅适用于有效损耗基准精确等于0，当前基准非零"})
                continue
            if not budget.available(140.0):
                rows.append({**row, "状态": "时间截断未评估", "原因": "为冻结输入、覆盖诊断及最终数值写出保留时间"})
                continue
            trial = {**case}
            fixed = {}
            width = case["width"]
            if index is not None:
                fixed[index] = base[index] * value if operation == "乘" else value
                row["请求固定参数值"] = fixed[index]
                if not BOUNDS[index, 0] <= fixed[index] <= BOUNDS[index, 1]:
                    rows.append({**row, "状态": "物理搜索域外", "原因": "不把请求扰动静默裁剪成不同幅度"})
                    continue
            elif operation == "偏振":
                trial["weight"] = value
            elif operation == "角度":
                trial["angles"] = ANGLES + value
            elif operation == "宽度":
                width *= value
            elif operation == "岭":
                trial["ridge"] = value
            elif operation == "单角":
                trial["angles"] = ANGLES[[value]]
                trial["values"] = case["values"][[value]]
            elif operation in ("下界", "上界"):
                condition = case["sigma"] >= value if operation == "下界" else case["sigma"] <= value
                trial["train"] = case["train"][condition[case["train"]]]
            trial = prepare(trial, width)
            try:
                candidate = search_pair(trial, budget, budget.deadline(4.0), grid_count=0, rounds=5,
                                        seeds=[fit["best"][model]["parameters"] for model in MODELS], fixed=fixed)
            except NoFeasibleCandidate as error:
                rows.append({**row, "状态": "请求条件不可行", "原因": str(error)})
                continue
            best = candidate["best"]["完整往返"]
            row.update({"状态": "已重估", "扰动后厚度_um": best["parameters"][0],
                "厚度变化_um": best["parameters"][0] - base[0], "相对变化_%": 100 * (best["parameters"][0] / base[0] - 1),
                "实际固定参数值": {PARAMETERS[fixed_index]: best["parameters"][fixed_index] for fixed_index in fixed},
                "固定值实际变化": {PARAMETERS[fixed_index]: best["parameters"][fixed_index] - base[fixed_index] for fixed_index in fixed},
                "重估参数": parameter_record(best), "两模型厚度差_um": best["parameters"][0] - candidate["best"]["两束"]["parameters"][0],
                "样本点数": len(trial["train"]), "训练波数范围_cm^-1": trial["sigma"][trial["train"][[0, -1]]],
                "偏振s权重": trial["weight"], "入射角_度": trial["angles"], "背景宽度_cm^-1": width,
                "搜索状态": candidate["trace"], "实际用时秒": budget.elapsed()})
            rows.append(row)
            budget.write("灵敏度.json", {"灵敏度_参数扰动": rows, "计划项数": len(sensitivity_specs()) * len(fits),
                "历史口径关系": "保留原15类损耗/衬底/角度/偏振/窗口口径并增加色散与背景；适用数随新估计器真实基准变化，不照抄54与2"})
        checkpoint(data, fits, rows, budget, "灵敏度分步完成")


def profile_library(case, parameters):
    thicknesses = np.linspace(0.5, 40.0, 241)
    spectra = {model: [] for model in MODELS}
    for thickness in thicknesses:
        trial = parameters.copy()
        trial[0] = thickness
        optics = physical(trial, case["sigma"], case["angles"], case["weight"])
        for model in MODELS:
            spectra[model].append(optics["spectra"][model])
    return thicknesses, {model: np.asarray(spectra[model]) - project(np.asarray(spectra[model]), case)
                         for model in MODELS}


def profile_estimate(observed, case, thicknesses, library):
    background = project(observed, case)
    target = observed - background
    train = case["train"]
    scale_squared = case["scale"] ** 2
    result = {}
    for model in MODELS:
        feature = library[model][:, :, train]
        norm = np.sum(feature ** 2, axis=2)
        gains = np.clip(np.sum(feature * target[None, :, train], axis=2)
                        / np.maximum(norm + case["ridge"] * len(train) * scale_squared[None, :], 1e-24), 0, GAIN_LIMIT)
        difference = (gains[:, :, None] * feature - target[None, :, train]) / case["scale"][None, :, None]
        losses = np.mean(difference ** 2, axis=(1, 2)) + case["ridge"] * np.mean(gains ** 2, axis=1)
        winner = int(np.argmin(losses))
        result[model] = {"厚度_um": float(thicknesses[winner]), "损失": float(losses[winner]),
                         "prediction": background + gains[winner, :, None] * library[model][winner]}
    return result


def resample_indices(size, block_length, generator):
    starts = generator.integers(0, size - block_length + 1, size=math.ceil(size / block_length))
    return np.concatenate([np.arange(start, start + block_length) for start in starts])[:size]


def supplemental_uncertainty(data, fits, budget, manifest):
    generator = np.random.default_rng(SEED)
    summaries = []
    for material in MATERIALS:
        if not budget.available(60.0):
            break
        full = fits[(material, 0)]
        case = prepare({**make_case(data, material, 1), "fold": 0, "train": np.arange(480)}, full["case"]["width"])
        parameters = full["best"]["完整往返"]["parameters"].copy()
        thicknesses, library = profile_library(case, parameters)
        protocol_path = FROZEN / f"{material}_重采样协议.json"
        atomic_json(protocol_path, {"材料": material, "背景宽度_cm^-1": case["width"],
            "固定光学参数": dict(zip(PARAMETERS, parameters)), "厚度搜索网格_um": thicknesses,
            "训练尺度_反射率比例": case["scale"], "幅值岭系数": case["ridge"], "增益上限": GAIN_LIMIT,
            "入射角_度": case["angles"], "s偏振权重": case["weight"],
            "训练范围": "480个抽样点全部使用，仅用于辅助区间，不作外层性能检验",
            "背景算法": "按floor((波数-1200)/宽度)分箱，取训练波数和反射率均值，箱间线性插值，两端保持节点值",
            "响应算法": "预测=背景(观测)+增益*(物理反射率-背景(物理反射率))",
            "增益算法": "clip(条纹特征与去背景观测内积/(特征平方和+岭系数*点数*尺度平方),0,20)",
            "剖面目标": "逐角标准化残差平方均值加岭系数乘逐角增益平方均值",
            "厚度真值是否传入搜索": False, "双角噪声": "采用相同块索引；块内保持原顺序，不循环绕回",
            "经验区间分位": [0.025, 0.975], "厚度区间离散扩宽_um": float((thicknesses[1] - thicknesses[0]) / 2),
            "生成侧最小增益": 0.1, "非匹配情景规则": "只在生成端修改损耗或色散，估计端仍固定此处光学参数"})
        manifest[protocol_path.name] = {"用途": "仅凭数据目录重构条件重采样估计器，不依赖结果目录参数",
            "生成方式": "复制当前训练估计器的必要参数及确定性算法定义", "参数": {"材料": material},
            "SHA256": hashlib.sha256(protocol_path.read_bytes()).hexdigest()}
        atomic_json(FROZEN / "输入清单.json", {"问题": 3, "变体": 2, "文件": manifest})
        original = profile_estimate(case["values"], case, thicknesses, library)
        center = original["完整往返"]["prediction"]
        residual = case["values"] - center
        residual -= np.mean(residual, axis=1)[:, None]
        draws, observations, indices_all = [], [], []
        for repetition in range(40):
            if not budget.available(55.0):
                break
            indices = resample_indices(480, 40, generator)
            observation = center + residual[:, indices]
            estimates = profile_estimate(observation, case, thicknesses, library)
            observations.append(observation)
            indices_all.append(indices)
            draws.append({model: estimates[model]["厚度_um"] for model in MODELS})
        if not draws:
            continue
        freeze_array(f"{material}_实测条件重采样.npz", {"波数_cm^-1": case["sigma"], "重采样反射率_比例": np.asarray(observations),
            "重采样源序号": np.asarray(indices_all), "原始反射率_比例": case["values"], "固定光学参数": parameters},
            "辅助条件厚度与配对差值的重采样，非八项实测核心输入", "双角同步移动块残差重采样",
            {"种子": SEED, "块长": 40, "计划重复数": 40, "完成重复数": len(draws)}, manifest)
        padding = float((thicknesses[1] - thicknesses[0]) / 2)
        pair_differences = [draw["完整往返"] - draw["两束"] for draw in draws]
        intervals = {model: [max(0.5, float(np.quantile([draw[model] for draw in draws], 0.025)) - padding),
                              min(40.0, float(np.quantile([draw[model] for draw in draws], 0.975)) + padding)] for model in MODELS}
        member = {"材料": material, "块长_抽样点": 40, "计划重复数": 40, "完成重复数": len(draws),
            "固定光学参数": dict(zip(PARAMETERS[1:], parameters[1:])), "条件厚度区间_um": intervals,
            "配对厚度差经验分位范围_um": [float(np.quantile(pair_differences, 0.025)) - 2 * padding,
                                             float(np.quantile(pair_differences, 0.975)) + 2 * padding],
            "条件重采样逐例": draws, "构造方法": "固定非厚度光学量；双角同步40点残差块，重估厚度与增益；2.5%至97.5%经验分位加半格离散误差",
            "光学参数不确定性处理": "仅由实测主结果的联合情景包络承担，不混同本区间",
            "名义覆盖率": 0.95, "实测厚度经验覆盖率": None, "实测无真值原因": "只能由下列独立合成检验条件构造"}
        block_sensitivity = []
        for block_length in (20, 60):
            extra_observations, extra_estimates = [], []
            for repetition in range(12):
                if not budget.available(45.0):
                    break
                indices = resample_indices(480, block_length, generator)
                observation = center + residual[:, indices]
                estimates = profile_estimate(observation, case, thicknesses, library)
                extra_observations.append(observation)
                extra_estimates.append(estimates["完整往返"]["厚度_um"])
            if extra_estimates:
                freeze_array(f"{material}_块长{block_length}灵敏度.npz", {"波数_cm^-1": case["sigma"],
                    "反射率_比例": np.asarray(extra_observations), "固定光学参数": parameters},
                    "条件厚度区间对残差相关块长的敏感性", "双角同步移动块残差重采样",
                    {"种子": SEED, "块长": block_length, "重复数": len(extra_estimates)}, manifest)
                block_sensitivity.append({"块长_点": block_length, "重复数": len(extra_estimates),
                    "条件厚度区间_um": [max(0.5, float(np.quantile(extra_estimates, 0.025)) - padding),
                                             min(40.0, float(np.quantile(extra_estimates, 0.975)) + padding)]})
        member["重采样块长灵敏度"] = block_sensitivity
        coverage_cases = []
        for generation_model in MODELS:
            for repetition in range(6):
                if not budget.available(35.0):
                    break
                truth = (3.3, 8.7, 14.1)[repetition % 3]
                truth_parameters = parameters.copy()
                truth_parameters[0] = truth
                scenario = "匹配光学情景"
                if repetition == 3:
                    truth_parameters[4] = 0.0
                    scenario = "低损耗失配" if parameters[4] != 0.0 else "零损耗未变控制"
                elif repetition == 4:
                    truth_parameters[4] = 2.0
                    scenario = "强损耗失配" if parameters[4] != 2.0 else "强损耗未变控制"
                elif repetition == 5:
                    truth_parameters[2] *= 0.8
                    scenario = "色散减20%失配" if parameters[2] != 0.0 else "零色散未变控制"
                truth_optics = physical(truth_parameters, case["sigma"], case["angles"], case["weight"])
                truth_spectrum = truth_optics["spectra"][generation_model]
                truth_feature = truth_spectrum - project(truth_spectrum, case)
                truth_mean = case["background"] + np.maximum(full["best"][generation_model]["gain"], 0.1)[:, None] * truth_feature
                noise_indices = resample_indices(480, 40, generator)
                observation = truth_mean + residual[:, noise_indices]
                estimates = profile_estimate(observation, case, thicknesses, library)
                simulation_center = estimates[generation_model]["prediction"]
                simulation_residual = observation - simulation_center
                simulation_residual -= np.mean(simulation_residual, axis=1)[:, None]
                repetitions, simulation_draws = [], []
                for inner_number in range(24):
                    if not budget.available(30.0):
                        break
                    inner_indices = resample_indices(480, 40, generator)
                    inner_observation = simulation_center + simulation_residual[:, inner_indices]
                    inner_estimates = profile_estimate(inner_observation, case, thicknesses, library)
                    repetitions.append(inner_estimates[generation_model]["厚度_um"])
                    simulation_draws.append(inner_observation)
                if not repetitions:
                    break
                interval = [max(0.5, float(np.quantile(repetitions, 0.025)) - padding),
                            min(40.0, float(np.quantile(repetitions, 0.975)) + padding)]
                case_number = len(coverage_cases)
                frozen_name = f"{material}_独立条件覆盖_{case_number}.npz"
                freeze_array(frozen_name, {"波数_cm^-1": case["sigma"], "反射率_比例": observation,
                    "重采样反射率_比例": np.asarray(simulation_draws), "真厚度_um": np.array(truth),
                    "真光学参数": truth_parameters, "生成均值_比例": truth_mean},
                    "已知厚度的条件覆盖与无高阶误判诊断，不用于实测核心指标", "固定真厚度、指定场模型及独立随机块噪声",
                    {"种子": SEED, "生成模型": generation_model, "真厚度_um": truth, "情景": scenario,
                     "实际重复数": len(repetitions)}, manifest)
                difference = estimates["完整往返"]["厚度_um"] - estimates["两束"]["厚度_um"]
                coverage_cases.append({"生成模型": generation_model, "情景": scenario, "真厚度_um": truth,
                    "估计厚度_um": estimates[generation_model]["厚度_um"], "条件区间_um": interval,
                    "覆盖真值": interval[0] <= truth <= interval[1], "重复数": len(repetitions),
                    "两模型厚度差_um": difference, "输入文件": f"数据/问题3_冻结合成输入/升格2/{frozen_name}"})
        member["独立合成逐例"] = coverage_cases
        member["独立合成完成数"] = len(coverage_cases)
        member["独立合成计划数"] = 12
        member["独立合成经验覆盖率"] = float(np.mean([entry["覆盖真值"] for entry in coverage_cases])) if coverage_cases else None
        member["分情景经验覆盖率"] = {scenario: float(np.mean([entry["覆盖真值"] for entry in coverage_cases if entry["情景"] == scenario]))
                                     for scenario in sorted({entry["情景"] for entry in coverage_cases})}
        member["分生成模型经验覆盖率"] = {model: float(np.mean([entry["覆盖真值"] for entry in coverage_cases if entry["生成模型"] == model]))
                                        for model in MODELS if any(entry["生成模型"] == model for entry in coverage_cases)}
        null_differences = [abs(entry["两模型厚度差_um"]) for entry in coverage_cases if entry["生成模型"] == "两束"]
        member["无高阶厚度差95%经验阈值_um"] = float(np.quantile(null_differences, 0.95)) if null_differences else None
        member["无高阶误判率_预设相对差阈值"] = {f"{threshold:g}%": float(np.mean([
            abs(entry["两模型厚度差_um"]) / entry["真厚度_um"] > threshold / 100
            for entry in coverage_cases if entry["生成模型"] == "两束"])) for threshold in (0.5, 1.0, 2.0)} if null_differences else {}
        member["局限"] = "仅12个预定情景及小样本重采样；阈值是诊断，不是仪器精度；不把匹配模型覆盖移植到实测真厚度"
        summaries.append(member)
        budget.write("重采样与覆盖.json", {"材料": summaries, "合成输入清单": "数据/问题3_冻结合成输入/升格2/输入清单.json"})
    budget.write("重采样与覆盖.json", {"材料": summaries, "完成材料数": len(summaries),
        "合成输入清单": "数据/问题3_冻结合成输入/升格2/输入清单.json",
        "未完成说明": "若少于2材料或每材料12独立情景，为时间截断；不填造覆盖率"})
    return summaries


def final_judgement(payload, uncertainty, fields, budget):
    judgments = []
    for material in MATERIALS:
        validation = payload["分材料验证"][material]
        full = next(record for record in payload["案例"] if record["材料"] == material and record["折号"] == 0)
        difference = full["完整往返"]["厚度_um"] - full["两束"]["厚度_um"]
        diagnostic = next((row for row in uncertainty if row["材料"] == material), {})
        threshold = diagnostic.get("无高阶厚度差95%经验阈值_um")
        field = next(row for row in fields if row["材料"] == material and row["折号"] == 0)
        for attachment in validation["逐附件跨折"]:
            support = attachment["跨折跨块全部正向"] and validation["均值与二次基线各折均优"]
            judgments.append({"材料": material, "附件": attachment["附件"],
                "多光束判断": "与可观测高阶贡献相容，非物理确证" if support else "当前数据不足以稳定支持高阶贡献，不等于不存在",
                "谱形条件支持": support, "同口径完整减两束厚度差_um": difference,
                "厚度影响诊断阈值_um": threshold,
                "厚度影响超经验阈值": abs(difference) > threshold if threshold is not None else None,
                "相对厚度差阈值情景": {f"{value:g}%": abs(difference) / max(full["两束"]["厚度_um"], 1e-12) > value / 100
                                           for value in (0.5, 1.0, 2.0)},
                "最大往返乘子模": field["最大往返乘子模"],
                "替代解释": "趋势/增益、经验色散与厚度分支均可影响残差；小角度双谱高相关不是高阶证明",
                "缺失物理证据": "仪器分辨率、时间相干、偏振及独立折射率均未提供"})
    silicon_carbide = next(record for record in payload["案例"] if record["材料"] == "碳化硅" and record["折号"] == 0)
    carbide_rows = [row for row in judgments if row["材料"] == "碳化硅"]
    correction = all(row["谱形条件支持"] and row["厚度影响超经验阈值"] is True for row in carbide_rows)
    budget.write("条件判定.json", {"逐附件判断": judgments, "碳化硅触发条件性修正": correction,
        "碳化硅正式保留厚度_um": payload["核心指标"]["碳化硅正式基准厚度_um"],
        "碳化硅同窗口新两束条件厚度_um": silicon_carbide["两束"]["厚度_um"],
        "碳化硅完整往返候选修正厚度_um": silicon_carbide["完整往返"]["厚度_um"],
        "处理": "仅在两个条件均成立时将完整场候选作为条件性修正；原正式基准始终单列且不改写",
        "验证结论限制": "只交付可计算诊断与数值，不代替独立解读或宣称G2通过"})


def assumption_audit(sensitivities, fields, uncertainty, budget):
    specifications = [
        ("A1V2", "同片10°和15°共享厚度，简化为每材料一个物理参数向量；两角背景和增益分别估计。",
         "数据档案:总览.独立晶圆数；数据档案:跨文件关系", ("只用10度重估", "只用15度重估"),
         "分别只用一个角度重估全部可变物理参数，报告相对共享厚度的变化。"),
        ("A2V2", "在1200至3800 cm^-1内，用参考折射率加倒波数平方项及常数衬底对比替代未知色散函数。",
         "数据档案:未随附件提供的建模输入；题面契约:3-5", ("参考折射率减3%", "参考折射率加3%", "色散减20%", "色散加20%", "衬底对比减20%", "衬底对比加20%"),
         "逐项固定请求值后重估厚度和其他自由参数；不裁剪越界扰动。"),
        ("A3V2", "将未知仪器连续背景近似为训练分箱均值的线性插值，训练支撑外冻结端点；同算子作用于物理谱。",
         "数据档案:质量问题与决策；数据档案:未随附件提供的建模输入", ("背景宽度减20%", "背景宽度加20%", "窗口下界加100", "窗口上界减100", "幅值岭系数置零", "幅值岭系数加倍"),
         "比较背景宽度±20%、窗口单端缩小100 cm^-1及增益正则强度改变后的重新估计。"),
        ("A4V2", "未测偏振与往返吸收时，基准使用s/p各半及非负常数有效损耗；不把有效损耗当实测消光系数。",
         "数据档案:未随附件提供的建模输入", ("偏振权重全s", "偏振权重全p", "有效损耗未扰动控制", "有效损耗减20%", "有效损耗加20%", "有效损耗绝对增加0→0.1", "有效损耗绝对增加0→0.2", "入射角同时减0.5度", "入射角同时加0.5度"),
         "全s/全p、角度±0.5°及损耗乘法/零点绝对扰动后重估厚度；非零基准明确标注零点情景不适用。"),
    ]
    results, ledger = {}, []
    for number, statement, source, labels, obligation in specifications:
        relevant = [row for row in sensitivities if row["情景"] in labels]
        completed = [row for row in relevant if row["状态"] == "已重估"]
        maximum = max((abs(row["相对变化_%"]) for row in completed), default=None)
        results[number] = {"对应情景": relevant, "完成重估数": len(completed), "计划状态数": len(relevant),
                           "最大绝对厚度相对变化_%": maximum}
        ledger.append({"假设号": number, "假设": statement, "依据": source, "灵敏度义务": obligation,
            "检验结果": f"完成{len(completed)}项重估，共记录{len(relevant)}项适用性状态；最大绝对相对变化为{maximum}%；未完成项不作通过判断（求解结果:假设检验.{number}）"})
    results["A5V2"] = {"分材料条件覆盖": uncertainty, "完成材料数": len(uncertainty)}
    ledger.append({"假设号": "A5V2", "假设": "仅在辅助条件区间中，固定非厚度光学量，将双角同步连续残差块视为可重复噪声；不用于实测核心点估计。",
        "依据": "数据档案:未随附件提供的建模输入.同条件重复测量", "灵敏度义务": "比较20/40/60点块长，并在匹配、低/强损耗、色散失配的已知厚度情景上报告经验覆盖。",
        "检验结果": f"输出{len(uncertainty)}个材料的已完成重采样和逐例覆盖；样本数与时间截断状态同时披露，不代表实测真厚度覆盖（求解结果:假设检验.A5V2）"})
    results["场与防泄漏"] = fields
    budget.write("假设检验.json", {"假设检验": results, "指标含义": "所有条目指向本次实际检验值或明确的未完成状态，未预置通过"})
    atomic_json(RESULT / "假设台账_问题3_升格2.json", ledger)


def main():
    budget = Budget()
    manifest, fits, sensitivities = {}, {}, []
    data = load_inputs(budget, manifest)
    for material in MATERIALS:
        for fold_number in (0, 1, 2):
            case = prepare(make_case(data, material, fold_number))
            fit = search_pair(case, budget, budget.deadline(16.0), grid_count=32, rounds=3)
            fits[(material, fold_number)] = fit
            checkpoint(data, fits, sensitivities, budget, "已形成真实数据最小可用条件估计")
    for key in list(fits):
        if not budget.available(450.0):
            break
        base_case = fits[key]["case"]
        width, selected, diagnostic = inner_selection(base_case, budget, budget.deadline(32.0))
        case = prepare(base_case, width)
        refined = search_pair(case, budget, budget.deadline(36.0), grid_count=96, rounds=8,
                              seeds=[fits[key]["best"][model]["parameters"] for model in MODELS])
        refined["selected"], refined["selection"] = selected, diagnostic
        fits[key] = refined
        budget.write(f"训练内选择_{key[0]}_折{key[1]}.json", diagnostic)
        checkpoint(data, fits, sensitivities, budget, "已完成当前案例的训练内特征选择")
    fields = physical_checks(fits, budget)
    harmonic_diagnostic(fits, budget)
    run_sensitivities(data, fits, budget, sensitivities)
    uncertainty = supplemental_uncertainty(data, fits, budget, manifest)
    payload = checkpoint(data, fits, sensitivities, budget, "完成；数值与验证缺口同时交付")
    final_judgement(payload, uncertainty, fields, budget)
    assumption_audit(sensitivities, fields, uncertainty, budget)
    prediction_rows = []
    for fit in fits.values():
        case = fit["case"]
        if case["fold"] == 0:
            continue
        for model in MODELS:
            prediction_rows.append({"材料": case["material"], "折号": case["fold"], "模型": model,
                "源行": case["rows"], "波数_cm^-1": case["sigma"], "观测反射率_比例": case["values"],
                "预测反射率_比例": fit["best"][model]["prediction"], "折分": case["spec"]})
    budget.write("留段预测.json", {"预测": prediction_rows, "指标含义": "480点含训练、校准、测试，外层评分只读测试块"})
    budget.write("同口径对照.json", {"案例": payload["案例"], "分材料验证": payload["分材料验证"],
        "核心指标": payload["核心指标"], "高阶比较": "两模型共享全部预处理与优化资源，只改变首返回场之后的递推分母"})
    budget.write("执行状态.json", {"状态": "正常结束", "软预算触发": budget.elapsed() >= SOFT_SECONDS,
        "硬预算内": budget.elapsed() < HARD_SECONDS, "完成案例数": len(fits), "核心指标数": len(payload["核心指标"]),
        "检验缺口数": sum(row["状态"] not in ("已重估", "不适用") for row in sensitivities),
        "未自判通过": True})
    print(json.dumps(plain({"问题": 3, "变体": 2, "核心指标": payload["核心指标"],
        "分材料验证": payload["分材料验证"], "实际用时秒": budget.elapsed()}), ensure_ascii=False, indent=2), flush=True)


if __name__ == "__main__":
    execution_start = time.monotonic()
    try:
        main()
    except Exception as error:
        atomic_json(RESULT / "执行异常.json", {"状态": "异常终止，保留已完成的分步结果", "异常类型": type(error).__name__,
            "原因": str(error), "实际用时秒": time.monotonic() - execution_start})
        raise
