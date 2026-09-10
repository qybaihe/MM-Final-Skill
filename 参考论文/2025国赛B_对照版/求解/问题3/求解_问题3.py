"""问题3：正式升格3入口、现存证据导表及有限成对数值补查。

输入只使用四个原始XLSX和问题2正式基准；不生成合成观测。波数用cm^-1，
厚度用微米，反射率由百分数除以100。脚本采用标准库读取XLSX，避免隐式插值。
主模型是 E_inf=r01+B/(1-u)，两束对照为 E_2=r01+B；两模型共享同一
厚度、经验色散和双角物理参数。正式响应为训练内选择的正交收缩响应；
下列旧函数中的二次基线、一次包络及18轮坐标搜索仅属历史实现。

运行预算：20分钟总上限、18分钟软截止。每个阶段原子写出结果；即使搜索被软截止，
也保留已经完成的折、预测、场核验和灵敏度。运行结束打印摘要。
默认执行版本化仲裁闭环；--paired-budget-review保留旧成对补查入口。
以下旧函数仅供显式历史复算，不生成当前正式声明。
--export-current-evidence 只读当前灵敏度和逐块评分，不重新拟合。
--paired-budget-review 冻结当前源行、响应、目标与物理边界，补查硅次折及碳化硅18个起点模型组合。
所有新运行写入隔离候选目录；未经独立复核不得覆盖当前正式结果。
"""

import argparse
import bisect
import cmath
import csv
import fcntl
import hashlib
import importlib.util
import json
import math
import os
import posixpath
import signal
import time
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

import numpy as np


ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "数据"
RESULT = ROOT / "求解/问题3/结果"
BASELINE = ROOT / "求解/问题2/结果/基准交接.json"
ROUTE = ROOT / "交接/路线侦察.json"
DECLARATION = ROOT / "交接/结果声明_问题3.json"
SENSITIVITY_RESULT_KEY = "灵敏度_参数扰动"
LOSS_LABELS = ("有效损耗未扰动控制", "有效损耗减20%", "有效损耗加20%",
               "有效损耗绝对增加0→0.1", "有效损耗绝对增加0→0.2")
PARAMETER_NAMES = ("厚度_um", "参考折射率", "色散系数", "衬底折射率对比", "有效往返损耗")
ANGLES = (10.0, 15.0)
MATERIAL_FILES = {"硅": (3, 4), "碳化硅": (1, 2)}
MATERIAL_ORDER = ("硅", "碳化硅")
BOUNDS = ((0.5, 40.0), (1.2, 6.0), (-1.0, 1.0), (-1.5, 2.5), (0.0, 3.0))
GAIN_LIMIT = 20.0
SOFT_SECONDS = 1080.0
HARD_SECONDS = 1180.0
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def atomic_json(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(path.name + f".{os.getpid()}.tmp")
    with temp.open("w", encoding="utf-8") as stream:
        json.dump(payload, stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write("\n")
        stream.flush()
        os.fsync(stream.fileno())
    os.replace(temp, path)


def write_csv(path, rows, fields):
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(path.name + f".{os.getpid()}.tmp")
    with temp.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
        stream.flush()
        os.fsync(stream.fileno())
    os.replace(temp, path)


def elapsed(t0):
    return time.monotonic() - t0


def read_sheet(path):
    """按原始XML顺序读Sheet1的两列数值，保留原始行号。"""
    with ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {x.get("Id"): x.get("Target") for x in rels}
        sheet = next(x for x in workbook.findall("m:sheets/m:sheet", NS) if x.get("name") == "Sheet1")
        rid = sheet.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        target = targets[rid]
        member = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
        root = ET.fromstring(archive.read(member))
        records = []
        for row in root.findall("m:sheetData/m:row", NS):
            number = int(row.get("r"))
            if number < 2:
                continue
            cells = {cell.get("r"): cell for cell in row.findall("m:c", NS)}
            values = []
            for column in ("A", "B"):
                cell = cells.get(f"{column}{number}")
                if cell is None or cell.get("t", "n") != "n" or cell.find("m:f", NS) is not None:
                    raise ValueError(f"{path.name}第{number}行不是原始数值")
                values.append(float(cell.findtext("m:v", namespaces=NS)))
            if not all(math.isfinite(x) for x in values) or values[0] <= 0:
                raise ValueError(f"{path.name}第{number}行存在非法波数或反射率")
            records.append((number, values[0], values[1] / 100.0))
    records.sort(key=lambda x: x[1])
    if len(records) != 7469 or len({x[1] for x in records}) != 7469:
        raise ValueError(f"{path.name}行数或波数唯一性不符合数据档案")
    return records


def quantile(values, probability):
    ordered = sorted(values)
    if not ordered:
        return 0.0
    position = (len(ordered) - 1) * probability
    low, high = math.floor(position), math.ceil(position)
    return ordered[low] + (ordered[high] - ordered[low]) * (position - low)


def dot(a, b):
    return sum(x * y for x, y in zip(a, b))


def solve_linear(matrix, target):
    """带部分选主元的高斯消元，矩阵维数不超过5。"""
    n = len(target)
    a = [list(row) + [target[i]] for i, row in enumerate(matrix)]
    for col in range(n):
        pivot = max(range(col, n), key=lambda row: abs(a[row][col]))
        if abs(a[pivot][col]) < 1e-12:
            raise ValueError("线性校准矩阵秩不足")
        a[col], a[pivot] = a[pivot], a[col]
        scale = a[col][col]
        a[col] = [value / scale for value in a[col]]
        for row in range(n):
            if row == col:
                continue
            factor = a[row][col]
            if factor:
                a[row] = [left - factor * right for left, right in zip(a[row], a[col])]
    return [a[row][-1] for row in range(n)]


def least_squares(columns, target, ridge=1e-12):
    p = len(columns)
    gram = [[0.0] * p for _ in range(p)]
    rhs = [0.0] * p
    for row in range(len(target)):
        values = [column[row] for column in columns]
        for i in range(p):
            rhs[i] += values[i] * target[row]
            for j in range(p):
                gram[i][j] += values[i] * values[j]
    for i in range(p):
        gram[i][i] += ridge
    return solve_linear(gram, rhs)


def poly_fit(x, y, degree):
    columns = [[value ** power for value in x] for power in range(degree + 1)]
    coefficient = least_squares(columns, y)
    residual = [observed - sum(coefficient[k] * value ** k for k in range(degree + 1))
                for value, observed in zip(x, y)]
    return coefficient, residual


def poly_predict(coefficient, x):
    return sum(coefficient[k] * x ** k for k in range(len(coefficient)))


def load_inputs(public_only=False):
    contract = json.loads(DECLARATION.read_text(encoding="utf-8"))["口径说明"]["独立复算规范"]
    baseline = (contract["数据与切分"] if public_only else
                json.loads(BASELINE.read_text(encoding="utf-8")))
    if baseline.get("共同窗口_cm^-1") != [1200, 3800] or baseline.get("抽样点数_每角度") != 480:
        raise ValueError("问题2正式基准的窗口或点数改变，不能静默复用")
    folds = baseline["折分"]
    scenarios = contract["硅历史估计器"]["初始情景"]
    if not public_only:
        route = json.loads(ROUTE.read_text(encoding="utf-8"))
        if scenarios != route["共同原型协议"]["共同光学情景"]["初始情景"]:
            raise ValueError("公开初始情景与历史来源不一致，须先明确估计口径")
    tables = {}
    hashes = {}
    for number in range(1, 5):
        path = DATA / f"附件{number}.xlsx"
        hashes[path.name] = hashlib.sha256(path.read_bytes()).hexdigest()
        tables[number] = read_sheet(path)
    baseline_hashes = {item["文件名"]: item["SHA256"] for item in contract["数据与切分"]["输入哈希"]}
    # 问题2正式基准只覆盖其使用的附件1、2；附件3、4是本问新增输入，
    # 不应因缺少于上游基准的哈希而被误判为旧基准。四个附件仍已在上面
    # 完成原始数值读取，后续继续执行7469点和逐行波数对齐校验。
    for name, expected in baseline_hashes.items():
        if hashes.get(name) != expected:
            raise ValueError(f"{name}与问题2正式基准哈希不一致，禁止消费旧基准")
    if not public_only and any(hashes.get(item["文件名"]) != item["SHA256"] for item in baseline["输入哈希"]):
        raise ValueError("问题2正式基准与原附件哈希不一致")
    coordinates = [record[1] for record in tables[1]]
    if any([record[1] for record in tables[number]] != coordinates for number in (2, 3, 4)):
        raise ValueError("四个附件波数未逐行对齐")
    full_keep = [index for index, sigma in enumerate(coordinates)
                 if 1200.0 <= sigma <= 3800.0]
    full_coordinates = [coordinates[index] for index in full_keep]
    if not full_coordinates:
        raise ValueError("1200—3800 cm^-1原始全量窗口为空")
    rows_by_file = {number: {record[0]: index for index, record in enumerate(tables[number])}
                    for number in range(1, 5)}
    reconstructed = [tables[1][full_keep[index * (len(full_keep) - 1) // 479]][0]
                     for index in range(480)]
    source_rows = [reconstructed, reconstructed] if public_only else baseline["抽样源行"]
    if (source_rows != [reconstructed, reconstructed]
            or folds != contract["数据与切分"]["折分"]
            or len(full_keep) != contract["数据与切分"]["全量点数_每角度"]):
        raise ValueError("公开源行规则、折分或全量点数与实际输入不一致")
    sampled = {}
    sample_coordinates = []
    for angle_index, row_list in enumerate(source_rows):
        if len(row_list) != 480:
            raise ValueError("问题2正式基准抽样源行不是480点")
        sample_coordinates = [tables[1][rows_by_file[1][row]][1] for row in row_list]
        if angle_index == 0:
            continue
        if [tables[2][rows_by_file[2][row]][1] for row in row_list] != sample_coordinates:
            raise ValueError("问题2正式基准两角源行不一致")
    for material, numbers in MATERIAL_FILES.items():
        sampled[material] = []
        for number, row_list in zip(numbers, source_rows):
            sampled[material].append([tables[number][rows_by_file[number][row]][2] for row in row_list])
    full_values = {}
    for material, numbers in MATERIAL_FILES.items():
        full_values[material] = [[tables[number][index][2] for index in full_keep]
                                 for number in numbers]
    return {"baseline": baseline, "folds": folds, "scenarios": scenarios,
            "tables": tables, "hashes": hashes, "coordinates": sample_coordinates, "sampled": sampled,
            "full_coordinates": full_coordinates, "full_values": full_values,
            "full_point_count": len(full_coordinates), "source_rows": source_rows,
            "contract": contract}


def n_film(parameters, sigma):
    return parameters[1] + parameters[2] * ((2000.0 / sigma) ** 2 - 1.0)


def valid(parameters, angles=ANGLES):
    if len(parameters) != 5:
        return False
    if any(not low <= value <= high for value, (low, high) in zip(parameters, BOUNDS)):
        return False
    max_sine = max(math.sin(math.radians(angle)) for angle in angles)
    for sigma in (1200.0, 3800.0):
        nf = n_film(parameters, sigma)
        ns = nf + parameters[3]
        if min(nf, ns) <= max_sine + 1e-8:
            return False
    return True


def interface_terms(parameters, sigma, angle):
    thickness, reference, dispersion, contrast, loss = parameters
    nf = n_film(parameters, sigma)
    ns = nf + contrast
    sine2 = math.sin(math.radians(angle)) ** 2
    q0 = math.cos(math.radians(angle))
    q1 = math.sqrt(nf * nf - sine2)
    q2 = math.sqrt(ns * ns - sine2)
    propagation = cmath.exp(complex(-loss * nf / q1,
                                     4.0 * math.pi * thickness * sigma * q1 / 10000.0))
    terms = []
    for eta0, eta1, eta2 in ((q0, q1, q2), (1.0 / q0, nf * nf / q1, ns * ns / q2)):
        r01 = (eta0 - eta1) / (eta0 + eta1)
        r10 = -r01
        r12 = (eta1 - eta2) / (eta1 + eta2)
        t01 = 2.0 * eta0 / (eta0 + eta1)
        t10 = 2.0 * eta1 / (eta0 + eta1)
        first = t01 * t10 * r12 * propagation
        ratio = r10 * r12 * propagation
        terms.append((r01, first, ratio, abs(propagation)))
    return terms


def physical_reflectance(parameters, coordinates, complete, angles=ANGLES, polarization_weight=0.5):
    if not valid(parameters, angles):
        return None
    output = []
    for angle in angles:
        spectrum = []
        for sigma in coordinates:
            fields = []
            for r01, first, ratio, attenuation in interface_terms(parameters, sigma, angle):
                if attenuation > 1.0 + 1e-10 or abs(ratio) >= 1.0:
                    return None
                field = r01 + first / (1.0 - ratio) if complete else r01 + first
                fields.append(abs(field) ** 2)
            intensity = polarization_weight * fields[0] + (1.0 - polarization_weight) * fields[1]
            if not math.isfinite(intensity):
                return None
            spectrum.append(intensity)
        output.append(spectrum)
    return output


def training_info(x, y, train_indices):
    tx = [x[index] for index in train_indices]
    ty = [y[index] for index in train_indices]
    trend, residual = poly_fit(tx, ty, 2)
    return {"train_indices": list(train_indices), "trend": trend, "residual": residual,
            "scale": max(quantile(residual, 0.75) - quantile(residual, 0.25), 0.0001),
            "train_mean": sum(ty) / len(ty)}


def make_case(material, coordinates, values, fold, guard=20.0):
    x = [(sigma - 2500.0) / 1300.0 for sigma in coordinates]
    if fold is None:
        train = list(range(len(coordinates)))
        spec = {"折号": 0, "训练块": list(range(1, 13)), "校准块": [], "测试块": []}
    else:
        spec = fold
        edges = [(coordinates[index - 1] + coordinates[index]) / 2.0 for index in range(40, 480, 40)]
        training_blocks = set(spec["训练块"])
        guards = [edge for block, edge in enumerate(edges, 1)
                  if (block in training_blocks) != (block + 1 in training_blocks)]
        train = []
        for index, sigma in enumerate(coordinates):
            block = bisect.bisect_right(edges, sigma) + 1
            if block in training_blocks and all(abs(sigma - edge) >= guard for edge in guards):
                train.append(index)
    infos = [training_info(x, values[angle_index], train) for angle_index in range(2)]
    for info, series in zip(infos, values):
        info["x"] = list(x)
        info["y"] = list(series)
    return {"材料": material, "折号": spec["折号"], "规格": spec, "coordinates": list(coordinates),
            "x": x, "values": values, "info": infos, "train_indices": train,
            "pool": [], "two": None, "full": None}


def calibrate(physical, info):
    indices = info["train_indices"]
    # info中的x、y在make_case中写入，使用训练子集完成五维线性校准。
    tx = [info["x"][index] for index in indices]
    ty = [info["y"][index] for index in indices]
    pp = [physical[index] for index in indices]
    columns = [[1.0] * len(indices), tx, [value * value for value in tx],
               [value * (1.0 - xx) / 2.0 for value, xx in zip(pp, tx)],
               [value * (1.0 + xx) / 2.0 for value, xx in zip(pp, tx)]]
    unconstrained = least_squares(columns, ty)
    candidates = [(min(GAIN_LIMIT, max(0.0, unconstrained[3])),
                   min(GAIN_LIMIT, max(0.0, unconstrained[4])))]
    candidates.extend((left, right) for left in (0.0, GAIN_LIMIT)
                      for right in (0.0, GAIN_LIMIT))
    best = None
    for gain0, gain1 in candidates:
        residual_target = [observed - gain0 * columns[3][row] - gain1 * columns[4][row]
                           for row, observed in enumerate(ty)]
        baseline = least_squares(columns[:3], residual_target)
        residuals = [observed - (sum(baseline[k] * tx[row] ** k for k in range(3))
                                 + gain0 * columns[3][row] + gain1 * columns[4][row])
                     for row, observed in enumerate(ty)]
        loss = dot(residuals, residuals) / len(indices) / (info["scale"] ** 2)
        candidate = {"baseline": baseline, "gains": (gain0, gain1), "loss": loss}
        if best is None or loss < best["loss"]:
            best = candidate
    return best


def evaluate(parameters, case, complete, angles=ANGLES, polarization_weight=0.5):
    if "response" in case:
        return evaluate_response(parameters, case, complete, angles, polarization_weight)
    physical = physical_reflectance(parameters, case["coordinates"], complete, angles, polarization_weight)
    if physical is None:
        return None
    models = [calibrate(spectrum, info) for spectrum, info in zip(physical, case["info"])]
    loss = sum(model["loss"] for model in models) / 2.0
    return {"parameters": tuple(parameters), "loss": loss, "models": models,
            "physical": physical, "angles": list(angles), "polarization_weight": polarization_weight}


def retain(pool, candidate, limit=3):
    if candidate is None:
        return pool
    candidates = sorted(pool + [candidate], key=lambda item: item["loss"])
    selected = []
    for item in candidates:
        if all(any(abs(item["parameters"][idx] - old["parameters"][idx]) > threshold
                   for idx, threshold in enumerate((0.3, 0.15, 0.05))) for old in selected):
            selected.append(item)
        if len(selected) >= limit:
            break
    return selected


def refine(seed, case, complete, deadline, angles=ANGLES, polarization_weight=0.5, fixed=None, trace=None):
    best = seed
    steps = [0.24, 0.12, 0.05, 0.18, 0.15]
    free = [index for index in range(5) if not fixed or index not in fixed]
    trace = {} if trace is None else trace
    trace.update({"停止原因": "迭代上限", "完成轮数": 0, "评估次数": 0,
                  "最大轮数": 18, "初始步长": list(steps), "自由参数索引": free})
    for _ in range(18):
        if time.monotonic() >= deadline:
            trace["停止原因"] = "时间预算"
            break
        improved = False
        for index in free:
            for direction in (-1.0, 1.0):
                if time.monotonic() >= deadline:
                    trace.update({"停止原因": "时间预算", "最终步长": list(steps)})
                    return best
                trial = list(best["parameters"])
                trial[index] = min(BOUNDS[index][1], max(BOUNDS[index][0], trial[index] + direction * steps[index]))
                candidate = evaluate(trial, case, complete, angles, polarization_weight)
                trace["评估次数"] += 1
                if candidate is not None and candidate["loss"] < best["loss"]:
                    best, improved = candidate, True
        if not improved:
            steps = [step * 0.5 for step in steps]
        trace["完成轮数"] += 1
        if max(steps[index] for index in free) < 0.001:
            trace["停止原因"] = "步长收敛"
            break
    trace["最终步长"] = list(steps)
    return best


def fit_case(case, scenarios, t0, deadline, honor_soft=True):
    trace = {"粗网格次数": 0, "粗网格计划次数": 96 * len(scenarios), "精修": []}
    case["搜索状态"] = trace
    for scene in scenarios:
        for grid_index in range(96):
            if (honor_soft and elapsed(t0) >= SOFT_SECONDS) or time.monotonic() >= deadline:
                break
            thickness = 0.5 + grid_index * 39.5 / 95.0
            parameters = (thickness, scene["参考折射率"], scene["色散系数"], 0.8, 0.0)
            case["pool"] = retain(case["pool"], evaluate(parameters, case, False))
            trace["粗网格次数"] += 1
        if honor_soft and elapsed(t0) >= SOFT_SECONDS:
            break
    if not case["pool"]:
        raise RuntimeError(f"{case['材料']}折{case['折号']}没有可行两束候选")
    case["two"] = case["pool"][0]
    full_pool = []
    for seed in case["pool"]:
        full_seed = list(seed["parameters"])
        full_seed[3], full_seed[4] = 0.8, 0.0
        full_pool = retain(full_pool, evaluate(full_seed, case, True))
    case["full"] = full_pool[0]
    for seed in full_pool:
        refinement = {}
        refined = refine(seed, case, True, deadline, trace=refinement)
        trace["精修"].append(refinement)
        case["full"] = min((case["full"], refined), key=lambda item: item["loss"])
    trace["搜索完整"] = (trace["粗网格次数"] == trace["粗网格计划次数"]
                         and all(item["停止原因"] != "时间预算" for item in trace["精修"]))
    return case


def evaluate_response(parameters, case, complete, angles=ANGLES, polarization_weight=0.5):
    if not valid(parameters, angles):
        return None
    coordinates = np.asarray(case["coordinates"])
    film = parameters[1] + parameters[2] * ((2000.0 / coordinates) ** 2 - 1.0)
    substrate = film + parameters[3]
    physical = []
    models = []
    degree, gain_degree = case["response"]
    for angle, info in zip(angles, case["info"]):
        sine_squared = math.sin(math.radians(angle)) ** 2
        normal_air = math.cos(math.radians(angle))
        normal_film = np.sqrt(film * film - sine_squared)
        normal_substrate = np.sqrt(substrate * substrate - sine_squared)
        propagation = np.exp(-parameters[4] * film / normal_film
            + 4j * math.pi * parameters[0] * coordinates * normal_film / 10000.0)
        intensities = []
        for air, layer, support in ((normal_air, normal_film, normal_substrate),
                (1.0 / normal_air, film * film / normal_film,
                 substrate * substrate / normal_substrate)):
            reflection = (air - layer) / (air + layer)
            back = (layer - support) / (layer + support)
            first = 4.0 * air * layer / (air + layer) ** 2 * back * propagation
            ratio = -reflection * back * propagation
            if np.any(np.abs(ratio) >= 1.0):
                return None
            field = reflection + (first / (1.0 - ratio) if complete else first)
            intensities.append(np.abs(field) ** 2)
        spectrum = polarization_weight * intensities[0] + (1.0 - polarization_weight) * intensities[1]
        cache_key = (degree, gain_degree)
        cache = info.setdefault("response_cache", {})
        if cache_key not in cache:
            indices = np.asarray(info["train_indices"])
            abscissa = np.asarray(info["x"])[indices]
            target = np.asarray(info["y"])[indices]
            background = np.column_stack([abscissa ** power for power in range(degree + 1)])
            inverse = np.linalg.pinv(background, rcond=1e-12)
            cache[cache_key] = (indices, abscissa, target, background, inverse,
                               target - background @ (inverse @ target))
        indices, abscissa, target, background, inverse, residual_target = cache[cache_key]
        response = (spectrum[indices, None] if gain_degree == 0 else
            np.column_stack((spectrum[indices] * (1.0 - abscissa) / 2.0,
                             spectrum[indices] * (1.0 + abscissa) / 2.0)))
        residual_response = response - background @ (inverse @ response)
        gram = residual_response.T @ residual_response
        right = residual_response.T @ residual_target
        free = np.linalg.lstsq(residual_response, residual_target, rcond=1e-12)[0]
        candidates = [np.clip(free, 0.0, GAIN_LIMIT)]
        if gain_degree == 1:
            for fixed_index in (0, 1):
                other = 1 - fixed_index
                for endpoint in (0.0, GAIN_LIMIT):
                    gains = np.zeros(2)
                    gains[fixed_index] = endpoint
                    gains[other] = np.clip((right[other] - gram[other, fixed_index] * endpoint)
                        / gram[other, other], 0.0, GAIN_LIMIT) if gram[other, other] > 1e-24 else 0.0
                    candidates.append(gains)
            candidates.extend(np.array([left, right_endpoint]) for left in (0.0, GAIN_LIMIT)
                              for right_endpoint in (0.0, GAIN_LIMIT))
        best_gains = min(candidates, key=lambda gains: float(np.sum(
            (residual_target - residual_response @ gains) ** 2)))
        coefficients = inverse @ (target - response @ best_gains)
        residual = target - background @ coefficients - response @ best_gains
        gains = [float(best_gains[0])] * 2 if gain_degree == 0 else best_gains.tolist()
        models.append({"baseline": coefficients.tolist(), "gains": gains,
            "loss": float(np.mean(residual ** 2) / info["scale"] ** 2)})
        physical.append(spectrum)
    return {"parameters": tuple(parameters), "loss": sum(model["loss"] for model in models) / 2.0,
        "models": models, "physical": physical, "angles": list(angles),
        "polarization_weight": polarization_weight, "response": list(case["response"])}


def search_fair_pair(case, scenarios, t0, deadline, grid_count=96, pool_limit=3):
    pools = {False: [], True: []}
    trace = {"粗网格次数_每模型": 0, "粗网格计划次数_每模型": grid_count * len(scenarios),
             "共同初值": [], "两束精修": [], "完整精修": [], "物理自由度_每模型": 5,
             "初值选择依据": "两模型各自训练损失候选的固定顺序并集，不读取校准或测试误差",
             "响应自由度_每角每模型": sum(case["response"]) + 2,
             "相同停止规则": "最多18轮或最大自由步长小于0.001", "精修候选": []}
    case["搜索状态"] = trace
    for scene in scenarios:
        for grid_index in range(grid_count):
            if pools[False] and (elapsed(t0) >= SOFT_SECONDS or time.monotonic() >= deadline):
                break
            parameters = (0.5 + grid_index * 39.5 / (grid_count - 1), scene["参考折射率"],
                          scene["色散系数"], 0.8, 0.0)
            for complete in (False, True):
                pools[complete] = retain(pools[complete], evaluate(parameters, case, complete), pool_limit)
            trace["粗网格次数_每模型"] += 1
        if elapsed(t0) >= SOFT_SECONDS or time.monotonic() >= deadline:
            break
    for complete in (False, True):
        for candidate in pools[complete]:
            seed = list(candidate["parameters"])
            if seed not in trace["共同初值"]:
                trace["共同初值"].append(seed)
    if not pools[False] or not pools[True]:
        raise RuntimeError(f"{case['材料']}折{case['折号']}没有共同可行候选")
    case["two"], case["full"] = pools[False][0], pools[True][0]
    for parameters in trace["共同初值"]:
        for complete, key, trace_key in ((False, "two", "两束精修"), (True, "full", "完整精修")):
            seed = evaluate(parameters, case, complete)
            refinement = {}
            refined = refine(seed, case, complete, deadline, trace=refinement)
            case[key] = min((case[key], seed, refined), key=lambda item: item["loss"])
            trace[trace_key].append(refinement)
            trace["精修候选"].append({"模型": "完整往返" if complete else "两束",
                "厚度_um": refined["parameters"][0], "训练标准化损失": refined["loss"],
                "完整物理参数": list(refined["parameters"]), "停止原因": refinement["停止原因"]})
    trace["搜索完整"] = (trace["粗网格次数_每模型"] == trace["粗网格计划次数_每模型"]
        and all(len(trace[key]) == len(trace["共同初值"]) for key in ("两束精修", "完整精修"))
        and all(item["停止原因"] != "时间预算"
                for key in ("两束精修", "完整精修") for item in trace[key]))
    trace["公平性成立"] = trace["搜索完整"]
    return case


def fit_fair_case(case, scenarios, t0, deadline):
    allowed = list(case["train_indices"])
    if len(allowed) > 480:
        allowed = [allowed[index * (len(allowed) - 1) // 479] for index in range(480)]
    coordinates = [case["coordinates"][index] for index in allowed]
    values = [[series[index] for index in allowed] for series in case["values"]]
    size = len(allowed)
    validation_slices = (list(range(size // 4, size // 4 + size // 8)),
                         list(range(5 * size // 8, 5 * size // 8 + size // 8)))
    inner_cases = []
    for inner_number, validation in enumerate(validation_slices, 1):
        lower, upper = coordinates[validation[0]], coordinates[validation[-1]]
        train = [index for index, coordinate in enumerate(coordinates)
                 if coordinate < lower - 20.0 or coordinate > upper + 20.0]
        inner = make_case(case["材料"], coordinates, values, None)
        inner["train_indices"] = train
        inner["info"] = [dict(training_info(inner["x"], series, train), x=inner["x"], y=series)
                         for series in values]
        inner_cases.append((inner, validation))
    diagnostic = {"材料": case["材料"], "折号": case["折号"], "实际用时秒": elapsed(t0),
        "构造": "只取本案例训练集，超过480点按序抽取480点；两段各约1/8，边界保护20 cm^-1",
        "外层测试用途": "已参与前轮诊断，仅探索性复查，不是独立终检",
        "内层切分": [{"内折": inner_number,
            "训练样本序号_原案例从零": [allowed[index] for index in inner["train_indices"]],
            "验证样本序号_原案例从零": [allowed[index] for index in validation]}
            for inner_number, (inner, validation) in enumerate(inner_cases, 1)],
        "候选": [], "默认响应阶数": [2, 1], "选择依据": "两内折、两角、两物理模型等权验证误差；同分先低自由度"}
    choices = [(degree, gain_degree) for degree in range(3) for gain_degree in range(2)]
    selection_deadline = min(deadline, time.monotonic() + max(0.0, deadline - time.monotonic()) * 0.6)
    for response in choices:
        if time.monotonic() >= selection_deadline:
            break
        row = {"背景阶数": response[0], "幅值阶数": response[1], "内折": []}
        for inner_number, (inner, validation) in enumerate(inner_cases, 1):
            if time.monotonic() >= selection_deadline:
                break
            inner["response"] = response
            search_fair_pair(inner, scenarios, t0, selection_deadline, grid_count=48, pool_limit=1)
            scores = {}
            for key, label in (("two", "两束"), ("full", "完整往返")):
                predicted = prediction(inner[key], inner)
                scores[label] = [math.sqrt(sum((predicted[angle_index][index] - series[index]) ** 2
                    for index in validation) / len(validation)) / inner["info"][angle_index]["scale"]
                    for angle_index, series in enumerate(inner["values"])]
            naive = {}
            for label in ("训练均值", "二次趋势"):
                predicted = prediction(baseline_candidate(inner, label), inner)
                naive[label] = [math.sqrt(sum((predicted[angle_index][index] - series[index]) ** 2
                    for index in validation) / len(validation)) / inner["info"][angle_index]["scale"]
                    for angle_index, series in enumerate(inner["values"])]
            row["内折"].append({"内折": inner_number, "物理模型误差_逐角": scores,
                "朴素基线误差_逐角": naive, "两束": parameter_record(inner["two"]),
                "完整往返": parameter_record(inner["full"]), "搜索状态": inner["搜索状态"]})
        row["完整"] = len(row["内折"]) == 2 and all(item["搜索状态"]["公平性成立"] for item in row["内折"])
        if row["完整"]:
            row["验证均值"] = sum(score for item in row["内折"]
                for scores in item["物理模型误差_逐角"].values() for score in scores) / 8.0
        diagnostic["候选"].append(row)
        diagnostic["实际用时秒"] = elapsed(t0)
        atomic_json(RESULT / f"训练内选择_{case['材料']}_折{case['折号']}.json", diagnostic)
    complete = len(diagnostic["候选"]) == len(choices) and all(row["完整"] for row in diagnostic["候选"])
    selected = min(diagnostic["候选"], key=lambda row: (row["验证均值"],
        row["背景阶数"] + row["幅值阶数"], row["背景阶数"])) if complete else None
    case["response"] = (selected["背景阶数"], selected["幅值阶数"]) if selected else (2, 1)
    diagnostic.update({"选型完整": complete, "所选响应阶数": list(case["response"]),
        "截断处理": "六候选两内折未全完成则用预设二次背景、一次幅值，不择取先完成的候选"})
    search_fair_pair(case, scenarios, t0, deadline)
    diagnostic["训练近优分支范围_um"] = {}
    for key, label in (("two", "两束"), ("full", "完整往返")):
        branches = [row for row in case["搜索状态"]["精修候选"] if row["模型"] == label]
        best = case[key]
        diagnostic["训练近优分支范围_um"][label] = {}
        for tolerance in (0.0, 0.05, 0.1):
            thicknesses = [best["parameters"][0]] + [row["厚度_um"] for row in branches
                if row["训练标准化损失"] <= best["loss"] * (1.0 + tolerance) + 1e-12]
            diagnostic["训练近优分支范围_um"][label][f"训练损失容差{100 * tolerance:g}%"] = [
                min(thicknesses), max(thicknesses)]
    case["搜索状态"]["响应选型完整"] = complete
    case["搜索状态"]["公平性成立"] = case["搜索状态"]["搜索完整"] and complete
    case["搜索状态"]["训练内响应选择"] = diagnostic
    diagnostic["实际用时秒"] = elapsed(t0)
    atomic_json(RESULT / f"训练内选择_{case['材料']}_折{case['折号']}.json", diagnostic)
    return case


def run_contract_review(fair=False):
    t0 = time.monotonic()
    output_path = RESULT / ("公平对照.json" if fair else "复算条件执行.json")
    payload = {"问题": 3, "状态": "开始", "实际用时秒": 0.0, "案例": [],
        SENSITIVITY_RESULT_KEY: [], "厚度条件范围_um": {},
        "用途": "建模侧实测重估，不是隔离红队独立复算证据；不覆盖历史答案",
        "合成输入": False, "硅全量条件厚度_um": {}, "模型选择": {},
        "模式": "训练内选响应、同五参数共同初值公平对照" if fair else "历史估计器公开条件重估",
        "估计器版本": "训练内响应选型与约束边界求解_v2" if fair else "历史公开规则_v1",
        "验证属性": "旧外层测试已参与开发，只作探索性复查；本轮无未使用独立外层段",
        "声明哈希": hashlib.sha256(DECLARATION.read_bytes()).hexdigest(),
        "指标含义": {"相对下降_%": "100×(两束留段误差−完整留段误差)/两束留段误差；允许负数",
            "厚度条件范围_um": "本次已完成全量、折间及扰动厚度极值，非置信区间",
            "纯场差异": "固定同一五参数和完整场校准，只切换往返阶数",
            "校准标准化均方根误差": "物理参数与响应仅在训练集估计，再在80个校准点评分；不读取测试响应",
            "条件模型选择": "完整公平两折的校准误差均值择优，同分或配对不完整取两束；测试误差仅作报告",
            "区间构造": "各模型各折各角分别取校准绝对残差的90%经验分位；只报告模型选择前的经验覆盖，不声称选择后共形保证",
            "公平性成立": "训练内响应六候选均完成；共同训练输入、初值、边界、响应复杂度和停止准则，且两边均未被时间截断",
            "训练内响应选择": "只在外层训练点内作两段留出；朴素基线仅诊断，六种背景/幅值阶数组合对两个物理模型等权选型",
            "训练近优分支范围_um": "同一响应下已精修候选在训练损失0%、5%、10%容差内的厚度极差，无概率含义",
            "相对朴素基线下降_%": "100×(1−物理模型标准化RMSE/同角同块基线标准化RMSE)，保留负值；零基线时无定义"}}
    atomic_json(output_path, payload)
    cases = []

    def flush():
        if fair and "切分证据" in payload:
            select_conditional_models(payload, payload["切分证据"]["折分"])
        for material in MATERIAL_ORDER:
            values = [row[model]["厚度_um"] for row in payload["案例"] if row["材料"] == material
                      for model in (("两束", "完整往返") if fair else ("完整往返",))]
            values.extend(row["扰动后厚度_um"] for row in payload[SENSITIVITY_RESULT_KEY]
                          if row["材料"] == material and row.get("状态") == "已重估")
            if values:
                payload["厚度条件范围_um"][material] = [min(values), max(values)]
            selected_range = payload["模型选择"].get(material, {}).get("厚度条件范围_um")
            if fair and selected_range:
                payload["厚度条件范围_um"][material] = selected_range
        if fair:
            payload["核心指标"] = {"碳化硅正式基准厚度_um": payload["碳化硅正式基准"]["厚度_um"]} if "碳化硅正式基准" in payload else {}
            if "碳化硅正式基准" in payload:
                payload["核心指标"]["碳化硅正式条件范围_um"] = payload["碳化硅正式基准"]["条件范围_um"]
            silicon = payload["模型选择"].get("硅", {})
            if "全量条件厚度_um" in silicon:
                payload["核心指标"]["硅全量条件厚度_um"] = silicon["全量条件厚度_um"]
                payload["核心指标"]["硅条件范围_um"] = silicon["厚度条件范围_um"]
        payload["实际用时秒"] = elapsed(t0)
        atomic_json(output_path, payload)

    try:
        data = load_inputs(public_only=True)
        payload["输入哈希"] = data["hashes"]
        payload["切分证据"] = {"样本源行": data["source_rows"], "折分": data["folds"],
                             "训练非训练保护带_cm^-1": 20.0}
        if fair:
            baseline = json.loads(BASELINE.read_text(encoding="utf-8"))
            if any(data["hashes"].get(item["文件名"]) != item["SHA256"] for item in baseline["输入哈希"]):
                raise ValueError("碳化硅正式基准输入摘要不一致")
            payload["碳化硅正式基准"] = {"厚度_um": baseline["全量参数"]["厚度_um"],
                "条件范围_um": baseline["条件范围_um"], "是否替换": False}
        materials = MATERIAL_ORDER if fair else ("硅",)
        payload["计划案例数"] = len(materials) * (1 + len(data["folds"]))
        for material in materials:
            for fold in (None, *data["folds"]):
                if elapsed(t0) >= SOFT_SECONDS:
                    break
                coordinates = data["full_coordinates"] if fold is None else data["coordinates"]
                values = data["full_values"][material] if fold is None else data["sampled"][material]
                case = make_case(material, coordinates, values, fold)
                deadline = min(t0 + SOFT_SECONDS, time.monotonic() + (150.0 if fair else 300.0))
                fitter = fit_fair_case if fair else fit_case
                fitter(case, data["scenarios"], t0, deadline)
                cases.append(case)
                record = {"材料": material, "折号": case["折号"], "每角点数": len(coordinates),
                    "训练样本序号_从零": case["train_indices"], "切分": case["规格"],
                    "训练尺度_比例": [info["scale"] for info in case["info"]],
                    "两束": parameter_record(case["two"]), "完整往返": parameter_record(case["full"]),
                    "搜索状态": case["搜索状态"], "五参数边界": BOUNDS,
                    "触边": {key: [PARAMETER_NAMES[index] for index, (low, high) in enumerate(BOUNDS)
                              if min(abs(candidate["parameters"][index] - low),
                                     abs(candidate["parameters"][index] - high)) <= 1e-10]
                             for key, candidate in (("两束", case["two"]), ("完整往返", case["full"]))}}
                if fold is None:
                    record["全量条件厚度_um"] = case["full"]["parameters"][0]
                    if material == "硅":
                        payload["硅全量条件厚度_um"] = {key: record[key]["厚度_um"]
                                                              for key in ("两束", "完整往返")}
                else:
                    record["校准标准化均方根误差"] = {
                        label: calibration_score(case, case[key]) for key, label in
                        (("two", "两束"), ("full", "完整往返"))}
                    scores = {}
                    for complete, key, label in ((False, "two", "两束"), (True, "full", "完整往返")):
                        scores[label] = score_prediction(case, case[key], label, [])
                    means = {key: sum(row["标准化均方根误差"] for row in rows) / len(rows)
                             for key, rows in scores.items()}
                    record.update({"分块": scores, "标准化均方根误差": means,
                        "相对下降_%": 100.0 * (1.0 - means["完整往返"] / means["两束"])
                                       if means["两束"] else 0.0,
                        "双向留角度": cross_angle_diagnostic(case),
                        "均值基线误差": baseline_score(case, "训练均值"),
                        "二次趋势基线误差": baseline_score(case, "二次趋势")})
                    record["朴素基线分块"] = {label: score_prediction(case, baseline_candidate(case, label), label, [])
                        for label in ("训练均值", "二次趋势")}
                    record["逐角逐块基线对照"] = [{"材料": material, "折号": case["折号"],
                        "模型": label, "附件": score["附件"], "入射角_度": score["入射角_度"],
                        "测试块": score["测试块"], "物理模型标准化均方根误差": score["标准化均方根误差"],
                        "朴素基线": naive, "基线标准化均方根误差": reference["标准化均方根误差"],
                        "相对朴素基线下降_%": 100.0 * (1.0 - score["标准化均方根误差"]
                            / reference["标准化均方根误差"]) if reference["标准化均方根误差"] > 0 else None}
                        for label, model_scores in scores.items() for score in model_scores
                        for naive, references in record["朴素基线分块"].items() for reference in references
                        if (score["附件"], score["测试块"]) == (reference["附件"], reference["测试块"])]
                record["场级数核验"] = field_check(case)
                if fair:
                    record["向量化场核验"] = {}
                    for complete, key, label in ((False, "two", "两束"), (True, "full", "完整往返")):
                        reference = physical_reflectance(case[key]["parameters"], coordinates, complete)
                        error = max(float(np.max(np.abs(spectrum - np.asarray(scalar))))
                            for spectrum, scalar in zip(case[key]["physical"], reference))
                        record["向量化场核验"][label] = {"逐点最大绝对差_反射率比例": error,
                            "每角核验点数": len(coordinates), "通过": error <= 1e-12}
                physical_two = physical_reflectance(case["full"]["parameters"], coordinates, False)
                fixed_two = prediction({**case["full"], "physical": physical_two}, case)
                fixed_full = prediction(case["full"], case)
                record["纯场差异"] = {"固定参数": list(case["full"]["parameters"]),
                    "固定各角响应": parameter_record(case["full"])["各角度校准"],
                    "各角原始场反射率均方根差_比例": [math.sqrt(sum((left - right) ** 2
                        for left, right in zip(full, two)) / len(coordinates))
                        for full, two in zip(case["full"]["physical"], physical_two)],
                    "各角固定响应预测均方根差_比例": [math.sqrt(sum((left - right) ** 2
                        for left, right in zip(full, two)) / len(coordinates))
                        for full, two in zip(fixed_full, fixed_two)]}
                payload["案例"].append(record)
                payload["状态"] = "已形成条件答案，继续验证"
                if fair:
                    select_conditional_models(payload, data["folds"])
                flush()
        controls = {}
        for material in materials:
            for scenario in data["contract"]["硅条件包络"]["扰动生成表"]:
                case = next((item for item in cases if item["材料"] == material
                             and item["折号"] == scenario["折号"]), None)
                if case is None or elapsed(t0) >= SOFT_SECONDS:
                    continue
                label = scenario["情景"]
                record = refit_sensitivity(case, label, t0, t0 + SOFT_SECONDS)
                if record is None:
                    continue
                if label == LOSS_LABELS[0]:
                    controls[(material, case["折号"])] = record
                control = controls.get((material, case["折号"]))
                if label in LOSS_LABELS and control and record["状态"] == control["状态"] == "已重估":
                    record["扣除控制厚度变化_um"] = record["扰动后厚度_um"] - control["扰动后厚度_um"]
                    record["优化继续推进_um"] = control["厚度变化_um"]
                    record["配对可比"] = all(item["停止状态"]["停止原因"] != "时间预算"
                                              for item in (record, control))
                payload[SENSITIVITY_RESULT_KEY].append(record)
                flush()
        payload["硅情景数"] = sum(row["材料"] == "硅" for row in payload["案例"]) + sum(
            row["材料"] == "硅" and row.get("状态") == "已重估" for row in payload[SENSITIVITY_RESULT_KEY])
        payload["硅31情景完整"] = payload["硅情景数"] == 31
        payload["搜索全部完整"] = len(payload["案例"]) == payload["计划案例数"] and all(
            row["搜索状态"]["搜索完整"] for row in payload["案例"])
        payload["计划灵敏度数"] = len(materials) * len(data["contract"]["硅条件包络"]["扰动生成表"])
        payload["灵敏度全部完整"] = len(payload[SENSITIVITY_RESULT_KEY]) == payload["计划灵敏度数"] and all(row.get("状态") == "已重估"
            and row["停止状态"]["停止原因"] != "时间预算" for row in payload[SENSITIVITY_RESULT_KEY])
        payload["分材料验证"] = {}
        for material in materials:
            rows = [row for row in payload["案例"] if row["材料"] == material and row["折号"] > 0]
            if not rows:
                continue
            paired = fair and len(rows) == 2 and all(row["搜索状态"].get("公平性成立", False) for row in rows)
            block_changes = [two["标准化均方根误差"] - full["标准化均方根误差"]
                for row in rows for two, full in zip(row["分块"]["两束"], row["分块"]["完整往返"])]
            payload["分材料验证"][material] = {
                "同自由度完整配对": paired,
                "各折相对下降_%": [row["相对下降_%"] for row in rows],
                "逐角逐块改善": block_changes,
                "两折均胜两束": paired and all(row["相对下降_%"] > 0 for row in rows),
                "跨折跨角块改善方向一致": paired and all(value > 0 for value in block_changes),
                "两折均胜均值基线": paired and all(row["标准化均方根误差"]["完整往返"] < row["均值基线误差"] for row in rows),
                "两折均胜二次趋势": paired and all(row["标准化均方根误差"]["完整往返"] < row["二次趋势基线误差"] for row in rows),
                "结论限制": "不胜出或方向不一致均保留条件厚度与范围，不触发碳化硅正式基准替换"}
            payload["分材料验证"][material]["逐附件跨折"] = [{
                "附件": number, "各折相对下降_%": [100.0 * (1.0 - sum(
                    item["标准化均方根误差"] for item in row["分块"]["完整往返"] if item["附件"] == number)
                    / max(sum(item["标准化均方根误差"] for item in row["分块"]["两束"]
                              if item["附件"] == number), 1e-15)) for row in rows],
                "配对完整": paired} for number in MATERIAL_FILES[material]]
        if fair:
            select_conditional_models(payload, data["folds"])
            for material, validation in payload["分材料验证"].items():
                selected = payload["模型选择"][material]["条件模型"]
                rows = [row for row in payload["案例"] if row["材料"] == material and row["折号"] > 0]
                validation["所选模型"] = selected
                validation["所选模型相对基线下降_%"] = {naive: [100.0 * (1.0 - row["标准化均方根误差"][selected]
                    / row[reference]) if row[reference] > 0 else None for row in rows]
                    for naive, reference in (("训练均值", "均值基线误差"), ("二次趋势", "二次趋势基线误差"))}
                validation["所选模型两折均胜两种朴素基线"] = validation["同自由度完整配对"] and all(
                    change is not None and change > 0 for changes in validation["所选模型相对基线下降_%"].values() for change in changes)
                validation["严格协议2数值条件"] = validation["所选模型两折均胜两种朴素基线"]
                validation["严格协议3数值条件"] = validation["跨折跨角块改善方向一致"]
                validation["独立优势已证明"] = False
        payload["状态"] = "正常结束，保留本次条件数值"
        payload["验证限制"] = "未完整的搜索或情景不宣称复现旧包络；公平性不成立时不得解释高阶收益"
        flush()
        print(json.dumps({key: payload[key] for key in ("状态", "实际用时秒", "硅情景数", "厚度条件范围_um", "模型选择")},
                         ensure_ascii=False), flush=True)
        return 0
    except Exception as error:
        payload.update({"状态": "中断，保留已完成条件数值", "错误": f"{type(error).__name__}: {error}"})
        flush()
        print(json.dumps({"状态": payload["状态"], "错误": payload["错误"], "实际用时秒": elapsed(t0)},
                         ensure_ascii=False), flush=True)
        return 1


def calibration_score(case, candidate):
    indices = [index for block in case["规格"]["校准块"] for index in block_indices(block)]
    predicted = prediction(candidate, case)
    scores = [math.sqrt(sum((predicted[angle_index][index] - observed[index]) ** 2
                            for index in indices) / len(indices)) / case["info"][angle_index]["scale"]
              for angle_index, observed in enumerate(case["values"])]
    return sum(scores) / len(scores)


def select_conditional_models(payload, folds):
    for material in MATERIAL_ORDER:
        rows = [row for row in payload["案例"] if row["材料"] == material and row["折号"] > 0]
        paired = len(rows) == len(folds) and all(row["搜索状态"]["公平性成立"] for row in rows)
        means = {model: sum(row["校准标准化均方根误差"][model] for row in rows) / len(rows)
                 for model in ("两束", "完整往返")} if rows else {}
        selected = "完整往返" if paired and means["完整往返"] < means["两束"] else "两束"
        full = next((row for row in payload["案例"] if row["材料"] == material and row["折号"] == 0), None)
        values = [row[model]["厚度_um"] for row in payload["案例"] if row["材料"] == material
                  for model in ("两束", "完整往返")]
        values.extend(row["扰动后厚度_um"] for row in payload[SENSITIVITY_RESULT_KEY]
                      if row["材料"] == material and row.get("状态") == "已重估")
        record = {"条件模型": selected, "两折公平完整": paired, "校准误差均值": means,
            "选择依据": "两折校准误差均值，测试响应不参与；同分或公平配对不完整取两束极限",
            "厚度条件范围_um": [min(values), max(values)] if values else [],
            "范围构造": "本次两模型全量、折间及已算完整场扰动厚度的联合包络，无真厚度不报告厚度覆盖",
            "置信限制": "不同光学条件可有不同厚度，选择不等于高阶物理确证；只报告选择前各模型反射率经验覆盖"}
        record["外层检验属性"] = "探索性复查：旧测试已参与前轮诊断，不据此声称独立优胜"
        unidentified = any(row[model].get("厚度未识别", False) for row in payload["案例"]
            if row["材料"] == material for model in ("两束", "完整往返"))
        if unidentified:
            record["厚度条件范围_um"] = list(BOUNDS[0])
            record["置信限制"] += "；存在双角响应幅值退化为零的情景，厚度在整个搜索盒未识别；须补测折射率或独立厚度/仪器响应"
        if full is not None:
            record["全量条件厚度_um"] = full[selected]["厚度_um"]
            record["全量搜索完整"] = full["搜索状态"]["搜索完整"]
        elif rows:
            record["留段条件厚度_um"] = rows[0][selected]["厚度_um"]
            record["置信限制"] += "；全量尚未完成，此值仅为首个已算训练折，不能标全量"
        payload["模型选择"][material] = record


def prediction(candidate, case):
    result = []
    for physical, model in zip(candidate["physical"], candidate["models"]):
        baseline, gains = model["baseline"], model["gains"]
        result.append([poly_predict(baseline, xx)
                       + (gains[0] * (1.0 - xx) + gains[1] * (1.0 + xx)) / 2.0 * value
                       for xx, value in zip(case["x"], physical)])
    return result


def parameter_record(candidate):
    p = candidate["parameters"]
    return {"厚度_um": p[0], "参考折射率": p[1], "色散系数": p[2], "衬底折射率对比": p[3],
            "响应阶数_背景与幅值": candidate.get("response", [2, 1]),
            "厚度未识别": all(abs(gain) <= 1e-10 for model in candidate["models"] for gain in model["gains"]),
            "有效往返损耗": p[4], "训练标准化损失": candidate["loss"],
            "各角度校准": [{"入射角_度": angle,
                            "背景实际次数": len(model["baseline"]) - 1,
                            "二次基线系数": list(model["baseline"]) + [0.0] * (3 - len(model["baseline"])),
                            "一次幅值端点": list(model["gains"])}
                           for angle, model in zip(candidate.get("angles", ANGLES), candidate["models"])]}


def block_indices(block):
    return list(range((block - 1) * 40, block * 40))


def conformal_half_width(predicted, observed, calibration_blocks):
    errors = sorted(abs(predicted[index] - observed[index])
                    for block in calibration_blocks for index in block_indices(block))
    return errors[math.ceil(0.9 * len(errors)) - 1] if errors else 0.0


def score_prediction(case, candidate, model_name, csv_rows):
    pred = prediction(candidate, case)
    spec = case["规格"]
    rows = []
    for angle_index, angle in enumerate(ANGLES):
        observed = case["values"][angle_index]
        width = conformal_half_width(pred[angle_index], observed, spec["校准块"])
        for block in spec["测试块"]:
            indices = block_indices(block)
            residuals = [pred[angle_index][i] - observed[i] for i in indices]
            score = math.sqrt(sum(value * value for value in residuals) / len(indices)) / case["info"][angle_index]["scale"]
            hits = sum(abs(value) <= width for value in residuals)
            row = {"材料": case["材料"], "折号": case["折号"], "模型": model_name,
                   "附件": MATERIAL_FILES[case["材料"]][angle_index], "入射角_度": angle,
                   "测试块": block, "点数": len(indices), "标准化均方根误差": score,
                   "均方根误差_百分点": 100.0 * math.sqrt(sum(value * value for value in residuals) / len(indices)),
                   "名义覆盖率": 0.9, "经验覆盖率": hits / len(indices), "区间半宽_比例": width,
                   "训练尺度_比例": case["info"][angle_index]["scale"]}
            rows.append(row)
            for index in indices:
                csv_rows.append({"材料": case["材料"], "折号": case["折号"], "模型": model_name,
                                 "附件": MATERIAL_FILES[case["材料"]][angle_index], "入射角_度": angle,
                                 "测试块": block, "波数_cm^-1": case["coordinates"][index],
                                 "实测反射率_比例": observed[index], "预测反射率_比例": pred[angle_index][index],
                                 "残差_比例": pred[angle_index][index] - observed[index],
                                 "区间下界_比例": pred[angle_index][index] - width,
                                 "区间上界_比例": pred[angle_index][index] + width})
    return rows


def baseline_score(case, mode):
    rows = score_prediction(case, baseline_candidate(case, mode), mode, [])
    return sum(row["标准化均方根误差"] for row in rows) / len(rows)


def baseline_candidate(case, mode):
    return {"physical": [[0.0] * len(case["coordinates"]) for info in case["info"]],
        "models": [{"baseline": [info["train_mean"]] if mode == "训练均值" else info["trend"],
                    "gains": [0.0, 0.0]} for info in case["info"]]}


def cross_angle_diagnostic(case):
    physical = case["full"]["physical"]
    scores = {}
    for source, target in ((0, 1), (1, 0)):
        fitted = case["full"]["models"][source]
        predicted = [poly_predict(fitted["baseline"], xx)
                     + (fitted["gains"][0] * (1.0 - xx) + fitted["gains"][1] * (1.0 + xx)) / 2.0 * value
                     for xx, value in zip(case["x"], physical[target])]
        observed = case["values"][target]
        indices = [index for block in case["规格"]["测试块"] for index in block_indices(block)]
        scores[f"{int(ANGLES[source])}度拟合预测{int(ANGLES[target])}度_标准化均方根误差"] = (
            math.sqrt(sum((predicted[i] - observed[i]) ** 2 for i in indices) / len(indices))
            / case["info"][target]["scale"])
    return scores


def field_check(case):
    parameters = case["full"]["parameters"]
    max_series_error = 0.0
    max_tail_bound = 0.0
    max_ratio = 0.0
    max_attenuation = 0.0
    for angle in ANGLES:
        for sigma in case["coordinates"][::24]:
            for r01, first, ratio, attenuation in interface_terms(parameters, sigma, angle):
                closed = r01 + first / (1.0 - ratio)
                partial = complex(r01)
                term = first
                for _ in range(512):
                    partial += term
                    term *= ratio
                    if abs(term) / max(1.0 - abs(ratio), 1e-15) < 1e-12:
                        break
                max_series_error = max(max_series_error, abs(partial - closed))
                max_tail_bound = max(max_tail_bound, abs(term) / max(1.0 - abs(ratio), 1e-15))
                max_ratio = max(max_ratio, abs(ratio))
                max_attenuation = max(max_attenuation, attenuation)
    return {"材料": case["材料"], "折号": case["折号"], "几何级数与闭式最大绝对误差": max_series_error,
            "尾项绝对上界": max_tail_bound, "最大往返乘子模": max_ratio,
            "最大传播因子模": max_attenuation,
            "通过": max_series_error <= max_tail_bound + 1e-10 and max_ratio < 1.0 and max_attenuation <= 1.0 + 1e-10}


def slice_case(case, lower, upper):
    keep = [index for index, sigma in enumerate(case["coordinates"]) if lower <= sigma <= upper]
    remap = {old: new for new, old in enumerate(keep)}
    clone = {key: value for key, value in case.items() if key not in ("coordinates", "x", "values", "info", "train_indices")}
    clone["coordinates"] = [case["coordinates"][i] for i in keep]
    clone["x"] = [case["x"][i] for i in keep]
    clone["values"] = [[series[i] for i in keep] for series in case["values"]]
    clone["train_indices"] = [remap[i] for i in case["train_indices"] if i in remap]
    clone["info"] = [training_info(clone["x"], series, clone["train_indices"]) for series in clone["values"]]
    for info, series in zip(clone["info"], clone["values"]):
        info["x"] = list(clone["x"])
        info["y"] = list(series)
    return clone


def refit_sensitivity(case, label, t0, deadline):
    base = list(case["full"]["parameters"])
    scenarios = []
    trial_case = case
    requested_range = [1200.0, 3800.0]
    if label == "衬底对比减20%":
        scenarios.append((base[3] * 0.8 if base[3] else -0.1, 3, ANGLES, 0.5))
    elif label == "衬底对比加20%":
        scenarios.append((base[3] * 1.2 if base[3] else 0.1, 3, ANGLES, 0.5))
    elif label in LOSS_LABELS:
        requested_loss = {LOSS_LABELS[0]: base[4], LOSS_LABELS[1]: base[4] * 0.8,
                          LOSS_LABELS[2]: base[4] * 1.2, LOSS_LABELS[3]: 0.1,
                          LOSS_LABELS[4]: 0.2}[label]
        if label in LOSS_LABELS[3:] and base[4] != 0.0:
            return None
        scenarios.append((requested_loss, 4, ANGLES, 0.5))
    elif label == "偏振权重全s":
        scenarios.append((1.0, None, ANGLES, 1.0))
    elif label == "偏振权重全p":
        scenarios.append((0.0, None, ANGLES, 0.0))
    elif label == "入射角同时减0.5度":
        scenarios.append((None, None, tuple(angle - 0.5 for angle in ANGLES), 0.5))
    elif label == "入射角同时加0.5度":
        scenarios.append((None, None, tuple(angle + 0.5 for angle in ANGLES), 0.5))
    elif label == "参考折射率减3%":
        scenarios.append((base[1] * 0.97, 1, ANGLES, 0.5))
    elif label == "参考折射率加3%":
        scenarios.append((base[1] * 1.03, 1, ANGLES, 0.5))
    elif label == "窗口下界加100":
        trial_case = slice_case(case, 1300.0, 3800.0)
        requested_range = [1300.0, 3800.0]
        scenarios.append((None, None, ANGLES, 0.5))
    elif label == "窗口上界减100":
        trial_case = slice_case(case, 1200.0, 3700.0)
        requested_range = [1200.0, 3700.0]
        scenarios.append((None, None, ANGLES, 0.5))
    else:
        return None
    fixed_value, fixed_index, angles, weight = scenarios[0]
    # 保留窗口情景的切窗对象；这里不能恢复为原始 case，否则窗口扰动不会改变重估样本。
    sample_range = ([trial_case["coordinates"][0], trial_case["coordinates"][-1]]
                    if trial_case["coordinates"] else [])
    sample_count = len(trial_case["coordinates"])
    trial = list(base)
    fixed = set()
    details = {"基准参数向量": base, "基准完整参数": parameter_record(case["full"]),
               "基准固定参数值": base[fixed_index] if fixed_index is not None else None,
               "请求固定参数值": fixed_value, "实际固定参数值": None}
    if fixed_index is not None:
        trial[fixed_index] = min(BOUNDS[fixed_index][1], max(BOUNDS[fixed_index][0], fixed_value))
        if fixed_index == 4:
            trial[fixed_index] = fixed_value
        details["实际固定参数值"] = trial[fixed_index]
        fixed.add(fixed_index)
    candidate = evaluate(trial, trial_case, True, angles, weight)
    if candidate is None:
        return {"材料": case["材料"], "折号": case["折号"], "情景": label, "状态": "不可行",
                **details, "停止状态": {"停止原因": "请求参数不可行，未裁剪损耗"},
                "请求窗口_cm^-1": requested_range, "样本范围_cm^-1": sample_range,
                "样本点数": sample_count}
    trace = {}
    candidate = refine(candidate, trial_case, True, deadline, angles, weight, fixed, trace)
    return {"材料": case["材料"], "折号": case["折号"], "情景": label, "状态": "已重估",
            **details, "重估参数向量": list(candidate["parameters"]),
            "重估完整参数": parameter_record(candidate), "停止状态": trace, "实际用时秒": elapsed(t0),
            "基准厚度_um": base[0], "扰动后厚度_um": candidate["parameters"][0],
            "厚度变化_um": candidate["parameters"][0] - base[0],
            "相对变化_%": 100.0 * (candidate["parameters"][0] / base[0] - 1.0),
            "训练标准化损失": candidate["loss"], "固定参数": fixed_index,
            "偏振s权重": weight, "入射角_度": list(angles),
            "请求窗口_cm^-1": requested_range, "样本范围_cm^-1": sample_range,
            "样本点数": sample_count}


def save_sensitivity_baselines(data, cases, t0, checks=None):
    payload = {"输入哈希": data["hashes"], "问题2基准哈希": hashlib.sha256(BASELINE.read_bytes()).hexdigest(),
               "情景来源哈希": hashlib.sha256(ROUTE.read_bytes()).hexdigest(), "参数顺序": PARAMETER_NAMES,
               "实际用时秒": elapsed(t0), "每折完整基准": []}
    for case in cases:
        if case["full"] is not None:
            payload["每折完整基准"].append({"材料": case["材料"], "折号": case["折号"],
                "参数向量": list(case["full"]["parameters"]), "完整参数": parameter_record(case["full"]),
                "训练索引": case["train_indices"], "折分": case["规格"], "波数_cm^-1": case["coordinates"],
                "训练尺度_比例": [info["scale"] for info in case["info"]],
                "恢复核对": (checks or {}).get((case["材料"], case["折号"]), "本轮拟合后、任何扰动前冻结")})
    atomic_json(RESULT / "灵敏度_完整基准.json", payload)


def restore_sensitivity_baselines(data, cases, t0):
    snapshot_path = RESULT / "灵敏度_完整基准.json"
    snapshot = json.loads(snapshot_path.read_text(encoding="utf-8")) if snapshot_path.exists() else {}
    if snapshot and (snapshot["输入哈希"] != data["hashes"]
            or snapshot["问题2基准哈希"] != hashlib.sha256(BASELINE.read_bytes()).hexdigest()
            or snapshot["情景来源哈希"] != hashlib.sha256(ROUTE.read_bytes()).hexdigest()):
        raise ValueError("完整折基准与当前附件、切分或初始情景不一致")
    saved = {(row["材料"], row["折号"]): row for row in snapshot.get("每折完整基准", [])}
    with (RESULT / "留段预测.csv").open(encoding="utf-8-sig", newline="") as stream:
        historical = list(csv.DictReader(stream))
    prior = json.loads((RESULT / "厚度结果.json").read_text(encoding="utf-8"))[SENSITIVITY_RESULT_KEY]
    checks = {}
    for case in cases:
        identity = (case["材料"], case["折号"])
        if identity in saved:
            row = saved[identity]
            if (row["训练索引"] != case["train_indices"] or row["折分"] != case["规格"]
                    or row["波数_cm^-1"] != case["coordinates"]):
                raise ValueError("恢复基准的训练索引、折分或波数不一致")
            case["full"] = evaluate(row["参数向量"], case, True)
        else:
            fit_case(case, data["scenarios"], t0, t0 + SOFT_SECONDS)
        if case["full"] is None:
            raise ValueError("恢复的完整基准不可行，不能用全量参数或单独厚度代替")
        predicted = prediction(case["full"], case)
        coordinates = {sigma: index for index, sigma in enumerate(case["coordinates"])}
        old_rows = [row for row in historical if (row["材料"], int(row["折号"])) == identity
                    and row["模型"] == "完整往返"]
        old_thickness = [row["基准厚度_um"] for row in prior
                         if (row["材料"], row["折号"]) == identity and "基准厚度_um" in row]
        if not old_rows or not old_thickness:
            raise ValueError("历史留段预测或折基准厚度缺失，不能声称恢复原基准")
        error = max(abs(predicted[ANGLES.index(float(row["入射角_度"]))][coordinates[float(row["波数_cm^-1"])]]
                        - float(row["预测反射率_比例"])) for row in old_rows)
        thickness_error = max(abs(case["full"]["parameters"][0] - value) for value in old_thickness)
        checks[identity] = {"来源": "完整参数快照" if identity in saved else "原算法与原切分确定性恢复",
                            "历史预测最大差_比例": error, "历史厚度最大差_um": thickness_error,
                            "匹配": error <= 1e-10 and thickness_error <= 1e-10}
        save_sensitivity_baselines(data, cases, t0, checks)
        if not checks[identity]["匹配"]:
            raise ValueError("恢复结果未复现历史基准；已保存差异，不把新拟合伪装成原基准")


def run_loss_sensitivity(cases, t0, deadline):
    rows, checks = [], []
    payload = {SENSITIVITY_RESULT_KEY: rows, "自动检查": checks, "状态": "计算中", "实际用时秒": elapsed(t0),
        "计划情景总数": sum(5 if case["full"]["parameters"][4] == 0.0 else 3
                            for case in cases if case["full"] is not None),
        "完整基准哈希": hashlib.sha256((RESULT / "灵敏度_完整基准.json").read_bytes()).hexdigest(),
        "指标含义": {"厚度变化_um": "重估减原基准，含优化推进；不能单独解释为损耗效应",
            "优化继续推进_um": "同初值、同18轮上限且固定原损耗的控制减原基准",
            "扣除控制厚度变化_um": "扰动减控制；配对可比为真时才作为条件损耗效应",
            "扣除控制相对变化_%": "100乘以扰动厚度与控制厚度之比减一；非统计置信区间"}}
    atomic_json(RESULT / "灵敏度_损耗复算.json", payload)
    for case in cases:
        if case["full"] is None:
            continue
        base_loss = case["full"]["parameters"][4]
        labels = LOSS_LABELS if base_loss == 0.0 else LOSS_LABELS[:3]
        control = None
        for label in labels:
            if time.monotonic() >= deadline:
                break
            record = refit_sensitivity(case, label, t0, deadline)
            if label == LOSS_LABELS[0]:
                control = record
            complete = record["状态"] == "已重估" and control["状态"] == "已重估"
            comparable = complete and all(item["停止状态"]["停止原因"] != "时间预算"
                                          for item in (control, record))
            record.update({"情景类别": "未扰动控制" if label == LOSS_LABELS[0] else
                           "相对扰动" if label in LOSS_LABELS[1:3] else "单侧绝对增加",
                           "配对可比": comparable})
            if complete:
                record.update({"控制厚度_um": control["扰动后厚度_um"],
                    "优化继续推进_um": control["厚度变化_um"],
                    "扣除控制厚度变化_um": record["扰动后厚度_um"] - control["扰动后厚度_um"],
                    "扣除控制相对变化_%": 100.0 * (record["扰动后厚度_um"] / control["扰动后厚度_um"] - 1.0),
                    "场级数核验": field_check({**case, "full": {"parameters": record["重估参数向量"]}})})
            expected = {LOSS_LABELS[0]: base_loss, LOSS_LABELS[1]: base_loss * 0.8,
                        LOSS_LABELS[2]: base_loss * 1.2, LOSS_LABELS[3]: 0.1, LOSS_LABELS[4]: 0.2}[label]
            fixed_ok = record["请求固定参数值"] == record["实际固定参数值"] == expected
            fixed_ok = fixed_ok and (not complete or record["重估参数向量"][4] == expected)
            zero_check = base_loss == 0.0 and label in LOSS_LABELS[1:3]
            identical = (all(abs(left - right) <= 1e-12 for left, right in
                            zip(record["重估参数向量"], control["重估参数向量"])) if comparable else None)
            check = {"材料": case["材料"], "折号": case["折号"], "情景": label,
                     "标签与实参一致": fixed_ok, "零基准相对扰动": zero_check,
                     "零基准实际为零": record["实际固定参数值"] == 0.0 if zero_check else None,
                     "零扰动与控制重估一致": identical if zero_check else None,
                     "配对可比": comparable}
            rows.append(record)
            checks.append(check)
            payload["实际用时秒"] = elapsed(t0)
            atomic_json(RESULT / "灵敏度_损耗复算.json", payload)
            if not fixed_ok or (zero_check and comparable and not identical):
                raise AssertionError("损耗标签、实际固定值或零扰动控制不一致，已保存自动检查")
    payload.update({"状态": "情景齐全" if len(rows) == payload["计划情景总数"] else "预算截止，保留部分数值",
        "实际用时秒": elapsed(t0), "零基准相对扰动实参": [{"材料": case["材料"], "折号": case["折号"],
            "实际固定值": [row["实际固定参数值"] for row in rows if row["材料"] == case["材料"]
                           and row["折号"] == case["折号"] and row["情景"] in LOSS_LABELS[1:3]],
            "应有固定值": [0.0, 0.0]} for case in cases
            if case["full"] is not None and case["full"]["parameters"][4] == 0.0]})
    atomic_json(RESULT / "灵敏度_损耗复算.json", payload)
    return rows


def save_sensitivity_rows(rows):
    fields = list(dict.fromkeys(key for row in rows for key in row))
    serialized = [{key: json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value
                   for key, value in row.items()} for row in rows]
    write_csv(RESULT / "灵敏度.csv", serialized, fields)


def rerun_loss_sensitivity():
    t0 = time.monotonic()
    original = json.loads((RESULT / "厚度结果.json").read_text(encoding="utf-8"))
    data = load_inputs()
    cases = [make_case(material, data["coordinates"], data["sampled"][material], fold)
             for material in MATERIAL_ORDER for fold in data["folds"]]
    restore_sensitivity_baselines(data, cases, t0)
    corrected = run_loss_sensitivity(cases, t0, t0 + SOFT_SECONDS)
    expected = sum(5 if case["full"]["parameters"][4] == 0.0 else 3 for case in cases)
    if len(corrected) == expected:
        retained = [row for row in original[SENSITIVITY_RESULT_KEY] if row["情景"] not in LOSS_LABELS]
        original[SENSITIVITY_RESULT_KEY] = retained + corrected
        original["损耗复算实际用时秒"] = elapsed(t0)
        atomic_json(RESULT / "厚度结果.json", original)
        save_sensitivity_rows(original[SENSITIVITY_RESULT_KEY])
    print(json.dumps({"问题": 3, "模式": "仅重算损耗灵敏度", "完成情景数": len(corrected),
                      "计划情景数": expected, "正式灵敏度已更新": len(corrected) == expected,
                      "分步结果": "灵敏度_损耗复算.json", "实际用时秒": elapsed(t0)}, ensure_ascii=False), flush=True)
    return 0 if len(corrected) == expected else 1


def build_judgement(cases, score_rows, field_rows, baseline):
    support = []
    for material in MATERIAL_ORDER:
        numbers = MATERIAL_FILES[material]
        material_rows = [row for row in score_rows if row["材料"] == material and row["模型"] in ("两束", "完整往返")]
        per_attachment = []
        for number in numbers:
            rows = [row for row in material_rows if row["附件"] == number]
            two = [row["标准化均方根误差"] for row in rows if row["模型"] == "两束"]
            full = [row["标准化均方根误差"] for row in rows if row["模型"] == "完整往返"]
            differences = [a - b for a, b in zip(two, full)]
            per_attachment.append({"附件": f"附件{number}", "两束减完整_逐块": differences,
                                  "两折平均改善": sum(differences) / len(differences) if differences else None,
                                  "两折是否均改善": bool(differences) and all(x > 0 for x in differences)})
        if material == "碳化硅":
            conclusion = "数据兼容高阶往返，但相干长度、空间重叠和仪器分辨率未提供，不能把必要条件写成物理确证。"
        else:
            conclusion = "完整往返支持在连续留段上折间不一致；仍用完整模型交付硅厚度，但不据此强判多光束已发生。"
        support.append({"材料": material, "附件证据": per_attachment, "结论": conclusion,
                        "场级数核验": [row for row in field_rows if row["材料"] == material]})
    carbon_base = baseline["全量参数"]["厚度_um"]
    return {"必要条件与可观测性": {
        "非零界面返回": "由复界面系数计算检查；不是反射率相关系数证明",
        "时间相干": "附件未提供，无法核验",
        "空间重叠": "附件未提供，无法核验",
        "仪器分辨率": "附件未提供；0.4821166711 cm^-1为平均采样步长，不是标称分辨率"},
            "分材料判定": support,
            "碳化硅条件分支": {"问题2正式基准厚度_um": carbon_base,
                                "正式修正是否触发": False,
                                "理由": "必要条件中的相干、空间重叠、分辨率均缺失，且本轮仅有预测兼容性而无独立检出率校准；保留问题2正式基准。",
                                "修正模型仍输出": True}}


def main():
    t0 = time.monotonic()
    RESULT.mkdir(parents=True, exist_ok=True)
    skeleton = {"问题": 3, "路线": "往返衰减场反演", "运行状态": "读取真实附件", "实际用时秒": 0.0,
                "核心指标": {}, "警告": [], "合成输入": False, SENSITIVITY_RESULT_KEY: []}
    if not (RESULT / "厚度结果.json").exists():
        atomic_json(RESULT / "厚度结果.json", skeleton)
    if not (RESULT / "条件判定.json").exists():
        atomic_json(RESULT / "条件判定.json", {"问题": 3, "运行状态": "尚未完成", "实际用时秒": 0.0})
    try:
        data = load_inputs()
        folds = data["folds"]
        cases = []
        for material in MATERIAL_ORDER:
            for fold in folds:
                cases.append(make_case(material, data["coordinates"], data["sampled"][material], fold))
        atomic_json(RESULT / "厚度结果.json", {**skeleton, "运行状态": "已读入四附件及问题2正式基准",
                                               "输入哈希": data["hashes"],
                                               "计划折分点数_每角度": 480,
                                               "全量原始点数_每角度": data["full_point_count"],
                                               "窗口_cm^-1": [1200, 3800], "实际用时秒": elapsed(t0)})
        for case in cases:
            if elapsed(t0) >= SOFT_SECONDS:
                skeleton["警告"].append("软截止已到，保留已完成折，不静默拼接未完成折")
                break
            fit_case(case, data["scenarios"], t0, t0 + SOFT_SECONDS)
            save_sensitivity_baselines(data, cases, t0)
            thickness_payload = []
            for item in cases:
                if item["full"] is not None:
                    thickness_payload.append({"材料": item["材料"], "折号": item["折号"],
                                              "两束": parameter_record(item["two"]),
                                              "完整往返": parameter_record(item["full"]),
                                              "厚度变化_um": item["full"]["parameters"][0] - item["two"]["parameters"][0]})
            atomic_json(RESULT / "厚度结果.json", {**skeleton, "运行状态": "折间拟合中",
                                                   "厚度及单位": thickness_payload, "实际用时秒": elapsed(t0)})
        score_rows, csv_rows, field_rows = [], [], []
        for case in cases:
            if case["full"] is None:
                continue
            score_rows.extend(score_prediction(case, case["two"], "两束", csv_rows))
            score_rows.extend(score_prediction(case, case["full"], "完整往返", csv_rows))
            field_rows.append(field_check(case))
        write_csv(RESULT / "留段预测.csv", csv_rows,
                  ["材料", "折号", "模型", "附件", "入射角_度", "测试块", "波数_cm^-1",
                   "实测反射率_比例", "预测反射率_比例", "残差_比例", "区间下界_比例", "区间上界_比例"])
        main_rows = [row for row in score_rows if row["模型"] == "完整往返"]
        two_rows = [row for row in score_rows if row["模型"] == "两束"]
        main_score = sum(row["标准化均方根误差"] for row in main_rows) / len(main_rows) if main_rows else None
        two_score = sum(row["标准化均方根误差"] for row in two_rows) / len(two_rows) if two_rows else None
        baseline_scores = {}
        for mode in ("训练均值", "二次趋势"):
            values = [baseline_score(case, mode) for case in cases if case["full"] is not None]
            baseline_scores[mode] = sum(values) / len(values) if values else None
        coverage = {"名义覆盖率": 0.9, "构造方法": "各折各角校准块绝对残差的向上取整90%经验分位对称区间",
                    "测试经验覆盖率": sum(row["经验覆盖率"] * row["点数"] for row in main_rows) / sum(row["点数"] for row in main_rows)
                    if main_rows else None,
                    "测试点数": sum(row["点数"] for row in main_rows),
                    "覆盖点数": sum(round(row["经验覆盖率"] * row["点数"]) for row in main_rows),
                    "限制": "谱内连续相关，经验覆盖不等于分布无关保证，也不是厚度覆盖率。"}
        atomic_json(RESULT / "同口径对照.json", {"问题": 3, "主方法": "往返衰减场反演",
            "主方法连续留段标准化均方根误差": main_score, "两束对照连续留段标准化均方根误差": two_score,
            "基线对照": baseline_scores, "双向留角度": [cross_angle_diagnostic(case) for case in cases if case["full"] is not None],
            "分块评分": score_rows, "场级数核验": field_rows, "实际用时秒": elapsed(t0)})
        atomic_json(RESULT / "区间覆盖.json", {"问题": 3, "经验覆盖率": coverage,
                                               "区间对象": "反射率点预测", "实际用时秒": elapsed(t0)})
        judgement = build_judgement(cases, score_rows, field_rows, data["baseline"])
        atomic_json(RESULT / "条件判定.json", {**judgement, "运行状态": "已形成必要条件与条件分支",
                                               "实际用时秒": elapsed(t0)})
        # 全量原始点是主答案口径，优先于可选的灵敏度扩展；其对象不能被
        # 问题2的480点抽样替代，也不能因前面的留段诊断耗尽软截止而跳过。
        full_data = []
        for material in MATERIAL_ORDER:
            if elapsed(t0) >= HARD_SECONDS:
                break
            all_case = make_case(material, data["full_coordinates"], data["full_values"][material], None)
            fit_case(all_case, data["scenarios"], t0, t0 + HARD_SECONDS, honor_soft=False)
            if all_case["full"] is not None:
                full_data.append({"材料": material, "数据对象": "1200—3800 cm^-1原始全量",
                                  "窗口_cm^-1": [1200, 3800],
                                  "每角度原始点数": data["full_point_count"],
                                  "全量完整往返": parameter_record(all_case["full"]),
                                  "全量两束": parameter_record(all_case["two"])})
            atomic_json(RESULT / "厚度结果.json", {**skeleton, "运行状态": "原始全量拟合中",
                                                   "输入哈希": data["hashes"],
                                                   "计划折分点数_每角度": 480,
                                                   "全量原始点数_每角度": data["full_point_count"],
                                                   "全量数据": full_data,
                                                   "实际用时秒": elapsed(t0)})
        sensitivity_rows = run_loss_sensitivity(cases, t0, t0 + SOFT_SECONDS)
        if elapsed(t0) < HARD_SECONDS:
            labels = ("衬底对比减20%", "衬底对比加20%",
                      "参考折射率减3%", "参考折射率加3%", "偏振权重全s", "偏振权重全p",
                      "入射角同时减0.5度", "入射角同时加0.5度", "窗口下界加100", "窗口上界减100")
            for case in cases:
                if case["full"] is None or elapsed(t0) >= SOFT_SECONDS:
                    break
                for label in labels:
                    if elapsed(t0) >= SOFT_SECONDS:
                        break
                    record = refit_sensitivity(case, label, t0, t0 + SOFT_SECONDS)
                    if record:
                        sensitivity_rows.append(record)
        save_sensitivity_rows(sensitivity_rows)
        base_thickness = data["baseline"]["全量参数"]["厚度_um"]
        completed_cases = sum(case["full"] is not None for case in cases)
        thickness_result = {"问题": 3, "运行状态": "已完成数值交付" if completed_cases == 4 else "已形成部分数值",
            "输入哈希": data["hashes"], "窗口_cm^-1": [1200, 3800],
            "计划折分点数_每角度": 480, "全量原始点数_每角度": data["full_point_count"],
            "硅": {"折间完整往返": [{"折号": case["折号"], "厚度_um": case["full"]["parameters"][0]}
                                      for case in cases if case["材料"] == "硅" and case["full"] is not None],
                    "全量条件估计": next((item for item in full_data if item["材料"] == "硅"), None),
                    "不确定性说明": "折间与灵敏度范围为条件情景，不是统计置信区间。"},
            "碳化硅": {"问题2正式基准厚度_um": base_thickness,
                      "问题2正式条件范围_um": data["baseline"]["条件范围_um"],
                      "条件分支是否替换": judgement["碳化硅条件分支"]["正式修正是否触发"],
                      "全量往返对照": next((item for item in full_data if item["材料"] == "碳化硅"), None)},
            SENSITIVITY_RESULT_KEY: sensitivity_rows, "场级数核验": field_rows,
            "不确定性协议": "反射率用校准块90%经验区间；厚度仅报告条件折间/参数灵敏度范围，实测无真值不报告厚度覆盖。",
            "实际用时秒": elapsed(t0)}
        if main_score is not None:
            thickness_result["核心指标"] = {"连续留段标准化均方根误差": main_score,
                                      "两束对照连续留段标准化均方根误差": two_score,
                                      "反射率测试经验覆盖率": coverage["测试经验覆盖率"]}
        atomic_json(RESULT / "厚度结果.json", thickness_result)
        skeleton["运行状态"] = thickness_result["运行状态"]
        skeleton["核心指标"] = thickness_result.get("核心指标", {})
        skeleton["实际用时秒"] = elapsed(t0)
        atomic_json(RESULT / "执行状态.json", skeleton)
        print(json.dumps({"问题": 3, "路线": "往返衰减场反演", "运行状态": skeleton["运行状态"],
                          "核心指标": skeleton["核心指标"], "实际用时秒": skeleton["实际用时秒"]},
                         ensure_ascii=False, allow_nan=False), flush=True)
        return 0
    except Exception as error:
        skeleton["运行状态"] = "中断，保留已完成数值"
        skeleton["失败原因"] = f"{type(error).__name__}: {error}"
        skeleton["实际用时秒"] = elapsed(t0)
        atomic_json(RESULT / "执行状态.json", skeleton)
        result_path = RESULT / "厚度结果.json"
        if result_path.exists():
            preserved = json.loads(result_path.read_text(encoding="utf-8"))
            preserved["运行状态"] = "中断，保留已完成数值"
            preserved["失败原因"] = f"{type(error).__name__}: {error}"
            preserved["实际用时秒"] = elapsed(t0)
            atomic_json(result_path, preserved)
        else:
            atomic_json(result_path, {**skeleton, "核心指标": skeleton.get("核心指标", {})})
        print(json.dumps(skeleton, ensure_ascii=False, allow_nan=False), flush=True)
        return 1


def load_v3_engine():
    source = ROOT / "求解/问题3/升格3/求解.py"
    specification = importlib.util.spec_from_file_location("question3_frozen_v3", source)
    engine = importlib.util.module_from_spec(specification)
    specification.loader.exec_module(engine)
    return engine


def current_evidence():
    paths = {name: ROOT / f"求解/问题3/结果/{name}.json" for name in ("公平对照", "灵敏度")}
    records = {name: json.loads(path.read_text(encoding="utf-8")) for name, path in paths.items()}
    hashes = {name: hashlib.sha256(path.read_bytes()).hexdigest() for name, path in paths.items()}
    if records["公平对照"]["估计器版本"] != "往返场_正交收缩响应_光程多起点_嵌套原块_v3":
        raise ValueError("当前证据不是正式升格3，禁止混用旧估计器")
    return records, hashes


def append_review_event(attempt, phenomenon, decision, evidence, category="科学尝试"):
    path = ROOT / "交接/实验记录.json"
    with (ROOT / "交接/.实验记录.lock").open("a", encoding="utf-8") as lock:
        fcntl.flock(lock, fcntl.LOCK_EX)
        records = json.loads(path.read_text(encoding="utf-8")) if path.exists() else []
        records.append({"类别": category, "问题": 3, "尝试": attempt, "现象": phenomenon,
                        "决定": decision, "依据": evidence})
        atomic_json(path, records)


def export_current_evidence():
    t0 = time.monotonic()
    records, hashes = current_evidence()
    output = RESULT / "回炉轮2候选/文图取数"
    cases = {(row["材料"], row["折号"]): row for row in records["公平对照"]["案例"]}
    rows, checks = [], []
    for index, original in enumerate(records["灵敏度"]["灵敏度_参数扰动"]):
        row = {key: value for key, value in original.items()
               if key not in ("训练源行", "重估完整参数", "未扰动控制参数")}
        row["src"] = f"求解/问题3/结果/灵敏度.json:灵敏度_参数扰动.{index}"
        row["原记录配对可比"] = row.pop("配对可比", None)
        if original["状态"] == "已重估":
            baseline = cases[(row["材料"], row["折号"])]["完整往返"]
            fitted = original["重估完整参数"]
            control = original["未扰动控制参数"]
            if not math.isclose(row["基准厚度_um"], baseline["厚度_um"], abs_tol=1e-10):
                raise ValueError("灵敏度基准与当前公平对照不一致")
            if "损耗" in row["情景"]:
                loss = baseline["有效往返损耗"]
                expected = {LOSS_LABELS[0]: loss, LOSS_LABELS[1]: 0.8 * loss,
                            LOSS_LABELS[2]: 1.2 * loss, LOSS_LABELS[3]: 0.1, LOSS_LABELS[4]: 0.2}[row["情景"]]
                actual = [row["请求固定参数值"], row["实际固定参数值"], fitted["参数向量"][4]]
                valid = all(math.isclose(value, expected, rel_tol=1e-10, abs_tol=1e-12) for value in actual)
                if loss == 0 and row["情景"] in LOSS_LABELS[1:3]:
                    valid &= actual == [0.0, 0.0, 0.0]
                if not valid or not math.isclose(control["参数向量"][4], loss, abs_tol=1e-12):
                    raise ValueError(f"损耗标签与实参不符：{row['src']}")
                checks.append({"材料": row["材料"], "折号": row["折号"], "情景": row["情景"],
                               "基准损耗": loss, "按标签应传值": expected, "请求_实际_重估值": actual,
                               "检查通过": True, "src": row["src"]})
            if control is not None:
                total = fitted["厚度_um"] - row["基准厚度_um"]
                advancement = control["厚度_um"] - row["基准厚度_um"]
                conditional = fitted["厚度_um"] - control["厚度_um"]
                if not (math.isclose(total, row["厚度变化_um"], abs_tol=1e-10)
                        and math.isclose(conditional, row["扣除控制厚度变化_um"], abs_tol=1e-10)
                        and math.isclose(total, advancement + conditional, abs_tol=1e-10)):
                    raise ValueError(f"厚度变化分解不一致：{row['src']}")
                row.update({"总厚度变化_um": total, "未扰动优化推进_um": advancement,
                            "扣除控制厚度变化_um": conditional,
                            "控制厚度_um": control["厚度_um"], "扰动后厚度_um": fitted["厚度_um"],
                            "成对优化器收敛": bool(row["停止状态"]["优化器收敛"]
                                                   and row["未扰动控制停止状态"]["优化器收敛"]),
                            "解释": "条件重估响应；两个未收敛过程的差不解释为固有灵敏度"})
        rows.append(row)
    blocks = [dict(row, src=f"求解/问题3/结果/公平对照.json:逐块评分.{index}")
              for index, row in enumerate(records["公平对照"]["逐块评分"])]
    payload = {"来源SHA256": hashes, "灵敏度表": rows, "损耗实参自动检查": checks, "逐材料逐块表": blocks,
               "说明": "仅从现存实际结果导表；src指向灵敏度.json，不重算、不引用旧793行或旧综合比较",
               "限制": "探索性复查；没有稳定高阶优势。未收敛候选不作固有敏感性解释。",
               "实际用时秒": elapsed(t0)}
    atomic_json(output / "当前文图数据.json", payload)
    fields = ("材料", "折号", "情景", "状态", "请求固定参数值", "实际固定参数值", "基准厚度_um",
              "扰动后厚度_um", "控制厚度_um", "总厚度变化_um", "未扰动优化推进_um",
              "扣除控制厚度变化_um", "成对优化器收敛", "src")
    write_csv(output / "灵敏度.csv", [{key: row.get(key) for key in fields} for row in rows], fields)
    append_review_event("按现存源结果生成问题3文图取数表", f"校验{len(checks)}条损耗实参，导出{len(rows)}个情景；没有重新拟合",
                        "仅导出候选文图数据，不修改论文或正式结果", "求解结果:当前文图数据:损耗实参自动检查", "流程事件")
    print(f"当前证据导表：{len(rows)}个情景、{len(blocks)}个逐块评分；未调用求解器", flush=True)
    return 0


def review_margins(engine, parameters, case):
    parameters = np.asarray(parameters, dtype=float)
    film = parameters[1] + parameters[2] * ((2000 / case["sigma"][[0, -1]]) ** 2 - 1)
    threshold = max(math.sin(math.radians(angle)) for angle in case["angles"]) + 2e-6
    propagation = np.concatenate((film - threshold, film + parameters[3] - threshold))
    margins = np.concatenate((parameters - engine.BOUNDS[:, 0], engine.BOUNDS[:, 1] - parameters, propagation))
    return {"物理盒下界余量": parameters - engine.BOUNDS[:, 0],
            "物理盒上界余量": engine.BOUNDS[:, 1] - parameters,
            "原光程盒下界余量": engine.encode(parameters), "原光程盒上界余量": 1 - engine.encode(parameters),
            "传播余量_膜两端_衬底两端": propagation,
            "传播约束端点_cm^-1": case["sigma"][[0, -1]],
            "最大约束违反量": max(0.0, -float(np.min(margins))),
            "优化传播阈值": threshold, "原场传播阈值": threshold - 1e-6}


def local_review(engine, parameters, case, model, config, deadline, fixed=None):
    fixed = {} if fixed is None else fixed
    probes = []
    for index in range(5):
        if index in fixed:
            continue
        for sign in (-1, 1):
            if time.monotonic() >= deadline:
                break
            unit = engine.encode(parameters)
            unit[index] += sign * 1e-6
            changed = engine.decode(unit)
            for fixed_index, value in fixed.items():
                changed[fixed_index] = value
            candidate = engine.evaluate(changed, case, model, config)
            probes.append({"原光程坐标维": index, "差分步长": sign * 1e-6,
                           "原目标返回值": 1e12 if candidate is None else candidate["loss"],
                           "触发不可行惩罚": candidate is None, "约束": review_margins(engine, changed, case)})
    return {"参数向量": parameters, "约束": review_margins(engine, parameters, case),
            "原坐标局部差分": probes, "说明": "这是实测局部数值诊断，不预断停止码4的唯一原因"}


def replay_v3_endpoint(engine, seed, case, model, config, deadline, maximum, fixed=None):
    fixed = {} if fixed is None else fixed
    saved_evaluate, saved_minimize = engine.evaluate, engine.minimize
    visited, terminal = [], {}

    def observe(parameters, *args, **kwargs):
        candidate = saved_evaluate(parameters, *args, **kwargs)
        visited.append({"参数向量": np.asarray(parameters).copy(),
                        "训练目标": None if candidate is None else candidate["loss"]})
        return candidate

    def minimize_record(objective, initial, **kwargs):
        def accepted(values):
            terminal["最后已接受迭代坐标"] = np.asarray(values).copy()

        def counted(values):
            terminal["最后请求坐标"] = np.asarray(values).copy()
            return objective(values)
        outcome = saved_minimize(counted, initial, callback=accepted, **kwargs)
        terminal.update({"优化器返回坐标": outcome.x, "停止码": int(outcome.status),
                         "停止信息": str(outcome.message), "迭代数": int(outcome.nit)})
        return outcome

    engine.evaluate, engine.minimize = observe, minimize_record
    try:
        best, trace = engine.refine(seed, case, model, config, deadline, maximum=maximum, fixed=fixed)
    finally:
        engine.evaluate, engine.minimize = saved_evaluate, saved_minimize
    initial = next((row["参数向量"] for row in visited if row["训练目标"] is not None), None)
    terminal["优化器返回终点"] = "优化器返回坐标" in terminal
    if initial is not None:
        for key in ("优化器返回坐标", "最后已接受迭代坐标", "最后请求坐标"):
            if key not in terminal:
                continue
            unit = engine.encode(initial)
            unit[[index for index in range(5) if index not in fixed]] = terminal[key]
            endpoint = engine.decode(unit)
            for index, value in fixed.items():
                endpoint[index] = value
            terminal[key + "诊断"] = local_review(engine, endpoint, case, model, config, deadline, fixed)
    return {"模型": model, "原共同初值": seed, "原算法停止记录": trace, "真实终点记录": terminal,
            "保留最优": None if best is None else engine.parameter_record(best), "全部访问": visited,
            "说明": "这是复演而非恢复历史终点；评估截断未返回OptimizeResult时，仅记录请求点及已接受迭代，不冒称停止终点"}


def paired_refine(engine, seed, case, model, config, deadline, maximum, fixed):
    lower, span = engine.BOUNDS[:, 0], np.diff(engine.BOUNDS, axis=1)[:, 0]
    free = [index for index in range(5) if index not in fixed]
    initial = np.clip(np.asarray(seed, dtype=float), lower, lower + span)
    for index, value in fixed.items():
        initial[index] = value
    anchor = np.array([20.0, 6.0, 0.0, 0.8, 0.2])
    for index, value in fixed.items():
        anchor[index] = value
    basis = np.zeros((4, 5))
    basis[:, 1] = 1
    basis[:, 2] = np.tile((2000 / case["sigma"][[0, -1]]) ** 2 - 1, 2)
    basis[2:, 3] = 1
    threshold = max(math.sin(math.radians(angle)) for angle in case["angles"]) + 2e-6
    jacobian = (basis * span)[..., free]
    calls, best, visited = 0, None, []
    maximum_iterations = {150: 30, 300: 60, 600: 120}[maximum]
    last = initial.copy()
    accepted_points = []

    def rebuild(values):
        parameters = initial.copy()
        parameters[free] = lower[free] + span[free] * values
        return parameters

    def objective(values):
        nonlocal calls, best, last
        if time.monotonic() >= deadline:
            raise engine.StopSearch("时间上限，未完成该成对单元")
        if calls >= maximum:
            raise engine.StopSearch("共同评估上限")
        calls += 1
        last = rebuild(values)
        margins = basis @ last - threshold
        evaluated = last.copy()
        if np.min(margins) < 0:
            anchor_margins = basis @ anchor - threshold
            violated = margins < 0
            fraction = np.max(-margins[violated] / (anchor_margins[violated] - margins[violated]))
            evaluated += min(1.0, fraction + 1e-10) * (anchor - evaluated)
        candidate = engine.evaluate(evaluated, case, model, config)
        if candidate is None:
            raise FloatingPointError("可行域延拓点计算失败；不返回1e12伪目标")
        extension = float(1000 * np.sum(((last - evaluated) / span) ** 2))
        visited.append({"请求参数": last.copy(), "可行参数": evaluated, "训练目标": candidate["loss"],
                        "求解器目标": candidate["loss"] + extension, "域外延拓惩罚": extension,
                        "请求点约束": review_margins(engine, last, case),
                        "约束": review_margins(engine, evaluated, case)})
        if best is None or candidate["loss"] < best["loss"]:
            best = candidate
        return candidate["loss"] + extension

    outcome, reason = None, "未开始"
    def accepted(values):
        accepted_points.append(rebuild(values))

    try:
        objective(((initial - lower) / span)[free])
        outcome = engine.minimize(objective, ((initial - lower) / span)[free], method="SLSQP",
            callback=accepted,
            bounds=[(0, 1)] * len(free), constraints={"type": "ineq",
                "fun": lambda values: basis @ rebuild(values) - threshold, "jac": lambda values: jacobian},
            options={"maxiter": maximum_iterations, "ftol": 1e-7, "eps": 1e-6})
        reason = "收敛" if outcome.success else f"优化器停止:{outcome.status}"
    except (engine.StopSearch, FloatingPointError, np.linalg.LinAlgError) as error:
        reason = str(error)
    endpoint = None if outcome is None else rebuild(outcome.x)
    violation = None if endpoint is None else review_margins(engine, endpoint, case)
    converged = bool(outcome is not None and outcome.success and violation["最大约束违反量"] <= 1e-7)
    retained_converged = bool(converged and best is not None
        and np.max(np.abs((best["parameters"] - endpoint) / span)) <= 1e-6
        and abs(best["loss"] - float(outcome.fun)) <= 1e-7)
    return {"共同初值": seed, "实际初值": initial, "固定参数索引": list(fixed), "固定参数值": list(fixed.values()),
            "共同评估上限": maximum, "共同迭代上限": maximum_iterations, "评估次数": calls,
            "迭代数": None if outcome is None else int(outcome.nit), "停止原因": reason,
            "停止码": None if outcome is None else int(outcome.status),
            "优化器收敛": converged, "保留点对应收敛终点": retained_converged,
            "优化器返回终点": outcome is not None,
            "终点参数": endpoint, "终点约束": violation,
            "终点求解器目标": None if outcome is None else float(outcome.fun),
            "最后已执行请求点": last, "最后已执行请求点约束": review_margins(engine, last, case),
            "已接受迭代点": accepted_points,
            "保留最优约束": None if best is None else review_margins(engine, best["parameters"], case),
            "保留最优": None if best is None else engine.parameter_record(best), "全部访问": visited,
            "终点口径": "仅OptimizeResult.x为返回终点；预算截断时终点为空，最后请求/已接受点/最优访问点分存",
            "修正": "原物理参数盒线性缩放；传播约束解析Jacobian；域外沿共同内点回缩加二次惩罚，可行域目标不变"}


def frozen_return_comparison(engine, case, record):
    groups = {"训练": case["train"], **{name: np.flatnonzero(np.isin(case["blocks"], case["spec"][name + "块"]))
                                     for name in ("校准", "测试")}}
    groups.update({f"测试块{block}": np.flatnonzero(case["blocks"] == block)
                   for block in case["spec"]["测试块"]})
    rows = []
    for anchor_model in engine.MODELS:
        anchor = record[anchor_model]
        degree, gain_degree, penalty = anchor["响应设置_背景阶_增益阶_惩罚"]
        background = np.vander(case["abscissa"], degree + 1, increasing=True)
        weights = (np.ones((len(background), 1)) if gain_degree == 0 else
                   np.column_stack(((1 - case["abscissa"]) / 2, (1 + case["abscissa"]) / 2)))
        predictions, physical = {}, {}
        for model in engine.MODELS:
            physical[model] = engine.fields(anchor["参数向量"], case, model == "完整往返")
            predictions[model] = np.asarray([response["训练响应均值_比例"] + response["训练响应尺度_比例"] *
                (background @ np.asarray(response["背景系数"]) +
                 ((spectrum[:, None] * weights - background @ np.asarray(response["物理投影系数"])) /
                  np.asarray(response["物理列尺度"])) @ np.asarray(response["标准化非负增益"]))
                for spectrum, response in zip(physical[model], anchor["各角响应"])])
        restored = engine.evaluate(anchor["参数向量"], case, anchor_model, (degree, gain_degree, penalty))
        if (restored is None or not np.allclose(restored["prediction"], predictions[anchor_model], rtol=0, atol=1e-10)
                or not math.isclose(restored["loss"], anchor["训练惩罚后目标"], rel_tol=0, abs_tol=1e-10)):
            raise ValueError("冻结响应或训练目标不能复现当前锚点")
        other_model = next(model for model in engine.MODELS if model != anchor_model)
        response_only = engine.evaluate(anchor["参数向量"], case, other_model, (degree, gain_degree, penalty))
        destination = record[other_model]
        if destination["响应设置_背景阶_增益阶_惩罚"] != list((degree, gain_degree, penalty)):
            raise ValueError("路径分解要求共同响应设置")
        optimized = engine.evaluate(destination["参数向量"], case, other_model, (degree, gain_degree, penalty))
        if (response_only is None or optimized is None
                or not math.isclose(optimized["loss"], destination["训练惩罚后目标"], rel_tol=0, abs_tol=1e-10)):
            raise ValueError("另一模型的同参数响应或现存最优候选无法重建")
        path_predictions = (predictions[anchor_model], predictions[other_model],
                            response_only["prediction"], optimized["prediction"])
        for name, members in groups.items():
            if not len(members):
                continue
            losses = np.asarray([engine.rmse(case, predicted, members) ** 2 for predicted in path_predictions])
            contributions = losses[:-1] - losses[1:]
            remainder = losses[0] - losses[-1] - np.sum(contributions, axis=0)
            if not np.allclose(remainder, 0, rtol=0, atol=1e-10):
                raise ArithmeticError("同一评价口径的收益分解不闭合")
            rows.append({"材料": case["material"], "折号": record["折号"], "锚点模型": anchor_model,
                "点集": name, "源行": case["rows"][members], "固定五参数": anchor["参数向量"],
                "固定响应": anchor["各角响应"], "响应设置": [degree, gain_degree, penalty],
                "训练尺度_比例": case["scales"], "重新拟合次数": 0,
                "纯物理场均方根差_反射率比例": np.sqrt(np.mean((physical["完整往返"][:, members] - physical["两束"][:, members]) ** 2, axis=1)),
                "冻结响应均方根误差_逐角": {model: engine.rmse(case, prediction, members)
                                            for model, prediction in predictions.items()},
                "收益路径分解": {"终点模型": other_model,
                    "四节点标准化MSE_逐角": losses,
                    "节点顺序": ["原锚点", "仅切返回阶数", "同物理参数仅重估训练响应", "另一模型现存重估候选"],
                    "仅返回阶数贡献_逐角": contributions[0], "响应重估贡献_逐角": contributions[1],
                    "物理参数重估贡献_逐角": contributions[2], "总收益_逐角": losses[0] - losses[-1],
                    "恒等式残差_逐角": remainder, "物理优化次数": 0,
                    "解释": "以正值为误差下降；双向锚点分别给出路径依赖的代数分解，不是因果效应或稳定高阶优势"},
                "说明": "仅切换返回阶数；五参数、响应投影、尺度和系数全部冻结，不声称这是重估厚度收益"})
    return rows


def same_start_stability(rows):
    groups = {}
    for row in rows:
        key = (row["材料"], row["折号"], row["起点序号"], row["模型"], row["情景"])
        groups.setdefault(key, {})[row["共同评估上限"]] = row
    comparisons = []
    for key, levels in groups.items():
        for earlier, later in ((150, 300), (300, 600)):
            older, newer = levels.get(earlier), levels.get(later)
            comparison = dict(zip(("材料", "折号", "起点序号", "模型", "情景"), key))
            comparison.update({"评估上限前后": [earlier, later], "局部数值稳定": False,
                               "状态": "未决：缺少级别或可行最优访问点"})
            if older and newer and older["保留最优"] and newer["保留最优"]:
                same_seed = np.allclose(older["实际初值"], newer["实际初值"], rtol=0, atol=1e-12)
                delta_d = newer["保留最优"]["厚度_um"] - older["保留最优"]["厚度_um"]
                delta_loss = newer["保留最优"]["训练惩罚后目标"] - older["保留最优"]["训练惩罚后目标"]
                stable = bool(same_seed and older["保留点对应收敛终点"] and newer["保留点对应收敛终点"]
                    and abs(delta_d) <= 0.01
                    and abs(delta_loss) <= 1e-6 * max(1, abs(newer["保留最优"]["训练惩罚后目标"])))
                comparison.update({"实际初值相同": bool(same_seed), "厚度差_um": delta_d,
                    "训练目标差": delta_loss, "局部数值稳定": stable,
                    "两级停止原因": [older["停止原因"], newer["停止原因"]],
                    "状态": "局部数值稳定，不保证全局最优" if stable else "未决：同预算或变化小不替代收敛"})
            comparisons.append(comparison)
    return comparisons


def review_carbide(engine, fair, cases, report, output, t0, flush):
    records = [record for record in fair["案例"] if record["材料"] == "碳化硅"]
    tasks = [(maximum, record, start_index, seed)
             for maximum in (150, 300, 600) for record in records
             for start_index, seed in enumerate(record["搜索状态"]["共同初值"])]
    rows, cumulative = report["碳化硅成对单元"], {}
    for record in records:
        for model in engine.MODELS:
            cumulative[f"折{record['折号']}/{model}"] = {"参数": record[model], "来源": "历史公平对照"}
    report["碳化硅含历史累计最优"] = cumulative
    report["碳化硅计划成员"] = [{"折号": record["折号"], "起点序号": start_index,
        "模型": model, "共同评估上限": maximum}
        for maximum, record, start_index, seed in tasks for model in engine.MODELS]
    report["碳化硅补查覆盖"] = {"计划成员数": 2 * len(tasks), "已保存成员数": 0,
        "实际执行完整": False, "保留点对应收敛终点数": 0, "说明": "尚无本轮碳化硅计算证据"}
    report["碳化硅同起点分级稳定性"] = []
    flush()
    for task_index, (maximum, record, start_index, seed) in enumerate(tasks):
        remaining = t0 + SOFT_SECONDS - time.monotonic()
        allowance = min(20.0, (remaining - 2.0) / (2 * (len(tasks) - task_index)))
        if allowance <= 0.05:
            break
        case = cases[(record["材料"], record["折号"])]
        config = tuple(record["完整往返"]["响应设置_背景阶_增益阶_惩罚"])
        if list(config) != record["两束"]["响应设置_背景阶_增益阶_惩罚"]:
            raise ValueError("碳化硅两模型的冻结响应配置不一致")
        for model in engine.MODELS:
            started = time.monotonic()
            trial = paired_refine(engine, seed, case, model, config,
                                  min(t0 + SOFT_SECONDS, started + allowance), maximum, {})
            trial.update({"材料": "碳化硅", "折号": record["折号"], "起点序号": start_index,
                "模型": model, "情景": "原五参数", "每成员时间上限秒": allowance,
                "本成员实际用时秒": time.monotonic() - started, "实际用时秒": elapsed(t0)})
            filename = f"碳化硅_折{record['折号']}_评估{maximum}_起点{start_index}_{model}.json"
            engine.atomic_json(output / filename, trial)
            rows.append({**{key: value for key, value in trial.items() if key != "全部访问"}, "文件": filename})
            retained = trial["保留最优"]
            key = f"折{record['折号']}/{model}"
            if retained and retained["训练惩罚后目标"] < cumulative[key]["参数"]["训练惩罚后目标"]:
                cumulative[key] = {"参数": retained, "来源": filename}
            report["碳化硅同起点分级稳定性"] = same_start_stability(rows)
            report["碳化硅补查覆盖"] = {"计划成员数": 2 * len(tasks), "已保存成员数": len(rows),
                "实际执行完整": len(rows) == 2 * len(tasks) and all("时间" not in item["停止原因"] for item in rows),
                "保留点对应收敛终点数": sum(item["保留点对应收敛终点"] for item in rows),
                "说明": "18个历史起点模型组合各作150/300/600补查；评估上限或时间截止不标为收敛"}
            flush()
            if retained:
                append_review_event(f"碳化硅第{record['折号']}组{model}同起点分级比较，评估上限{maximum}",
                    f"厚度{retained['厚度_um']:.12g}微米，目标{retained['训练惩罚后目标']:.12g}，实际评估{trial['评估次数']}次，收敛{trial['保留点对应收敛终点']}",
                    "按同一起点相邻评估额度比较，不以全部耗尽额度证明数值稳定", "求解结果:补查结果:碳化硅成对单元")
            else:
                append_review_event("碳化硅有限补查当前成员未形成新候选", trial["停止原因"],
                    "保存停止状态并保留既有答案，不宣称已收敛", "求解结果:补查结果:碳化硅成对单元", "流程事件")


def run_paired_budget_review():
    t0 = time.monotonic()
    engine = load_v3_engine()
    sources, hashes = current_evidence()
    output = RESULT / "回炉轮3候选/成对补查" / f"{time.strftime('%Y%m%d_%H%M%S')}_{os.getpid()}"
    fair, sensitivity = sources["公平对照"], sources["灵敏度"]["灵敏度_参数扰动"]
    report = {"问题": 3, "状态": "冻结既有真实数值，未替换正式结果", "源证据SHA256": hashes,
              "核心指标": fair["核心指标"], "既有核心指标": fair["核心指标"],
              "固定返回阶数对照": [], "历史复演": [], "成对单元": [], "碳化硅成对单元": [],
              "分级汇总": [], "同起点控制分解": [], "实际用时秒": 0.0,
              "指标含义": {"训练目标": "双角标准化训练MSE加原响应惩罚，不是留段误差",
                           "最大约束违反量": "物理盒和传播不等式负余量的最大值；零表示满足",
                           "局部数值稳定": "收敛且相邻评估级别的厚度和目标达到预设差值条件，不保证全局最优",
                           "硅折2已算可行候选包络_um": "原候选与本轮可行访问点极值，不是概率区间；无真厚度不能计算厚度覆盖"},
              "说明": "150/300/600为每次目标评估上限，迭代上限30/60/120；完整执行不等于收敛。只报告已算条件包络。",
              "结论状态": "未决：本候选不自动替换正式结果，独立复核与文图同步尚未发生"}

    def flush():
        report["实际用时秒"] = elapsed(t0)
        engine.atomic_json(output / "补查结果.json", report)

    flush()
    engine.atomic_json(output / "原证据冻结.json", sources)
    try:
        if engine.fingerprint(Path(engine.__file__)) != fair["实现SHA256"]:
            raise ValueError("正式升格3源码哈希不同，禁止将新实现冒充历史复演")
        inputs = engine.load_inputs()
        if inputs["hashes"] != fair["输入哈希"] or inputs["baseline_hash"] != fair["上游基准哈希"]:
            raise ValueError("源附件或上游基准已改变")
        report["原精修逐项核查"] = [{"材料": record["材料"], "折号": record["折号"],
            "模型": model, "起点序号": start_index, "共同初值": record["搜索状态"]["共同初值"][start_index],
            "历史停止记录": trace, "原文件是否保存终点": False,
            "src": f"公平对照:案例.{case_index}.搜索状态.优化记录.{model}.{start_index}"}
            for case_index, record in enumerate(fair["案例"])
            for model, traces in record["搜索状态"]["优化记录"].items()
            for start_index, trace in enumerate(traces)]
        carbide_traces = [row for row in report["原精修逐项核查"] if row["材料"] == "碳化硅"]
        report["碳化硅原18次核查"] = {"精修次数": len(carbide_traces),
            "评估上限停止数": sum(row["历史停止记录"]["停止原因"] == "共同评估上限" for row in carbide_traces),
            "优化器收敛数": sum(row["历史停止记录"].get("优化器收敛", False) for row in carbide_traces),
            "说明": "只读历史停止记录，不冒称重算或恢复历史终点"}
        flush()
        cases = {}
        report["冻结案例合同"] = []
        for record in fair["案例"]:
            fold = record["折号"]
            if fold and record["切分"] != inputs["folds"][fold - 1]:
                raise ValueError("校准或测试切分不符原合同")
            case = engine.build_case(record["材料"], inputs["tables"],
                inputs["full_indices"] if fold == 0 else inputs["selected"], inputs["edges"],
                None if fold == 0 else record["切分"])
            if (not np.array_equal(case["rows"][case["train"]], record["训练源行"])
                    or not np.allclose(case["scales"], record["训练尺度_比例"], rtol=0, atol=1e-12)):
                raise ValueError("源行或训练尺度不一致")
            cases[(record["材料"], fold)] = case
            report["冻结案例合同"].append({"材料": record["材料"], "折号": fold,
                "源行": case["rows"], "训练源行": case["rows"][case["train"]],
                "训练尺度_比例": case["scales"], "切分": record["切分"],
                "响应设置": {model: record[model]["响应设置_背景阶_增益阶_惩罚"] for model in engine.MODELS},
                "五参数边界": engine.BOUNDS, "目标": "冻结升格3双角标准化训练MSE及响应惩罚"})
            report["固定返回阶数对照"].extend(frozen_return_comparison(engine, case, record))
            flush()
        record = next(row for row in fair["案例"] if row["材料"] == "硅" and row["折号"] == 2)
        worst = next(row for row in sensitivity if row["材料"] == "硅" and row["折号"] == 2
                     and row["情景"] == "衬底对比加20%")
        case, config = cases[("硅", 2)], tuple(record["完整往返"]["响应设置_背景阶_增益阶_惩罚"])
        if list(config) != record["两束"]["响应设置_背景阶_增益阶_惩罚"]:
            raise ValueError("两模型响应设置不一致")
        contrast = record["完整往返"]["衬底折射率对比"]
        if not math.isclose(worst["请求固定参数值"], 1.2 * contrast, abs_tol=1e-12):
            raise ValueError("最坏扰动源实参不匹配")
        report["原停止记录"] = {"基准": record["搜索状态"], "扰动": worst["停止状态"],
                                 "控制": worst["未扰动控制停止状态"]}
        report["补查合同"] = {"五参数边界": engine.BOUNDS, "响应设置": config,
            "训练源行": case["rows"][case["train"]], "训练尺度_比例": case["scales"],
            "校准源行": case["rows"][np.isin(case["blocks"], case["spec"]["校准块"])],
            "测试源行": case["rows"][np.isin(case["blocks"], case["spec"]["测试块"])],
            "选型": "冻结已选响应，不用本轮外测试重选；数值补查为探索性，非独立验证",
            "响应自由度": 2 * (config[0] + config[1] + 2), "未固定时物理自由度": 5,
            "固定衬底对比时物理自由度": 4, "局部稳定阈值": {"厚度变化_um": 0.01, "目标相对变化": 1e-6},
            "总软截止秒": SOFT_SECONDS, "硅阶段截止秒": 720.0,
            "时间分配": "硅最多前720秒，碳化硅保留后360秒；每对相同时间上限，实际评估量分列，时间停止不标为收敛",
            "补查实现SHA256": engine.fingerprint(Path(__file__)),
            "局部探针": "双向原坐标1e-6探针单列，不计入优化评估额度；不能据域外探针断言实际差分失效"}
        cumulative = {f"{model}/原五参数": {"参数": record[model], "来源": "历史公平对照"} for model in engine.MODELS}
        cumulative.update({"完整往返/固定原对比控制": {"参数": worst["未扰动控制参数"], "来源": "历史灵敏度"},
                           "完整往返/衬底对比加20%": {"参数": worst["重估完整参数"], "来源": "历史灵敏度"}})
        for start_index, seed in enumerate(record["搜索状态"]["共同初值"]):
            for model in engine.MODELS:
                replay = replay_v3_endpoint(engine, seed, case, model, config, min(t0 + 1080, time.monotonic() + 15), 150)
                filename = f"原算法_{start_index}_{model}.json"
                engine.atomic_json(output / filename, replay)
                report["历史复演"].append({"文件": filename, "停止记录": replay["原算法停止记录"]})
                flush()
                append_review_event(f"按原光程坐标复演硅次折{model}第{start_index + 1}个共同起点",
                    f"目标评估{replay['原算法停止记录']['评估次数']}次，收敛标记{replay['原算法停止记录'].get('优化器收敛', False)}",
                    "分别保留最优访问点和真实终点的约束余量，不把执行完整当作收敛", "求解结果:补查结果:历史复演")
        for name, value in (("固定原对比控制", contrast), ("衬底对比加20%", 1.2 * contrast)):
            replay = replay_v3_endpoint(engine, record["完整往返"]["参数向量"], case, "完整往返", config,
                                       min(t0 + 1080, time.monotonic() + 15), 110, {3: value})
            filename = f"原算法_{name}.json"
            engine.atomic_json(output / filename, replay)
            report["历史复演"].append({"文件": filename, "停止记录": replay["原算法停止记录"]})
            flush()
            append_review_event(f"按原光程坐标复演硅次折{name}",
                f"固定衬底对比{value:.12g}，目标评估{replay['原算法停止记录']['评估次数']}次，收敛标记{replay['原算法停止记录'].get('优化器收敛', False)}",
                "保留原停止终点及约束诊断，与修正后的成对结果分列", "求解结果:补查结果:历史复演")
        seeds = list(record["搜索状态"]["共同初值"]) + [record[model]["参数向量"] for model in engine.MODELS]
        seeds += [worst[key]["参数向量"] for key in ("重估完整参数", "未扰动控制参数")]
        common = []
        for source in seeds:
            seed = np.clip(np.asarray(source), engine.BOUNDS[:, 0], engine.BOUNDS[:, 1])
            minimum_contrast = min(seed[3], contrast, 1.2 * contrast, 0.0)
            required = max(math.sin(math.radians(angle)) for angle in case["angles"]) + 2.2e-5
            required -= np.min(seed[2] * ((2000 / case["sigma"][[0, -1]]) ** 2 - 1)) + minimum_contrast
            seed[1] = max(seed[1], required)
            if seed[1] > engine.BOUNDS[1, 1]:
                raise ValueError("共同内点修复越过物理边界，不能静默裁剪")
            if not any(np.allclose(seed, prior, rtol=0, atol=1e-10) for prior in common):
                common.append(seed)
        report["共同起点修复"] = {"原池": seeds, "修正后共同池": common,
                                  "规则": "仅提升共同参考折射率到三个固定条件均传播的内点；原候选完整冻结，不删除"}
        conditions = (("原五参数", {}), ("固定原对比控制", {3: contrast}), ("衬底对比加20%", {3: 1.2 * contrast}))
        for stage_index, maximum in enumerate((150, 300, 600)):
            stage = []
            for start_index, seed in enumerate(common):
                remaining_members = ((3 - stage_index) * len(common) - start_index) * 6
                allowance = min(12.0, (t0 + 715.0 - time.monotonic()) / remaining_members)
                if allowance <= 0.05:
                    break
                group = []
                for name, fixed in conditions:
                    for model in engine.MODELS:
                        started = time.monotonic()
                        trial_deadline = min(t0 + 720, started + allowance)
                        trial = paired_refine(engine, seed, case, model, config,
                            trial_deadline, maximum, fixed)
                        trial.update({"材料": "硅", "折号": 2, "起点序号": start_index,
                                      "模型": model, "情景": name, "每成员时间上限秒": allowance,
                                      "本成员实际可用秒": max(0.0, trial_deadline - started),
                                      "本成员实际用时秒": time.monotonic() - started,
                                      "实际用时秒": elapsed(t0)})
                        filename = f"评估{maximum}_起点{start_index}_{name}_{model}.json"
                        engine.atomic_json(output / filename, trial)
                        summary = {key: value for key, value in trial.items() if key != "全部访问"}
                        summary["文件"] = filename
                        group.append(summary)
                        report["成对单元"].append(summary)
                        report["同起点分级稳定性"] = same_start_stability(report["成对单元"])
                        flush()
                        retained = trial["保留最优"]
                        if retained:
                            key = f"{model}/{name}"
                            if key not in cumulative or retained["训练惩罚后目标"] < cumulative[key]["参数"]["训练惩罚后目标"]:
                                cumulative[key] = {"参数": retained, "来源": filename}
                            append_review_event(f"同起点比较硅次折{model}的{name}，目标评估上限{maximum}",
                                f"厚度{retained['厚度_um']:.12g}微米，训练目标{retained['训练惩罚后目标']:.12g}，实际评估{trial['评估次数']}次，收敛标记{trial['保留点对应收敛终点']}",
                                "保留全部有限候选，按连续两级厚度与目标变化判断局部数值稳定性", "求解结果:补查结果:成对单元")
                        else:
                            append_review_event("成对数值补查当前成员未形成新候选", trial["停止原因"],
                                "保留既有真实答案及已算候选，不宣称数值稳定", "求解结果:补查结果:成对单元", "流程事件")
                stage.extend(group)
                for model in engine.MODELS:
                    paired = {row["情景"]: row for row in group if row["模型"] == model and row["保留最优"]}
                    if all(name in paired for name in ("固定原对比控制", "衬底对比加20%")):
                        control, changed = paired["固定原对比控制"], paired["衬底对比加20%"]
                        unperturbed = math.isclose(record[model]["衬底折射率对比"], contrast, rel_tol=0, abs_tol=1e-12)
                        report["同起点控制分解"].append({"模型": model, "共同评估上限": maximum,
                            "起点序号": start_index, "共同初值": seed,
                            "控制目标": control["保留最优"]["训练惩罚后目标"],
                            "扰动目标": changed["保留最优"]["训练惩罚后目标"],
                            "总厚度变化_um": changed["保留最优"]["厚度_um"] - record[model]["厚度_um"],
                            "控制相对旧候选总变化_um": control["保留最优"]["厚度_um"] - record[model]["厚度_um"],
                            "控制优化推进_um": control["保留最优"]["厚度_um"] - record[model]["厚度_um"] if unperturbed else None,
                            "控制是本模型未扰动条件": unperturbed, "原模型衬底对比": record[model]["衬底折射率对比"],
                            "共同固定控制对比": contrast,
                            "扣除控制厚度变化_um": changed["保留最优"]["厚度_um"] - control["保留最优"]["厚度_um"],
                            "共同自由参数初值相同": bool(np.allclose(np.asarray(control["实际初值"])[[0, 1, 2, 4]],
                                np.asarray(changed["实际初值"])[[0, 1, 2, 4]], rtol=0, atol=1e-12)),
                            "成对收敛": control["保留点对应收敛终点"] and changed["保留点对应收敛终点"],
                            "控制记录": control["文件"], "扰动记录": changed["文件"],
                            "解释": "完整场为本模型未扰动控制；两束仅为共同对比条件，不能把改固定值的效应称作纯优化推进。任一未收敛仍非固有灵敏度"})
                flush()
            winners = {}
            for name, fixed in conditions:
                for model in engine.MODELS:
                    candidates = [row for row in stage if row["情景"] == name and row["模型"] == model and row["保留最优"]]
                    if candidates:
                        winners[f"{model}/{name}"] = min(candidates, key=lambda row: row["保留最优"]["训练惩罚后目标"])
            changes = {}
            for model in engine.MODELS:
                control, changed = (winners.get(f"{model}/{name}") for name in ("固定原对比控制", "衬底对比加20%"))
                if control and changed:
                    unperturbed = math.isclose(record[model]["衬底折射率对比"], contrast, rel_tol=0, abs_tol=1e-12)
                    changes[model] = {"总厚度变化_um": changed["保留最优"]["厚度_um"] - record[model]["厚度_um"],
                        "控制相对旧候选总变化_um": control["保留最优"]["厚度_um"] - record[model]["厚度_um"],
                        "控制优化推进_um": control["保留最优"]["厚度_um"] - record[model]["厚度_um"] if unperturbed else None,
                        "控制是本模型未扰动条件": unperturbed,
                        "扣除控制厚度变化_um": changed["保留最优"]["厚度_um"] - control["保留最优"]["厚度_um"],
                        "成对收敛": control["保留点对应收敛终点"] and changed["保留点对应收敛终点"],
                        "解释": "共同初值池分别择优的条件响应，最优起点可能不同；同起点结果另列，未收敛差值不是固有灵敏度"}
            report["分级汇总"].append({"共同评估上限": maximum, "各条件最优": winners, "控制分解": changes,
                "含历史累计最优": dict(cumulative), "累计最优说明": "保留旧候选，不因新搜索更差而删除；累计最优不自动带有收敛资格",
                "成对执行完整": len(stage) == len(common) * 6 and all("时间" not in row["停止原因"] for row in stage)})
            flush()
        comparisons = []
        for earlier, later in zip(report["分级汇总"], report["分级汇总"][1:]):
            for name, newer in later["各条件最优"].items():
                older = earlier["各条件最优"].get(name)
                if older:
                    delta_d = newer["保留最优"]["厚度_um"] - older["保留最优"]["厚度_um"]
                    delta_objective = newer["保留最优"]["训练惩罚后目标"] - older["保留最优"]["训练惩罚后目标"]
                    same_seed = np.allclose(older["实际初值"], newer["实际初值"], rtol=0, atol=1e-12)
                    stable = (same_seed and earlier["成对执行完整"] and later["成对执行完整"] and older["保留点对应收敛终点"]
                              and newer["保留点对应收敛终点"] and abs(delta_d) <= 0.01
                              and abs(delta_objective) <= 1e-6 * max(1, abs(newer["保留最优"]["训练惩罚后目标"])))
                    comparisons.append({"条件": name, "评估上限前后": [earlier["共同评估上限"], later["共同评估上限"]],
                                        "最优起点相同": bool(same_seed), "厚度差_um": delta_d,
                                        "训练目标差": delta_objective, "局部数值稳定": bool(stable)})
        report["分级数值稳定性"] = comparisons
        decompositions = {(row["模型"], row["起点序号"], row["共同评估上限"]): row
                          for row in report["同起点控制分解"]}
        report["扣除控制变化的分级比较"] = []
        for model in engine.MODELS:
            for start_index in range(len(common)):
                for earlier, later in ((150, 300), (300, 600)):
                    older = decompositions.get((model, start_index, earlier))
                    newer = decompositions.get((model, start_index, later))
                    if older and newer:
                        report["扣除控制变化的分级比较"].append({"模型": model, "起点序号": start_index,
                            "评估上限前后": [earlier, later],
                            "条件响应变化_um": newer["扣除控制厚度变化_um"] - older["扣除控制厚度变化_um"],
                            "两级成对收敛": older["成对收敛"] and newer["成对收敛"],
                            "说明": "同一起点的控制差随预算变化；任一未收敛时不解释成物理敏感性"})
        values = [record[model]["厚度_um"] for model in engine.MODELS] + [worst["扰动后厚度_um"], worst["未扰动控制参数"]["厚度_um"]]
        values.extend(visit["可行参数"][0] for row in report["成对单元"]
                      for visit in json.loads((output / row["文件"]).read_text(encoding="utf-8"))["全部访问"])
        report["硅折2已算可行候选包络_um"] = [float(min(values)), float(max(values))]
        report["硅补查覆盖"] = {"计划成员数": len(common) * 18, "已保存成员数": len(report["成对单元"]),
            "保留点对应收敛终点数": sum(row["保留点对应收敛终点"] for row in report["成对单元"]),
            "实际执行完整": len(report["成对单元"]) == len(common) * 18
                and all("时间" not in row["停止原因"] for row in report["成对单元"])}
        flush()
        review_carbide(engine, fair, cases, report, output, t0, flush)
        report["状态"] = "有限预算补查停止；是否完整与是否收敛分列，未决项不自动消解"
    except Exception as error:
        report["状态"] = f"补查停止，保留既有答案与已算候选：{type(error).__name__}: {error}"
        flush()
        append_review_event("有限成对补查异常停止", report["状态"], "保留已保存候选与原正式答案，不自动宣告通过",
                            "求解结果:补查结果:状态", "流程事件")
        print(report["状态"], flush=True)
        return 1
    flush()
    print(json.dumps({"状态": report["状态"], "实际用时秒": report["实际用时秒"],
                      "输出目录": str(output), "结论状态": report["结论状态"],
                      "硅折2已算可行候选包络_um": report["硅折2已算可行候选包络_um"]}, ensure_ascii=False), flush=True)
    return 0


CLOSURE_VERSION = "往返场_实际端点_历史候选整合_原始起点精修_v3r3"
CLOSURE_INPUT = DATA / "问题3_公开复算输入/冻结条件.json"


def closure_backend(independent, deadline):
    engine = load_v3_engine()
    specification = importlib.util.spec_from_file_location("question3_independent_field", ROOT / "求解/问题3/复算.py")
    reference = importlib.util.module_from_spec(specification)
    specification.loader.exec_module(reference)
    reference.np, reference.minimize = np, engine.minimize
    reference.PARAM_LOW, reference.PARAM_HIGH = engine.BOUNDS[:, 0], engine.BOUNDS[:, 1]
    reference.TRANSFORM_LOW, reference.TRANSFORM_SCALE = engine.OPTICAL_LOW, engine.OPTICAL_SPAN
    reference.DEADLINE = deadline

    def problem(case, config):
        key = ("独立物理响应", tuple(config), case.get("weight", 0.5), case.get("gain_cap", engine.GAIN_CAP))
        if key not in case["cache"]:
            data = {"波数": case["sigma"], "反射率": case["values"], "窗口": case["sigma"][[0, -1]]}
            polarization = {0.0: "p", 0.5: "平均", 1.0: "s"}[case.get("weight", 0.5)]
            case["cache"][key] = reference.FieldProblem(data, case["train"], config,
                angles=case["angles"], polarization=polarization, gain_upper=case.get("gain_cap", engine.GAIN_CAP))
        return case["cache"][key]

    def evaluate(parameters, case, model, config):
        candidate = problem(case, config).evaluate(np.asarray(parameters), model, details=True)
        if candidate is None:
            return None
        return {"parameters": candidate["参数"], "loss": candidate["目标"], "prediction": candidate["预测"],
                "response": [{"标准化非负增益": gains} for gains in candidate["增益"]], "config": config, "model": model}

    original = engine.evaluate
    if independent:
        engine.evaluate = evaluate
    engine.closure_other = original if independent else evaluate
    engine.closure_problem = problem
    engine.closure_reference = reference
    return engine


def closure_repair(engine, original, case, fixed):
    parameters = np.clip(np.asarray(original, dtype=float), engine.BOUNDS[:, 0], engine.BOUNDS[:, 1])
    for index, value in fixed.items():
        if not engine.BOUNDS[index, 0] <= value <= engine.BOUNDS[index, 1]:
            raise ValueError("固定值在物理盒外，不夹紧请求条件")
        parameters[index] = value
    anchor = np.array([20.0, 6.0, 0.0, 0.8, 0.2])
    for index, value in fixed.items():
        anchor[index] = value
    margin = review_margins(engine, parameters, case)["传播余量_膜两端_衬底两端"]
    if np.min(margin) < 0:
        anchor_margin = review_margins(engine, anchor, case)["传播余量_膜两端_衬底两端"]
        if np.min(anchor_margin) <= 0:
            raise ValueError("固定条件下公开可行内点不存在")
        negative = margin < 0
        fraction = min(1.0, float(np.max(-margin[negative] / (anchor_margin[negative] - margin[negative]))) + 1e-10)
        parameters += fraction * (anchor - parameters)
    return parameters


def closure_search(engine, case, config, model, label, fixed, raw_seeds, historical, deadline, save):
    pool, coarse, trials, checks, rejected = [], [], [], [], []
    span = np.diff(engine.BOUNDS, axis=1)[:, 0]

    def consider(parameters, source, converged=False, retained=True, stopping=None):
        parameters = np.asarray(parameters, dtype=float)
        if any(abs(parameters[index] - value) > 1e-9 for index, value in fixed.items()):
            return None
        margins = review_margins(engine, parameters, case)
        if margins["最大约束违反量"] > 1e-7:
            rejected.append({"来源": source, "原因": "物理盒或优化传播约束违反超过1e-7", "约束": margins})
            return None
        candidate = engine.evaluate(parameters, case, model, config)
        if candidate is None:
            return None
        row = {"参数向量": parameters, "厚度_um": float(parameters[0]), "训练目标": candidate["loss"],
               "来源": source, "收敛": bool(converged), "范围资格": retained, "停止状态": stopping,
               "约束": margins, "固定参数索引": list(fixed), "固定参数值": list(fixed.values())}
        pool.append(row)
        return row

    for row in historical:
        if time.monotonic() >= deadline:
            break
        free_family = row["情景"] == "原五参数" and not row["固定参数索引"]
        if label == "原五参数" and not free_family:
            continue
        if not fixed and not free_family and row["情景"] != label:
            continue
        consider(row["参数向量"], row["来源"], row["收敛"], stopping=row["停止原因"])
    for index, seed in enumerate(raw_seeds):
        if time.monotonic() >= deadline and pool:
            break
        repaired = closure_repair(engine, seed, case, fixed)
        row = consider(repaired, f"原始谱投影起点/{index}", retained=False)
        if row is not None:
            coarse.append((row, index, repaired))
    if not pool:
        repaired = closure_repair(engine, [5.0, 3.0, 0.0, 0.8, 0.2], case, fixed)
        consider(repaired, "公开固定保底起点", retained=False)
    if not pool:
        raise ValueError("公开固定条件没有可行候选")
    ordered = sorted(pool, key=lambda row: (row["训练目标"], tuple(row["参数向量"])))
    for row in ordered[:3]:
        if time.monotonic() >= deadline:
            break
        counterpart = engine.closure_other(row["参数向量"], case, model, config)
        current = engine.evaluate(row["参数向量"], case, model, config)
        checks.append({"参数向量": row["参数向量"], "本侧约束": row["约束"],
            "对侧可行": counterpart is not None,
            "目标绝对差": None if counterpart is None else abs(current["loss"] - counterpart["loss"]),
            "预测最大绝对差_比例": None if counterpart is None else float(np.max(np.abs(current["prediction"] - counterpart["prediction"]))),
            "说明": "精修前的独立场与响应同参核对，不代替独立重新优化"})
    ranked = sorted(coarse, key=lambda item: (item[0]["训练目标"], item[1]))
    roots = ranked[:3]
    for branch in ordered:
        if len(roots) >= 6 or branch["训练目标"] > 1.1 * ordered[0]["训练目标"] + 1e-12 or not ranked:
            break
        nearest = min(ranked, key=lambda item: np.linalg.norm((item[2] - branch["参数向量"]) / span))
        if all(nearest[1] != selected[1] for selected in roots):
            roots.append(nearest)
    save({"阶段": "历史候选重评与同参核对", "同参核对": checks, "已接纳候选": pool,
          "拒绝候选": rejected, "原始起点数": len(raw_seeds), "精修根序号": [item[1] for item in roots]})
    for maximum in (150, 300):
        for _, seed_index, repaired in roots:
            if time.monotonic() >= deadline:
                break
            trial = paired_refine(engine, repaired, case, model, config, deadline, maximum, fixed)
            trial.update({"原始起点序号": seed_index, "未修复原始起点": raw_seeds[seed_index]})
            best = trial["保留最优"]
            if best is not None:
                consider(best["参数向量"], f"原始起点/{seed_index}/评估上限/{maximum}",
                    trial["保留点对应收敛终点"], stopping=trial["停止原因"])
            if trial["终点参数"] is not None:
                consider(trial["终点参数"], f"原始起点/{seed_index}/评估上限/{maximum}/返回终点",
                    trial["优化器收敛"], stopping=trial["停止原因"])
            trials.append(trial)
            save({"阶段": "原始起点有限精修", "同参核对": checks, "已接纳候选": pool,
                  "拒绝候选": rejected, "分级精修": trials})
            phenomenon = (f"保留厚度{best['厚度_um']:.15g}微米，训练目标{best['训练惩罚后目标']:.15g}，"
                          f"保留点对应收敛终点{trial['保留点对应收敛终点']}" if best is not None
                          else f"已作{trial['评估次数']}次目标评估，无可行保留点")
            append_review_event(f"在{case['material']}第{case['spec']['折号']}折、{label}条件下，"
                f"以训练谱投影第{seed_index + 1}个起点拟合{model}，目标评估上限{maximum}次",
                phenomenon, "保留可行较优点，并按本条件训练目标筛选近优分支；未收敛差不解释为固有灵敏度",
                "求解结果:分级精修/保留最优；分级精修/保留点对应收敛终点")
    winner = min(pool, key=lambda row: (row["训练目标"], tuple(row["参数向量"])))
    winner_record = engine.parameter_record(engine.evaluate(winner["参数向量"], case, model, config))
    eligible = [row for row in pool if row["范围资格"] or row is winner]
    ranges = {f"训练目标容差{tolerance}%": [row for row in eligible
        if row["训练目标"] <= winner["训练目标"] * (1 + tolerance / 100) + 1e-12] for tolerance in (0, 5, 10)}
    stability = []
    for seed_index in {trial["原始起点序号"] for trial in trials}:
        levels = [trial for trial in trials if trial["原始起点序号"] == seed_index]
        if len(levels) != 2 or any(trial["保留最优"] is None for trial in levels):
            continue
        earlier, later = (trial["保留最优"] for trial in levels)
        delta = later["厚度_um"] - earlier["厚度_um"]
        objective_delta = later["训练惩罚后目标"] - earlier["训练惩罚后目标"]
        stability.append({"原始起点序号": seed_index, "厚度变化_um": delta, "目标变化": objective_delta,
            "局部数值稳定": bool(all(trial["保留点对应收敛终点"] for trial in levels)
                and abs(delta) <= 0.01 and abs(objective_delta) <= 1e-6 * max(1.0, abs(later["训练惩罚后目标"])))})
    summary = {"材料": case["material"], "折号": case["spec"]["折号"], "模型": model, "情景": label,
        "固定参数索引": list(fixed), "固定参数值": list(fixed.values()), "最优": winner,
        "近优成员": ranges, "同起点分级稳定性": stability, "同参核对": checks,
        "厚度未识别": winner_record["厚度未识别"], "补充条件": winner_record["补充条件"],
        "分级端点": [{"原始起点序号": trial["原始起点序号"], "评估上限": trial["共同评估上限"],
                     "参数向量": trial["保留最优"]["参数向量"], "厚度_um": trial["保留最优"]["厚度_um"],
                     "训练目标": trial["保留最优"]["训练惩罚后目标"], "收敛": trial["保留点对应收敛终点"],
                     "约束": trial["保留最优约束"], "停止状态": trial["停止原因"]}
                    for trial in trials if trial["保留最优"] is not None],
        "精修完成": len(trials) == 2 * len(roots) and bool(roots)
                    and all("时间上限" not in trial["停止原因"] and trial["保留最优"] is not None for trial in trials),
        "范围规则": "仅同一目标和固定条件内0/5/10%近优的历史候选、精修保留点及终点；粗网格仅保底最优可入选",
        "停止状态": "达到本情景时间限额" if time.monotonic() >= deadline else "有限起点及两级评估完成"}
    save(dict(summary, 分级精修=trials, 已接纳候选=pool, 拒绝候选=rejected))
    return summary


def run_arbitration_closure(independent=False, compare_path=None):
    t0 = time.monotonic()
    backend = "独立公式" if independent else "建模公式"
    output = RESULT / "回炉轮3候选/仲裁闭环" / f"{time.strftime('%Y%m%d_%H%M%S')}_{os.getpid()}_{backend}"
    prior = json.loads((RESULT / "厚度结果.json").read_text(encoding="utf-8"))
    report = {"问题": 3, "估计器版本": CLOSURE_VERSION, "计算实现": backend, "核心指标": dict(prior["核心指标"]),
        "历史核心指标": dict(prior["核心指标"]), "数值来源": "新结果形成前保留原八项历史条件估计", "案例": [],
        "条件重估": [], "逐块评分": [], "预测区间": [], "同起点控制分解": [], "范围成员": [],
        "共同起点池条件响应": [], "原始初值生成完整": True, "原始初值": [],
        "八键独立复算": [], "实际用时秒": 0.0, "结论状态": "FAIL保留；未发生新数值验收",
        "置信说明": "低置信条件估计；自由厚度与固定光学条件分列。有限搜索、多解及旧测试探索性均不因收敛而消失。无厚度真值，厚度范围没有可验证的概率覆盖率。",
        "指标含义": {"训练目标": "仅训练数据的双角标准化惩罚后均方误差", "条件范围_um": "固定条件与近优分支联合极值，不是概率区间",
            "经验覆盖率": "反射率留出块覆盖，不是厚度覆盖；厚度无真值无法检验概率覆盖"}}

    def flush():
        report["实际用时秒"] = elapsed(t0)
        atomic_json(output / "厚度结果.json", load_v3_native(report))

    def alarm_handler(signum, frame):
        raise TimeoutError("1180秒硬截止，保留最近已写块")

    flush()
    previous_handler = signal.signal(signal.SIGALRM, alarm_handler)
    signal.setitimer(signal.ITIMER_REAL, HARD_SECONDS)
    try:
        package = json.loads(CLOSURE_INPUT.read_text(encoding="utf-8"))
        manifest = json.loads(CLOSURE_INPUT.with_name("输入清单.json").read_text(encoding="utf-8"))
        if hashlib.sha256(CLOSURE_INPUT.read_bytes()).hexdigest() != manifest["文件"][0]["SHA256"]:
            raise ValueError("公开条件输入哈希不符")
        engine = closure_backend(independent, t0 + SOFT_SECONDS)
        inputs = engine.load_inputs()
        if inputs["hashes"] != package["原始输入哈希"] or package["版本"] != CLOSURE_VERSION:
            raise ValueError("附件哈希或估计器版本不同")
        report["公开条件"] = {"文件": str(CLOSURE_INPUT.relative_to(ROOT)), "SHA256": manifest["文件"][0]["SHA256"],
            "闭环源码SHA256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
            "原场源码SHA256": engine.fingerprint(Path(engine.__file__)),
            "独立场源码SHA256": engine.fingerprint(ROOT / "求解/问题3/复算.py"),
            "传播约束端点_cm^-1": package["传播约束端点_cm^-1"], "训练截窗": "仅改变训练成员，不截短波数轴或重抽源行",
            "固定参数锚点": "冻结条件各案例原完整场五参数；不随本轮自由主线移动", "近优容差_百分数": [0, 5, 10],
            "优化约束容差": 1e-7, "独立比较相对容差": 0.01, "初值": "训练谱投影原始起点，历史头条及扰动上限不作初值",
            "范围版本变化": "同名八键仅在本版本隔离候选内使用；不覆盖v3正式结果", "保留证据": package["保留证据"]}
        cases, seeds, records = {}, {}, {}
        for record in package["案例"]:
            key = (record["材料"], record["折号"])
            fold = record["折号"]
            case = engine.build_case(record["材料"], inputs["tables"], inputs["full_indices"] if fold == 0 else inputs["selected"],
                inputs["edges"], None if fold == 0 else inputs["folds"][fold - 1])
            if not np.array_equal(case["rows"][case["train"]], record["训练源行"]):
                raise ValueError("训练源行与公开条件不同")
            if not np.array_equal(case["sigma"][[0, -1]], package["传播约束端点_cm^-1"]):
                raise ValueError("实际波数首末值变化，不自动换名义端点")
            cases[key], records[key] = case, record
            if independent:
                reference = engine.closure_reference
                reference.DEADLINE = min(t0 + 120, time.monotonic() + 15)
                try:
                    raw = reference.field_seeds(engine.closure_problem(case, tuple(record["响应设置"])))
                except reference.BudgetExpired:
                    raw = [np.array([5.0, 3.0, 0.0, 0.8, 0.2])]
                finally:
                    reference.DEADLINE = t0 + SOFT_SECONDS
            else:
                raw = engine.spectral_seeds(case, min(t0 + 120, time.monotonic() + 15))
            report["原始初值生成完整"] &= len(raw) == 61
            seeds[key] = raw
            report["原始初值"].append({"材料": key[0], "折号": fold, "五参数起点": raw})
            engine.atomic_json(output / f"原始初值_{key[0]}_{fold}.json", {"原始初值": raw,
                "生成方式": "既有v3训练谱投影384厚度网格，非头条参数回填", "实际用时秒": elapsed(t0)})
        stage_deadline = t0 + (360 if independent else 480)
        for position, (key, case) in enumerate(cases.items()):
            if time.monotonic() >= stage_deadline:
                break
            config = tuple(records[key]["响应设置"])
            best = {}
            for model in engine.MODELS:
                history = [row for row in package["历史候选"] if (row["材料"], row["折号"]) == key and row["模型"] == model]
                filename = f"自由_{key[0]}_{key[1]}_{model}.json"
                summary = closure_search(engine, case, config, model, "原五参数", {}, seeds[key], history,
                    min(stage_deadline, time.monotonic() + max(1.0, (stage_deadline - time.monotonic()) / max(1, 2 * (6 - position)))),
                    lambda value, filename=filename: engine.atomic_json(output / filename, dict(value, 实际用时秒=elapsed(t0))))
                report["案例"].append(summary)
                best[model] = engine.evaluate(summary["最优"]["参数向量"], case, model, config)
                flush()
            if key[1] and len(best) == 2:
                rows, intervals = engine.score_case(case, best, "完整往返")
                report["逐块评分"].extend(rows)
                report["预测区间"].extend(intervals)
                flush()
        condition_deadline = t0 + (660 if independent else 1020)
        jobs = []
        for key, case in cases.items():
            if key[0] != "硅":
                continue
            anchor = np.asarray(records[key]["固定参数锚点"])
            config = tuple(records[key]["响应设置"])
            for spec in engine.sensitivity_specs({"parameters": anchor}):
                label, index, requested, kind = spec
                if kind == "仅零损耗" and abs(anchor[4]) > 1e-8:
                    continue
                fixed = {} if index is None else {index: float(requested)}
                if index is not None and not engine.BOUNDS[index, 0] <= requested <= engine.BOUNDS[index, 1]:
                    report["条件重估"].append({"材料": key[0], "折号": key[1], "情景": label, "状态": "请求固定值出盒，不纳入", "请求值": requested})
                    continue
                changed, changed_config = dict(case), config
                if kind == "偏振":
                    changed["weight"] = requested
                elif kind == "角度":
                    changed["angles"] = tuple(angle + requested for angle in ANGLES)
                elif kind in ("下界", "上界"):
                    keep = case["sigma"][case["train"]] >= requested if kind == "下界" else case["sigma"][case["train"]] <= requested
                    changed["train"] = case["train"][keep]
                elif kind == "收缩":
                    changed_config = (*config[:2], config[2] * requested)
                elif kind == "增益上界":
                    changed["gain_cap"] = requested
                jobs.append((key, engine.prepare_case(changed), changed_config, label, fixed,
                             {} if index is None else {index: float(anchor[index])}))
        jobs.sort(key=lambda job: (0 if job[0] == ("硅", 2) and job[3] == "衬底对比加20%" else 1, job[0][1], job[3]))
        controls = {}
        for number, (key, case, config, label, fixed, control_fixed) in enumerate(jobs):
            if time.monotonic() >= condition_deadline:
                break
            history = [row for row in package["历史候选"] if (row["材料"], row["折号"]) == key and row["模型"] == "完整往返"]
            control_key = (key, tuple(control_fixed.items()))
            budget = max(0.5, (condition_deadline - time.monotonic()) / max(1, 2 * (len(jobs) - number)))
            if control_key not in controls:
                control = closure_search(engine, cases[key], tuple(records[key]["响应设置"]), "完整往返",
                    "固定原值控制" if control_fixed else "原五参数", control_fixed, seeds[key], history,
                    min(condition_deadline, time.monotonic() + budget),
                    lambda value, number=number: engine.atomic_json(output / f"控制_{number}.json", dict(value, 实际用时秒=elapsed(t0))))
                controls[control_key] = control
            changed = closure_search(engine, case, config, "完整往返", label, fixed, seeds[key], history,
                min(condition_deadline, time.monotonic() + budget),
                lambda value, number=number: engine.atomic_json(output / f"条件_{number}.json", dict(value, 实际用时秒=elapsed(t0))))
            report["条件重估"].append(changed)
            control = controls[control_key]
            control_members = control["近优成员"]["训练目标容差10%"]
            changed_members = changed["近优成员"]["训练目标容差10%"]
            differences = [perturbed["厚度_um"] - original["厚度_um"] for perturbed in changed_members for original in control_members]
            report["共同起点池条件响应"].append({"材料": key[0], "折号": key[1], "情景": label,
                "控制厚度_um": control["最优"]["厚度_um"], "扰动厚度_um": changed["最优"]["厚度_um"],
                "扣除控制厚度变化_um": changed["最优"]["厚度_um"] - control["最优"]["厚度_um"],
                "竞争分支差值范围_um": [min(differences), max(differences)],
                "成对收敛": control["最优"]["收敛"] and changed["最优"]["收敛"],
                "说明": "同一原始起点池，条件内分别择优；非同一胜出起点，不称固有物理灵敏度；差值范围无概率意义"})
            paired_controls = {(row["原始起点序号"], row["评估上限"]): row for row in control["分级端点"]}
            for endpoint in changed["分级端点"]:
                original = paired_controls.get((endpoint["原始起点序号"], endpoint["评估上限"]))
                if original is None:
                    continue
                anchor_depth = records[key]["固定参数锚点"][0]
                report["同起点控制分解"].append({"材料": key[0], "折号": key[1], "情景": label,
                    "原始起点序号": endpoint["原始起点序号"], "评估上限": endpoint["评估上限"],
                    "控制": original, "扰动": endpoint, "锚点厚度_um": anchor_depth,
                    "总厚度变化_um": endpoint["厚度_um"] - anchor_depth,
                    "未扰动优化推进_um": original["厚度_um"] - anchor_depth,
                    "扣除控制厚度变化_um": endpoint["厚度_um"] - original["厚度_um"],
                    "成对收敛": original["收敛"] and endpoint["收敛"],
                    "说明": "原始五参数起点相同；各固定条件的可行修复逐项公开，未收敛差不解释成固有敏感性"})
            flush()
        report["计划条件数"] = len(jobs)
        report["已算条件数"] = sum("最优" in row for row in report["条件重估"])
        report["固定原值控制"] = list(controls.values())
        for row in report["案例"] + report["条件重估"] + report["固定原值控制"]:
            if row.get("材料") == "硅" and "近优成员" in row:
                report["范围成员"].extend(dict(member, 材料="硅", 折号=row["折号"], 模型=row["模型"], 情景=row["情景"])
                    for member in row["近优成员"]["训练目标容差10%"])
        silicon = {row["折号"]: row for row in report["案例"] if row["材料"] == "硅" and row["模型"] == "完整往返"}
        if len(silicon) == 3 and report["范围成员"]:
            for fold, metric in ((0, "硅全量完整往返厚度_um"), (1, "硅折1完整往返条件厚度_um"), (2, "硅折2完整往返条件厚度_um")):
                report["核心指标"][metric] = silicon[fold]["最优"]["厚度_um"]
            depths = [member["厚度_um"] for member in report["范围成员"]]
            report["核心指标"].update({"硅已计算条件包络下限_um": min(depths), "硅已计算条件包络上限_um": max(depths)})
            report["数值来源"] = "硅五键为v3r3已算条件集合；碳化硅三键承接原基准，独立重算状态另列"
        report["分材料验证"] = engine.summarize_validation(report["逐块评分"]) if report["逐块评分"] else {}
        report["外测试口径"] = "全部原有角块不变；选型冻结后评分，仅探索性，不能重新用测试选分支"
        report["条件集合完整"] = report["已算条件数"] == len(jobs) and len(report["案例"]) == 12 and report["原始初值生成完整"]
        report["搜索规则完整"] = report["条件集合完整"] and all(row["精修完成"]
            for row in report["案例"] + report["条件重估"] + report["固定原值控制"] if "最优" in row)
        flush()
        if independent and time.monotonic() < t0 + 1050:
            reference = engine.closure_reference
            reference.checkpoint = lambda: engine.atomic_json(output / "碳化硅独立过程.json",
                {"核心指标": reference.VALUES, "证据": reference.EVIDENCE, "实际用时秒": elapsed(t0)})
            from openpyxl import load_workbook
            reference.load_workbook = load_workbook
            data_spec = package["碳化硅独立数据规范"]
            raw_records = reference.load_inputs(data_spec)
            reference.silicon_carbide_metrics(raw_records, data_spec, inputs["edges"])
            for metric, value in reference.VALUES.items():
                if metric.startswith("碳化硅") and value is not None:
                    report["核心指标"][metric] = value
            report["碳化硅独立完整"] = all(reference.VALUES.get(metric) is not None for metric in report["核心指标"] if metric.startswith("碳化硅"))
        if compare_path:
            compared = json.loads(Path(compare_path).read_text(encoding="utf-8"))
            same = compared.get("估计器版本") == CLOSURE_VERSION and compared.get("公开条件") == report["公开条件"]
            same &= compared.get("计算实现") == "建模公式" and independent
            original_pools = {(row["材料"], row["折号"]): np.asarray(row["五参数起点"])
                              for row in compared.get("原始初值", [])}
            seed_matches = []
            for row in report["原始初值"]:
                original = original_pools.get((row["材料"], row["折号"]))
                current = np.asarray(row["五参数起点"])
                seed_matches.append(bool(original is not None and original.shape == current.shape
                                         and np.allclose(original, current, rtol=0, atol=1e-10)))
            report["原始初值逐案例一致"] = seed_matches
            same &= len(seed_matches) == 6 and all(seed_matches)
            for metric, value in report["核心指标"].items():
                other = compared["核心指标"][metric]
                difference = abs(value - other) / max(abs(other), 1e-12)
                recomputed = report.get("碳化硅独立完整", False) if metric.startswith("碳化硅") else report["搜索规则完整"] and compared.get("搜索规则完整", False)
                report["八键独立复算"].append({"指标": metric, "本侧值_um": value, "对侧值_um": other,
                    "相对差": difference, "同公开条件": same, "真实完成独立计算": bool(independent and recomputed),
                    "百分之一内": bool(same and independent and recomputed and difference <= 0.01)})
        report["状态"] = "有限精修结束；保留竞争分支与验证失败，不自动替换正式八键或消解仲裁"
    except Exception as error:
        report["状态"] = f"部分完成：{type(error).__name__}: {error}"
    finally:
        signal.setitimer(signal.ITIMER_REAL, 0)
        signal.signal(signal.SIGALRM, previous_handler)
        flush()
        append_review_event("按公开实际端点与冻结条件执行有限候选整合", report["状态"],
            "保留真实八项答案和未决状态，不覆盖正式产物", "求解结果:仲裁闭环/厚度结果", "流程事件")
        print(json.dumps({"状态": report["状态"], "核心指标": report["核心指标"], "实际用时秒": report["实际用时秒"],
                          "候选目录": str(output)}, ensure_ascii=False), flush=True)
    return 0


def load_v3_native(value):
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, dict):
        return {key: load_v3_native(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [load_v3_native(item) for item in value]
    return value


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="问题3往返衰减场反演")
    modes = parser.add_mutually_exclusive_group()
    modes.add_argument("--loss-sensitivity-only", action="store_true", help="兼容别名：只从现存灵敏度取数，不重新拟合")
    modes.add_argument("--export-current-evidence", action="store_true", help="从实际灵敏度与逐块评分导表并核对损耗实参")
    modes.add_argument("--paired-budget-review", action="store_true", help="固定升格3口径，进行有限成对分级数值补查")
    modes.add_argument("--public-contract-only", action="store_true", help="只读原附件与公开规范重估硅31情景，不覆盖历史答案")
    modes.add_argument("--fair-comparison-only", action="store_true", help="正式升格3复算，候选隔离保存")
    modes.add_argument("--historical-only", action="store_true", help="显式重算历史不对称估计器，不用于高阶独立收益")
    modes.add_argument("--arbitration-closure", action="store_true", help="实际端点、公开原始初值和同条件候选整合；隔离产出v3r3")
    modes.add_argument("--closure-independent", action="store_true", help="独立场与响应实现重算v3r3及碳化硅原基准，不写红队报告")
    parser.add_argument("--compare-closure", help="独立重算后按1%%比较的v3r3厚度结果文件")
    arguments = parser.parse_args()
    if arguments.loss_sensitivity_only or arguments.export_current_evidence:
        raise SystemExit(export_current_evidence())
    if arguments.paired_budget_review:
        raise SystemExit(run_paired_budget_review())
    if arguments.historical_only:
        RESULT = RESULT / "回炉轮2候选/历史不对称"
        raise SystemExit(main())
    if arguments.public_contract_only:
        RESULT = RESULT / "回炉轮2候选/历史公开条件"
        raise SystemExit(run_contract_review(fair=False))
    if arguments.fair_comparison_only:
        official_engine = load_v3_engine()
        official_engine.RESULT = RESULT / "回炉轮2候选/升格3复算"
        raise SystemExit(official_engine.main())
    raise SystemExit(run_arbitration_closure(arguments.closure_independent, arguments.compare_closure))
