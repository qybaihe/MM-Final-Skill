"""问题3变体1：共享相位谐波回归；独立运行，所有数值产物写入本目录结果。"""

import hashlib
import json
import math
import os
import time
from pathlib import Path

for variable in ("OPENBLAS_NUM_THREADS", "OMP_NUM_THREADS", "MKL_NUM_THREADS"):
    os.environ[variable] = "1"

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import least_squares


ROOT = Path(__file__).resolve().parents[3]
RESULT = Path(__file__).resolve().parent / "结果"
ANGLES = (10.0, 15.0)
MATERIALS = {"硅": (3, 4), "碳化硅": (1, 2)}
ORDERS = (1, 3)
RIDGES = (0.0001, 0.01, 0.1)
BOUNDS = np.array([[0.5, 40.0], [1.2, 6.0], [-1.0, 1.0]])
SOFT_SECONDS = 1080.0
WORK_SECONDS = 1020.0
PARAMETER_NAMES = ("厚度_um", "参考折射率", "色散系数")
CORE_MAPPING = {
    "硅全量完整往返厚度_um": "硅全量条件厚度_um",
    "硅折1完整往返条件厚度_um": "硅折1条件厚度_um",
    "硅折2完整往返条件厚度_um": "硅折2条件厚度_um",
    "硅已计算条件包络下限_um": "硅已计算条件包络下限_um",
    "硅已计算条件包络上限_um": "硅已计算条件包络上限_um",
    "碳化硅正式基准厚度_um": "碳化硅正式基准厚度_um",
    "碳化硅正式条件范围下限_um": "碳化硅正式条件范围下限_um",
    "碳化硅正式条件范围上限_um": "碳化硅正式条件范围上限_um",
}


class BudgetStop(Exception):
    pass


def plain(value):
    if isinstance(value, np.ndarray):
        return plain(value.tolist())
    if isinstance(value, np.generic):
        return plain(value.item())
    if isinstance(value, dict):
        return {str(key): plain(item) for key, item in value.items()}
    if isinstance(value, (tuple, list)):
        return [plain(item) for item in value]
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("禁止写出非有限数值")
    return value


def save(name, value, started):
    RESULT.mkdir(parents=True, exist_ok=True)
    if isinstance(value, dict):
        value = {**value, "实际用时秒": time.monotonic() - started}
    destination = RESULT / name
    temporary = destination.with_suffix(".tmp")
    with temporary.open("w", encoding="utf-8") as stream:
        json.dump(plain(value), stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write("\n")
        stream.flush()
        os.fsync(stream.fileno())
    temporary.replace(destination)


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def experiment(attempt, observation, decision, evidence):
    path = ROOT / "交接/实验记录.json"
    entries = read_json(path) if path.exists() else []
    if not isinstance(entries, list):
        raise ValueError("实验记录必须为数组，禁止覆盖异常结构")
    entries.append({"类别": "科学尝试", "问题": 3, "尝试": attempt,
                    "现象": observation, "决定": decision, "依据": evidence})
    temporary = path.with_suffix(".升格1.tmp")
    with temporary.open("w", encoding="utf-8") as stream:
        json.dump(entries, stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write("\n")
        stream.flush()
        os.fsync(stream.fileno())
    temporary.replace(path)


def load_inputs():
    archive = read_json(ROOT / "交接/数据档案.json")
    upstream = read_json(ROOT / "求解/问题2/结果/基准交接.json")
    entries = {entry["文件名"]: entry for entry in archive["文件档案"]}
    tables, hashes = {}, []
    for number in range(1, 5):
        filename = f"附件{number}.xlsx"
        path = ROOT / "数据" / filename
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        if digest != entries[filename]["文件哈希"]:
            raise ValueError(f"{filename}与数据档案哈希不一致")
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            records = [(row_number, float(row[0]), float(row[1]) / 100.0)
                       for row_number, row in enumerate(
                           workbook["Sheet1"].iter_rows(min_row=2, values_only=True), 2)]
        finally:
            workbook.close()
        table = np.asarray(sorted(records, key=lambda record: record[1]))
        if table.shape != (7469, 3) or not np.all(np.isfinite(table)):
            raise ValueError(f"{filename}应有7469条有限数值观测")
        tables[number] = table
        hashes.append({"文件名": filename, "SHA256": digest})
    for number in (2, 3, 4):
        if not np.array_equal(tables[1][:, :2], tables[number][:, :2]):
            raise ValueError("四附件的原始行号或波数未逐行对齐")
    digest_map = {entry["文件名"]: entry["SHA256"] for entry in hashes}
    if any(digest_map[entry["文件名"]] != entry["SHA256"]
           for entry in upstream["输入哈希"]):
        raise ValueError("问题2正式基准不是当前附件的结果")
    if upstream["共同窗口_cm^-1"] != [1200, 3800] or upstream["抽样点数_每角度"] != 480:
        raise ValueError("共同口径发生变化，不能静默沿用")
    if [spec["折号"] for spec in upstream["折分"]] != [1, 2]:
        raise ValueError("必须保留原两折的编号和顺序")
    for spec in upstream["折分"]:
        groups = [set(spec[name]) for name in ("训练块", "校准块", "测试块")]
        if [len(group) for group in groups] != [8, 2, 2] or set.union(*groups) != set(range(1, 13)):
            raise ValueError("原折分应为互斥的8个训练块、2个校准块、2个测试块")
    keep = np.flatnonzero((tables[1][:, 1] >= 1200) & (tables[1][:, 1] <= 3800))
    if keep.size != 5392:
        raise ValueError("共同窗口应为每角5392点")
    sample = keep[np.arange(480) * (keep.size - 1) // 479]
    source_rows = tables[1][sample, 0].astype(int).tolist()
    if upstream["抽样源行"] != [source_rows, source_rows]:
        raise ValueError("抽样源行与正式基准不一致")
    datasets = {}
    for material, numbers in MATERIALS.items():
        datasets[material] = {}
        for label, indices in (("抽样", sample), ("全量", keep)):
            datasets[material][label] = {
                "sigma": tables[1][indices, 1],
                "values": np.stack([tables[number][indices, 2] for number in numbers]),
                "rows": tables[1][indices, 0].astype(int),
                "angles": ANGLES,
                "blocks": np.arange(480) // 40 + 1 if label == "抽样" else None,
            }
    specification = {
        "输入哈希": hashes, "共同窗口_cm^-1": [1200, 3800],
        "抽样点数_每角度": 480, "全量点数_每角度": 5392,
        "抽样源行": source_rows, "训练保护带_cm^-1": 20,
        "折分": upstream["折分"], "真实数据无合成输入": True,
        "搜索盒": dict(zip(PARAMETER_NAMES, BOUNDS.tolist())),
        "谐波阶数": ORDERS, "正则候选": RIDGES,
        "评分": "每角训练二次趋势残差IQR标准化；角度—测试块RMSE等权平均",
        "使用边界": "旧测试仅为探索性复查；未声称有新的独立外测试",
    }
    return datasets, upstream, specification


def training_indices(data, allowed_blocks, guard=20.0):
    if data["blocks"] is None:
        return np.arange(data["sigma"].size)
    allowed = set(allowed_blocks)
    selected = np.isin(data["blocks"], list(allowed))
    for boundary in range(1, 12):
        if (boundary in allowed) != (boundary + 1 in allowed):
            edge = np.mean(data["sigma"][[40 * boundary - 1, 40 * boundary]])
            selected &= np.abs(data["sigma"] - edge) >= guard
    return np.flatnonzero(selected)


def phase(parameters, sigma, angle):
    thickness, reference, dispersion = parameters
    refractive = reference + dispersion * ((2000.0 / sigma) ** 2 - 1.0)
    sine_squared = math.sin(math.radians(angle)) ** 2
    if np.any(refractive <= math.sqrt(sine_squared) + 0.01):
        return None
    return 4.0 * math.pi * thickness * 0.0001 * sigma * np.sqrt(refractive ** 2 - sine_squared)


def design(parameters, sigma, angle, order):
    argument = phase(parameters, sigma, angle)
    if argument is None:
        return None
    abscissa = (sigma - 2500.0) / 1300.0
    columns = [np.ones_like(sigma), abscissa, abscissa ** 2]
    for harmonic in range(1, order + 1):
        cosine, sine = np.cos(harmonic * argument), np.sin(harmonic * argument)
        columns.extend((cosine, sine, abscissa * cosine, abscissa * sine))
    return np.column_stack(columns)


def scales_and_baselines(data, indices):
    abscissa = (data["sigma"] - 2500.0) / 1300.0
    background = np.column_stack((np.ones_like(abscissa), abscissa, abscissa ** 2))
    scales, trends, means = [], [], []
    for observed in data["values"]:
        coefficient = np.linalg.lstsq(background[indices], observed[indices], rcond=None)[0]
        trend = background @ coefficient
        residual = observed[indices] - trend[indices]
        scales.append(max(float(np.quantile(residual, 0.75) - np.quantile(residual, 0.25)), 0.0001))
        trends.append(trend)
        means.append(np.full_like(observed, np.mean(observed[indices])))
    return np.asarray(scales), np.asarray(trends), np.asarray(means)


class PhaseFit:
    def __init__(self, data, indices, order, ridge, deadline, fixed=None, bounds=None):
        self.data, self.indices = data, np.asarray(indices)
        self.order, self.ridge, self.deadline = order, ridge, deadline
        self.fixed = fixed or {}
        self.bounds = BOUNDS.copy() if bounds is None else np.asarray(bounds)
        self.scales = scales_and_baselines(data, indices)[0]
        self.penalty = np.array([0.0, 0.0, 0.0] + [weight for harmonic in range(1, order + 1)
                               for weight in (harmonic ** 4, harmonic ** 4,
                                              4 * harmonic ** 4, 4 * harmonic ** 4)])
        self.pool, self.calls, self.stopped = [], 0, False

    def check(self):
        if time.monotonic() >= self.deadline:
            self.stopped = True
            raise BudgetStop

    def evaluate(self, parameters):
        self.calls += 1
        parameters = np.asarray(parameters, dtype=float).copy()
        for position, value in self.fixed.items():
            parameters[position] = value
        if np.any(parameters < self.bounds[:, 0]) or np.any(parameters > self.bounds[:, 1]):
            return None
        coefficients, residuals = [], []
        for angle_index, angle in enumerate(self.data["angles"]):
            if phase(parameters, self.data["sigma"], angle) is None:
                return None
            matrix = design(parameters, self.data["sigma"][self.indices], angle, self.order)
            observed = self.data["values"][angle_index, self.indices]
            normal = matrix.T @ matrix + len(observed) * self.ridge * np.diag(self.penalty)
            try:
                coefficient = np.linalg.solve(normal, matrix.T @ observed)
            except np.linalg.LinAlgError:
                augmented = np.vstack((matrix, np.diag(np.sqrt(len(observed) * self.ridge * self.penalty))))
                coefficient = np.linalg.lstsq(augmented, np.r_[observed, np.zeros(len(self.penalty))], rcond=None)[0]
            residuals.extend(((matrix @ coefficient - observed) / (self.scales[angle_index] * math.sqrt(len(observed))),
                              np.sqrt(self.ridge * self.penalty) * coefficient / self.scales[angle_index]))
            coefficients.append(coefficient)
        vector = np.concatenate(residuals) / math.sqrt(len(self.data["angles"]))
        candidate = {"parameters": parameters, "coefficients": np.asarray(coefficients),
                     "order": self.order, "ridge": self.ridge, "scales": self.scales,
                     "loss": float(vector @ vector), "residual": vector, "bounds": self.bounds}
        if not self.pool or candidate["loss"] < self.pool[0]["loss"] * 1.25:
            self.pool.append(candidate)
            self.pool.sort(key=lambda item: item["loss"])
            retained = []
            for item in self.pool:
                if all(np.linalg.norm((item["parameters"] - other["parameters"]) / [1.0, 0.5, 0.2]) > 0.08
                       for other in retained):
                    retained.append(item)
                if len(retained) == 10:
                    break
            self.pool = retained
        return candidate

    def run(self, seeds=(), scan=True, grid_count=129):
        first = np.clip(np.array([5.0, 3.0, 0.0]), self.bounds[:, 0], self.bounds[:, 1])
        self.evaluate(first)
        try:
            for seed in seeds:
                self.check()
                adjusted = np.clip(np.asarray(seed, dtype=float), self.bounds[:, 0], self.bounds[:, 1])
                if 1 in self.fixed:
                    adjusted[0] *= adjusted[1] / self.fixed[1]
                self.evaluate(np.clip(adjusted, self.bounds[:, 0], self.bounds[:, 1]))
            if scan:
                scenarios = [(reference, dispersion) for reference in (2.0, 3.0, 4.0)
                             for dispersion in (-0.5, 0.0, 0.5)] + [(1.2, 0.0), (6.0, 0.0)]
                optical_grid = np.linspace(0.6, 240.0, grid_count)
                traversal = sorted(range(grid_count), key=lambda index: int(f"{index:016b}"[::-1], 2))
                scan_end = time.monotonic() + max(0.0, self.deadline - time.monotonic()) * 0.6
                for grid_index in traversal:
                    if time.monotonic() >= scan_end:
                        break
                    for reference, dispersion in scenarios:
                        self.check()
                        self.evaluate([optical_grid[grid_index] / reference, reference, dispersion])
            starts = [candidate["parameters"].copy() for candidate in self.pool[:4]]
            if starts:
                for multiplier in (1.0 / 3.0, 0.5, 2.0, 3.0):
                    alias = starts[0].copy()
                    alias[0] *= multiplier
                    if self.bounds[0, 0] <= alias[0] <= self.bounds[0, 1]:
                        starts.append(alias)
            free = [position for position in range(3) if position not in self.fixed]
            size = len(self.data["angles"]) * (len(self.indices) + len(self.penalty))
            for initial in starts:
                self.check()

                def objective(values):
                    self.check()
                    parameters = initial.copy()
                    parameters[free] = values
                    candidate = self.evaluate(parameters)
                    return candidate["residual"] if candidate is not None else np.full(size, 1e6)

                least_squares(objective, initial[free], bounds=(self.bounds[free, 0], self.bounds[free, 1]),
                              x_scale=np.array([10.0, 3.0, 0.5])[free], max_nfev=65,
                              ftol=1e-7, xtol=1e-7, gtol=1e-7)
        except BudgetStop:
            self.stopped = True
        if not self.pool:
            raise ValueError("预先固定的物理可行起点未产生有限拟合")
        best = self.pool[0].copy()
        best.update({"calls": self.calls, "stopped": self.stopped,
                     "branches": [record(candidate, branches=False) for candidate in self.pool]})
        return best


def record(candidate, branches=True):
    result = {**dict(zip(PARAMETER_NAMES, candidate["parameters"])),
              "参考光学厚度_um": candidate["parameters"][0] * candidate["parameters"][1],
              "谐波阶数": candidate["order"], "正则强度": candidate["ridge"],
              "训练目标_无量纲": candidate["loss"], "分角线性系数": candidate["coefficients"],
              "训练尺度_反射率比例": candidate["scales"],
              "搜索边界命中": {name: bool(min(abs(value - lower), abs(value - upper)) <= 0.001 * (upper - lower))
                             for name, value, (lower, upper) in zip(PARAMETER_NAMES, candidate["parameters"], candidate.get("bounds", BOUNDS))},
              "系数顺序": "常数,x,x²；逐阶cos,sin,xcos,xsin；x=(波数-2500)/1300"}
    if branches:
        result.update({"目标计算次数": candidate.get("calls", 0),
                       "触发局部时间截止": candidate.get("stopped", False),
                       "保留分支": candidate.get("branches", [])})
    return result


def predict(candidate, data):
    return np.stack([design(candidate["parameters"], data["sigma"], angle, candidate["order"]) @ coefficient
                     for angle, coefficient in zip(data["angles"], candidate["coefficients"])])


def block_score(data, indices, predicted, scales):
    scores = []
    for block in sorted(set(data["blocks"][indices].tolist())):
        members = indices[data["blocks"][indices] == block]
        scores.extend(np.sqrt(np.mean((predicted[:, members] - data["values"][:, members]) ** 2, axis=1)) / scales)
    return float(np.mean(scores))


def deadline(started, seconds):
    return min(started + WORK_SECONDS, time.monotonic() + seconds)


def fit(data, indices, order, ridge, started, seconds, seeds=(), scan=True, fixed=None, bounds=None, grid_count=129):
    return PhaseFit(data, indices, order, ridge, deadline(started, seconds), fixed, bounds).run(seeds, scan, grid_count)


def select_inside(data, spec, started):
    available = sorted(spec["训练块"])
    validations = [available[len(available) // 3:len(available) // 3 + 2],
                   available[2 * len(available) // 3:2 * len(available) // 3 + 2]]
    observations = []
    for inner_number, held in enumerate(validations, 1):
        inner_blocks = [block for block in available if block not in held]
        indices = training_indices(data, inner_blocks)
        validation = np.flatnonzero(np.isin(data["blocks"], held))
        for order in ORDERS:
            if time.monotonic() >= started + WORK_SECONDS - 25.0:
                return None, observations
            anchor = fit(data, indices, order, 0.01, started, 6.0)
            for ridge in RIDGES:
                candidate = anchor if ridge == 0.01 else fit(data, indices, order, ridge, started, 1.5,
                                                           [anchor["parameters"]], scan=False)
                score = block_score(data, validation, predict(candidate, data), candidate["scales"])
                observations.append({"内折": inner_number, "训练块": inner_blocks, "验证块": held,
                                     "谐波阶数": order, "正则强度": ridge, "验证误差_无量纲": score,
                                     "训练点数_每角": len(indices), "条件厚度_um": candidate["parameters"][0]})
    averages = {(order, ridge): float(np.mean([row["验证误差_无量纲"] for row in observations
                 if row["谐波阶数"] == order and row["正则强度"] == ridge]))
                for order in ORDERS for ridge in RIDGES}
    ridge = min(RIDGES, key=lambda value: np.mean([averages[order, value] for order in ORDERS]))
    order = min(ORDERS, key=lambda value: (averages[value, ridge], value))
    return {"order": order, "ridge": ridge}, observations


def case_payload(case):
    return {"材料": case["material"], "折号": case["spec"]["折号"],
            "规格": case["spec"], "选中谐波阶数": case["selected"],
            "选择说明": case["selection_note"], "内层选型": case.get("inner", []),
            "初始模型": case.get("initial", {}),
            "训练源行": case["data"]["rows"][case["indices"]],
            "模型": {f"谐波{order}阶": record(candidate) for order, candidate in case["models"].items()}}


def core_metrics(cases, upstream, sensitivities):
    core = {"碳化硅正式基准厚度_um": upstream["全量参数"]["厚度_um"],
            "碳化硅正式条件范围下限_um": upstream["条件范围_um"][0],
            "碳化硅正式条件范围上限_um": upstream["条件范围_um"][1]}
    members = []
    for case in cases:
        if case["material"] != "硅":
            continue
        selected = case["models"][case["selected"]]
        label = "硅全量条件厚度_um" if case["spec"]["折号"] == 0 else f"硅折{case['spec']['折号']}条件厚度_um"
        core[label] = float(selected["parameters"][0])
        for order, candidate in case["models"].items():
            members.append({"来源": f"硅折{case['spec']['折号']}谐波{order}阶", "厚度_um": candidate["parameters"][0]})
            for branch in candidate.get("branches", []):
                if branch["训练目标_无量纲"] <= candidate["loss"] * 1.05 + 1e-12:
                    members.append({"来源": f"硅折{case['spec']['折号']}谐波{order}阶近优分支", "厚度_um": branch["厚度_um"]})
    members.extend({"来源": row["扰动"], "厚度_um": row["重估厚度_um"]}
                   for row in sensitivities if row["材料"] == "硅")
    if members:
        core["硅已计算条件包络下限_um"] = float(min(member["厚度_um"] for member in members))
        core["硅已计算条件包络上限_um"] = float(max(member["厚度_um"] for member in members))
    return core, members


def checkpoint(cases, upstream, sensitivities, started, stage):
    core, members = core_metrics(cases, upstream, sensitivities)
    aligned = {old_name: core[new_name] for old_name, new_name in CORE_MAPPING.items() if new_name in core}
    save("厚度结果.json", {"问题": 3, "方法族": "共享相位谐波回归与变量投影",
         "阶段": stage, "核心指标": core, "案例": [case_payload(case) for case in cases],
         "硅条件包络成员": members, "包络定义": "全量、两折、两阶模型、5%近优分支及实际完成扰动的极差；非概率置信区间"}, started)
    save("结果声明_问题3_升格1.json", {"问题": 3, "核心指标": core,
         "主线键对齐指标": aligned,
         "口径说明": {"主线键到本法键": CORE_MAPPING,
                      "对齐指标": "沿用旧键仅作机器比较索引，不表示重新使用完整往返场估计器；新估计器名称以核心指标和方法族为准。",
                      "改名理由": "前三键的完整往返是旧估计器名称；本法不冒充复场反演。材料、全量/折号、单位和范围用途不变。",
                      "数据": "四份真实XLSX；无合成观测；窗口1200至3800 cm^-1",
                      "全量": "5392点/角真实观测重新估计，480点只作初始化和共同留段评价",
                      "不确定性": "条件包络非概率区间；预测区间覆盖另见区间覆盖；真实厚度无标签，厚度覆盖不可识别",
                      "碳化硅": "正式基准始终原样保留，若触发相容性及厚度影响条件，附列本法修正候选"},
         "置信": "条件性答案；旧测试为探索性；不以复算一致性冒充真实测厚精度",
         "完成核心数": len(core), "预期核心数": 8, "阶段": stage}, started)
    save("执行状态.json", {"正常结束": False, "阶段": stage, "已形成核心数": len(core),
         "全局软截止秒": SOFT_SECONDS, "搜索停止秒": WORK_SECONDS}, started)


def evaluate_cases(cases, started):
    rows, predictions, summaries = [], [], []
    for case in cases:
        spec = case["spec"]
        if spec["折号"] == 0:
            continue
        data, indices = case["data"], case["indices"]
        scales, quadratic, mean = scales_and_baselines(data, indices)
        fitted = {f"谐波{order}阶": predict(candidate, data) for order, candidate in case["models"].items()}
        fitted.update({"所选模型": fitted[f"谐波{case['selected']}阶"], "训练均值": mean, "二次趋势": quadratic})
        calibration = np.flatnonzero(np.isin(data["blocks"], spec["校准块"]))
        for model_name, predicted in fitted.items():
            for angle_index, angle in enumerate(data["angles"]):
                errors = np.sort(np.abs(predicted[angle_index, calibration] - data["values"][angle_index, calibration]))
                rank = min(len(errors), math.ceil((len(errors) + 1) * 0.9))
                width = float(errors[rank - 1])
                for block in spec["测试块"]:
                    members = np.flatnonzero(data["blocks"] == block)
                    residual = predicted[angle_index, members] - data["values"][angle_index, members]
                    hits = int(np.count_nonzero(np.abs(residual) <= width))
                    rmse = float(np.sqrt(np.mean(residual ** 2)))
                    rows.append({"材料": case["material"], "折号": spec["折号"], "模型": model_name,
                                 "附件": MATERIALS[case["material"]][angle_index], "入射角_度": angle, "测试块": block,
                                 "标准化均方根误差": rmse / scales[angle_index], "均方根误差_百分点": rmse * 100,
                                 "训练尺度_比例": scales[angle_index], "校准点数": len(errors), "校准次序": rank,
                                 "名义覆盖率": 0.9, "经验覆盖率": hits / len(members), "覆盖点数": hits,
                                 "测试点数": len(members), "区间全宽_百分点": 200 * width,
                                 "距最近训练波数均值_cm^-1": float(np.mean(np.min(np.abs(
                                     data["sigma"][members, None] - data["sigma"][indices]), axis=1))),
                                 "预测超出零一比例点数": int(np.count_nonzero((predicted[angle_index, members] < 0)
                                                                                         | (predicted[angle_index, members] > 1)))})
                    predictions.append({"材料": case["material"], "折号": spec["折号"], "模型": model_name,
                                        "入射角_度": angle, "测试块": block, "源行": data["rows"][members],
                                        "波数_cm^-1": data["sigma"][members], "实测反射率_比例": data["values"][angle_index, members],
                                        "预测反射率_比例": predicted[angle_index, members],
                                        "区间下界_比例": predicted[angle_index, members] - width,
                                        "区间上界_比例": predicted[angle_index, members] + width})
        averages = {name: float(np.mean([row["标准化均方根误差"] for row in rows
                    if row["材料"] == case["material"] and row["折号"] == spec["折号"] and row["模型"] == name]))
                    for name in fitted}
        summaries.append({"材料": case["material"], "折号": spec["折号"], "逐模型误差": averages,
                          "所选相对均值下降_%": 100 * (1 - averages["所选模型"] / max(averages["训练均值"], 1e-15)),
                          "所选相对二次趋势下降_%": 100 * (1 - averages["所选模型"] / max(averages["二次趋势"], 1e-15)),
                          "三阶相对一阶下降_%": 100 * (1 - averages["谐波3阶"] / max(averages["谐波1阶"], 1e-15))})
    save("同口径对照.json", {"逐块评分": rows, "逐折汇总": summaries,
         "比较边界": "一阶是自由相位单谐波基线，不等同原Fresnel两束场；比较高阶的增量，不借用旧场模型误差",
         "评分单位": "每角训练二次趋势残差IQR标准化后，角度—测试块等权平均"}, started)
    save("留段预测.json", {"记录": predictions}, started)
    save("区间覆盖.json", {"构造": "仅用各角80个外校准绝对残差，第ceil((80+1)*0.9)=73个顺序统计量作半宽",
         "名义覆盖率": 0.9, "逐块覆盖": rows,
         "距离说明": "波数非时间；按测试块列离训练支持的距离和区间宽度，不虚构时间预测步长",
         "局限": "连续光谱残差不保证交换性，90%仅名义目标；经验覆盖如实报告，不许当有限样本覆盖保证",
         "厚度覆盖": {"可检验": False, "原因": "附件没有真实厚度和重复测量；本法厚度范围为条件极差，不设名义覆盖率"}}, started)
    return rows, summaries


def sensitivity(cases, started, upstream):
    rows = []
    for case in cases:
        if case["spec"]["折号"] != 0:
            continue
        candidate = case["models"][case["selected"]]
        base = candidate["parameters"]
        scenarios = [("未扰动重估控制", {})]
        scenarios += [(f"参考折射率固定为{reference:g}", {"fixed": {1: reference}})
                      for reference in (1.2, 2.0, 3.0, 4.0, 6.0)]
        scenarios += [("参考折射率减20%", {"fixed": {1: max(1.2, base[1] * 0.8)}}),
                      ("参考折射率加20%", {"fixed": {1: min(6.0, base[1] * 1.2)}}),
                      ("常折射率", {"fixed": {2: 0.0}}),
                      ("色散向下扰动", {"fixed": {2: max(-1.0, base[2] - max(abs(base[2]) * 0.2, 0.1))}}),
                      ("色散向上扰动", {"fixed": {2: min(1.0, base[2] + max(abs(base[2]) * 0.2, 0.1))}}),
                      ("双角各减0.5度", {"angles": tuple(angle - 0.5 for angle in ANGLES)}),
                      ("双角各加0.5度", {"angles": tuple(angle + 0.5 for angle in ANGLES)}),
                      ("窗口1300至3700", {"window": (1300.0, 3700.0)}),
                      ("窗口1200至3500", {"window": (1200.0, 3500.0)}),
                      ("正则减20%", {"ridge": candidate["ridge"] * 0.8}),
                      ("正则加20%", {"ridge": candidate["ridge"] * 1.2}),
                      ("二阶谐波", {"order": 2}), ("四阶谐波", {"order": 4}),
                      ("搜索盒向外扩20%", {"bounds": np.array([[0.4, 48.0], [0.96, 7.2], [-1.2, 1.2]])})]
        for label, configuration in scenarios:
            if time.monotonic() >= started + WORK_SECONDS - 15.0:
                save("灵敏度.json", {"灵敏度_参数扰动": rows, "提前收敛": True,
                     "未完成原因": "接近全局软截止；不把未计算情景记为已通过"}, started)
                return rows
            data = dict(case["data"])
            data["angles"] = configuration.get("angles", ANGLES)
            indices = case["indices"]
            if "window" in configuration:
                lower, upper = configuration["window"]
                indices = indices[(data["sigma"][indices] >= lower) & (data["sigma"][indices] <= upper)]
            altered = fit(data, indices, configuration.get("order", candidate["order"]),
                          configuration.get("ridge", candidate["ridge"]), started, 3.0,
                          [base], scan=False, fixed=configuration.get("fixed"), bounds=configuration.get("bounds"))
            rows.append({"材料": case["material"], "扰动": label, "配置": {
                         "固定光学参数": {PARAMETER_NAMES[position]: value for position, value in configuration.get("fixed", {}).items()},
                         "角度_度": data["angles"], "谐波阶数": altered["order"], "正则强度": altered["ridge"],
                         "窗口_cm^-1": configuration.get("window", [1200, 3800])},
                         "基准厚度_um": base[0], "重估厚度_um": altered["parameters"][0],
                         "厚度变化_%": 100 * (altered["parameters"][0] / base[0] - 1),
                         "重估参数": record(altered), "训练点数_每角": len(indices)})
            save("灵敏度.json", {"灵敏度_参数扰动": rows, "提前收敛": False,
                 "旧场参数适用性": "衬底折射率、偏振权重和有效往返损耗已不进入本估计器；用角度、谐波阶次和复系数递推检验替代，不移植旧54项数值"}, started)
            checkpoint(cases, upstream, rows, started, "逐项灵敏度")
            experiment(f"对{case['material']}施加{label}并重新估计厚度",
                       f"原厚度{base[0]:.9g}微米，改变条件后为{altered['parameters'][0]:.9g}微米",
                       "保留重估结果构成条件包络，不将条件变化解释为概率误差",
                       "求解结果:灵敏度_参数扰动")
    for case in cases:
        if case["spec"]["折号"] == 0:
            continue
        for guard in (16.0, 24.0):
            if time.monotonic() >= started + WORK_SECONDS - 10.0:
                break
            candidate = case["models"][case["selected"]]
            indices = training_indices(case["data"], case["spec"]["训练块"], guard)
            altered = fit(case["data"], indices, candidate["order"], candidate["ridge"], started, 2.0,
                          [candidate["parameters"]], scan=False)
            rows.append({"材料": case["material"], "扰动": f"折{case['spec']['折号']}保护带{guard:g}cm^-1",
                         "基准厚度_um": candidate["parameters"][0], "重估厚度_um": altered["parameters"][0],
                         "厚度变化_%": 100 * (altered["parameters"][0] / candidate["parameters"][0] - 1),
                         "重估参数": record(altered), "训练点数_每角": len(indices)})
            save("灵敏度.json", {"灵敏度_参数扰动": rows, "提前收敛": False}, started)
            experiment(f"将{case['material']}第{case['spec']['折号']}组保护带改为{guard:g}波数单位",
                       f"训练保留{len(indices)}点每角，重估厚度{altered['parameters'][0]:.9g}微米",
                       "保持原测试成员，保留保护带改变后的厚度变化",
                       "求解结果:灵敏度_参数扰动")
    return rows


def cross_angle(cases, started):
    records = []
    for case in cases:
        if case["spec"]["折号"] == 0:
            continue
        for source in (0, 1):
            if time.monotonic() >= started + WORK_SECONDS - 12.0:
                return records
            target = 1 - source
            data = case["data"]
            source_data = {**data, "values": data["values"][[source]], "angles": (ANGLES[source],)}
            target_data = {**data, "values": data["values"][[target]], "angles": (ANGLES[target],)}
            source_fit = fit(source_data, case["indices"], 3, 0.01, started, 4.0)
            predictor = PhaseFit(target_data, case["indices"], 3, 0.01, deadline(started, 1.0))
            adapted = predictor.evaluate(source_fit["parameters"])
            for name, candidate in (("严格移植响应", source_fit), ("仅目标角训练段适配响应", adapted)):
                predicted = predict(candidate, target_data)
                target_scales = scales_and_baselines(target_data, case["indices"])[0]
                for block in case["spec"]["测试块"]:
                    members = np.flatnonzero(data["blocks"] == block)
                    records.append({"材料": case["material"], "折号": case["spec"]["折号"],
                                    "来源角_度": ANGLES[source], "目标角_度": ANGLES[target], "测试块": block,
                                    "规范": name, "冻结厚度_um": source_fit["parameters"][0],
                                    "标准化均方根误差": block_score(target_data, members, predicted, target_scales)})
            save("交叉印证.json", {"留角度诊断": records,
                 "隔离": "源角单独按固定三阶和0.01正则估参，无双角起点；目标角只在自身外训练段拟线性系数；测试只计分"}, started)
            experiment(f"在{case['material']}第{case['spec']['折号']}组由{ANGLES[source]:g}度预测{ANGLES[target]:g}度",
                       f"源角冻结厚度{source_fit['parameters'][0]:.9g}微米，目标角分别计算严格移植与响应适配两种误差",
                       "将共享相位检验与目标角响应适配分开，不让目标测试反射率参与估计",
                       "求解结果:留角度诊断")
    return records


def algebra_and_judgement(cases, scores, sensitivities, cross_records, started):
    diagnostics, algebra = [], []
    for case in cases:
        if case["spec"]["折号"] != 0:
            continue
        data = case["data"]
        candidate = case["models"][3]
        linear = case["models"][1]
        unpenalized = PhaseFit(data, case["indices"], 3, 0.0, deadline(started, 1.0)).evaluate(candidate["parameters"])
        abscissa = (data["sigma"] - 2500.0) / 1300.0
        real_prediction = predict(candidate, data)
        for angle_index, angle in enumerate(ANGLES):
            coefficients = candidate["coefficients"][angle_index]
            argument = phase(candidate["parameters"], data["sigma"], angle)
            complex_prediction = coefficients[0] + coefficients[1] * abscissa + coefficients[2] * abscissa ** 2
            complex_terms = []
            for harmonic in range(1, 4):
                offset = 3 + 4 * (harmonic - 1)
                amplitudes = ((coefficients[offset] + coefficients[offset + 2] * abscissa)
                              - 1j * (coefficients[offset + 1] + coefficients[offset + 3] * abscissa))
                complex_terms.append(amplitudes)
                complex_prediction = complex_prediction + np.real(amplitudes * np.exp(1j * harmonic * argument))
            noise = float(np.sqrt(np.mean((real_prediction[angle_index] - data["values"][angle_index]) ** 2)))
            selected = np.linspace(0, len(abscissa) - 1, 9, dtype=int)
            first, second, third = np.asarray(complex_terms)[:, selected]
            recurrence = np.abs(third * first - second ** 2) / np.maximum(np.abs(third * first) + np.abs(second ** 2), 1e-12)
            attenuation = np.abs(second) / np.maximum(np.abs(first), 1e-12)
            high_amplitude = np.sqrt(np.mean((np.abs(second) ** 2 + np.abs(third) ** 2) / 2))
            strength = float(high_amplitude / max(noise, 1e-12))
            energy = np.array([np.mean(np.abs(term) ** 2) for term in complex_terms])
            fundamental_share = float(energy[0] / max(float(np.sum(energy)), 1e-12))
            raw_coefficients = unpenalized["coefficients"][angle_index]
            raw_terms = []
            for harmonic in range(1, 4):
                offset = 3 + 4 * (harmonic - 1)
                raw_terms.append((raw_coefficients[offset] + raw_coefficients[offset + 2] * abscissa[selected])
                                 - 1j * (raw_coefficients[offset + 1] + raw_coefficients[offset + 3] * abscissa[selected]))
            raw_first, raw_second, raw_third = raw_terms
            raw_recurrence = np.abs(raw_third * raw_first - raw_second ** 2) / np.maximum(
                np.abs(raw_third * raw_first) + np.abs(raw_second ** 2), 1e-12)
            raw_attenuation = np.abs(raw_second) / np.maximum(np.abs(raw_first), 1e-12)
            groups = [row for row in scores if row["材料"] == case["material"] and row["入射角_度"] == angle and row["模型"] == "谐波3阶"]
            paired = []
            for group in groups:
                contrasts = {name: next(row["标准化均方根误差"] for row in scores
                             if row["材料"] == case["material"] and row["入射角_度"] == angle
                             and row["折号"] == group["折号"] and row["测试块"] == group["测试块"] and row["模型"] == name)
                             for name in ("谐波1阶", "训练均值", "二次趋势")}
                paired.append({"折号": group["折号"], "测试块": group["测试块"],
                               "三阶误差": group["标准化均方根误差"], "对照误差": contrasts,
                               "三个对照均改善": all(group["标准化均方根误差"] < value for value in contrasts.values())})
            direction = bool(paired) and all(item["三个对照均改善"] for item in paired)
            thresholds = [{"递推阈值": threshold, "幅值噪声比阈值": minimum,
                           "相容性支持": bool(direction and np.max(recurrence) <= threshold
                                              and np.max(raw_recurrence) <= threshold and np.all(raw_attenuation < 1)
                                              and np.all(attenuation < 1) and strength >= minimum and fundamental_share >= 0.1)}
                          for threshold in (0.2, 0.25, 0.3) for minimum in (0.8, 1.0, 1.2)]
            supported = next(row["相容性支持"] for row in thresholds if row["递推阈值"] == 0.25 and row["幅值噪声比阈值"] == 1.0)
            difference = float(candidate["parameters"][0] - linear["parameters"][0])
            relative = 100 * abs(difference) / linear["parameters"][0]
            diagnostics.append({"材料": case["material"], "附件": MATERIALS[case["material"]][angle_index],
                                "入射角_度": angle, "递推检查波数_cm^-1": data["sigma"][selected],
                                "复谐波递推偏差": recurrence, "二阶与一阶幅值比": attenuation,
                                "冻结相位去正则递推偏差": raw_recurrence, "冻结相位去正则幅值比": raw_attenuation,
                                "一阶能量占比": fundamental_share, "最大能量谐波阶次": int(np.argmax(energy) + 1),
                                "基频歧义提示": "一阶能量低于10%，不能排除将倍频当基频" if fundamental_share < 0.1 else "一阶可见仍不证明厚度分支唯一",
                                "高阶幅值噪声比": strength, "逐块基线对照": paired, "阈值敏感性": thresholds,
                                "相容性支持": supported,
                                "多光束判断": "存在与多光束相容的高阶支持，非充分证明" if supported else "未获得稳定可观测高阶支持，不能判定不存在多光束",
                                "高阶加入厚度差_um": difference, "高阶加入厚度绝对变化_%": relative,
                                "厚度影响情景": [{"相对变化阈值_%": threshold, "超过阈值": relative > threshold} for threshold in (0.5, 1.0, 2.0)],
                                "替代解释": "色散失配、慢变响应和非几何高阶；偏振混合可破坏单复倍率近似"})
            reduced = design(candidate["parameters"], data["sigma"], angle, 1) @ coefficients[:7]
            algebra.append({"材料": case["material"], "入射角_度": angle,
                            "实三角与复系数实现最大差": float(np.max(np.abs(complex_prediction - real_prediction[angle_index]))),
                            "高阶置零退化最大差": float(np.max(np.abs(reduced - (design(candidate["parameters"], data["sigma"], angle, 3) @ np.r_[coefficients[:7], np.zeros(8)])))),
                            "厚度厘米微米换算相位最大差": float(np.max(np.abs(argument - 4 * math.pi * (candidate["parameters"][0] / 10000) * data["sigma"] * np.sqrt((candidate["parameters"][1] + candidate["parameters"][2] * ((2000 / data["sigma"]) ** 2 - 1)) ** 2 - math.sin(math.radians(angle)) ** 2))))})
    silicon_carbide = next(case for case in cases if case["material"] == "碳化硅" and case["spec"]["折号"] == 0)
    relevant = [row for row in diagnostics if row["材料"] == "碳化硅"]
    trigger = all(row["相容性支持"] and row["高阶加入厚度绝对变化_%"] > 1.0 for row in relevant)
    save("条件判定.json", {"逐附件判断": diagnostics, "代数核验": algebra,
         "碳化硅条件触发": trigger, "碳化硅三阶修正候选厚度_um": silicon_carbide["models"][3]["parameters"][0],
         "碳化硅结论": "附列本法高阶修正候选；保留问题2原基准作对照" if trigger else "保留问题2正式基准，不把谱形差异直接当成厚度精度改善",
         "未获观测信息": "相干长度、仪器分辨率、偏振态、实测色散及真实厚度均未提供",
         "必要条件": "至少两次非零返回；往返幅值衰减；时间相干与空间重叠；高阶贡献须高于噪声且谱仪能分辨",
         "厚度精度边界": "比较的是条件估计的改变，不是真实误差改善；1%仅预设情景阈值"}, started)
    checks = {"假设_共享相位": cross_records,
              "假设_色散与折射率": [row for row in sensitivities if "折射率" in row["扰动"] or "色散" in row["扰动"]],
              "假设_响应与阶次": [row for row in sensitivities if "正则" in row["扰动"] or "阶谐波" in row["扰动"]],
              "假设_波段与角度": [row for row in sensitivities if "窗口" in row["扰动"] or "双角" in row["扰动"]],
              "假设_保护带": [row for row in sensitivities if "保护带" in row["扰动"]]}
    save("假设检验.json", {"检验项": checks, "说明": "数组非空只表示完成对应重估，不代表通过；逐项数值由运行产生"}, started)
    obligations = [
        ("V1-A1", "同片两角共享厚度与色散，仅分开线性响应，简化跨角反演。", "数据档案:总览.独立晶圆数=2/每块晶圆入射角度=[10,15]", "双向源角独立拟合，比较严格移植与目标角训练响应适配。", "假设_共享相位"),
        ("V1-A2", "在1200至3800波数窗口以三参数厚度—参考折射率—经验色散替代未知连续物性函数。", "数据档案:未随附件提供的建模输入.折射率和消光系数随波长的参数或测量", "固定参考折射率情景及±20%，色散置零及双向扰动，重新估厚。", "假设_色散与折射率"),
        ("V1-A3", "各角二次背景和一次谐波幅值包络简化未知仪器响应；有限谐波截断不宣称无限场已知。", "数据档案:未随附件提供的建模输入.校准方法及参考光谱", "二阶、四阶以及正则±20%重新拟合，另核查去正则递推。", "假设_响应与阶次"),
        ("V1-A4", "主解采用标称角度和共同窗口，省略未知角度误差与低波数异常区响应。", "数据档案:总览.每块晶圆入射角度/总览.超过百分之百总点数=262", "双角±0.5度和两个缩窄窗口重新估计厚度。", "假设_波段与角度"),
        ("V1-A5", "以连续分块及20波数保护带降低相邻谱点泄漏，未假定残差已经独立。", "上游结果解读:问题2正式基准.训练保护带_cm^-1=20", "保护带改为16和24，在原两折训练成员内重新估计。", "假设_保护带"),
    ]
    ledger = []
    for number, assumption, source, duty, key in obligations:
        count = len(checks[key])
        result = f"实际完成{count}项比较，数值见对应结果；不预设稳健性通过" if count else "时间截止前未形成此项比较；该简化仅条件性生效，不作稳健性背书"
        ledger.append({"假设号": number, "假设": assumption, "依据": source, "灵敏度义务": duty,
                       "检验结果": f"{result}（求解结果:检验项.{key}）"})
    save("假设台账_问题3_升格1.json", ledger, started)


def main():
    started = time.monotonic()
    datasets, upstream, specification = load_inputs()
    save("输入规范.json", specification, started)
    cases, sensitivities = [], []
    for material in MATERIALS:
        full_spec = {"折号": 0, "训练块": list(range(1, 13)), "校准块": [], "测试块": []}
        for spec in [full_spec, *upstream["折分"]]:
            sampled = datasets[material]["抽样"]
            indices = training_indices(sampled, spec["训练块"])
            models = {order: fit(sampled, indices, order, 0.01, started, 5.0, grid_count=65) for order in ORDERS}
            data = sampled
            if spec["折号"] == 0:
                data = datasets[material]["全量"]
                indices = np.arange(len(data["sigma"]))
                models = {order: fit(data, indices, order, 0.01, started, 3.0,
                                    [candidate["parameters"]], scan=False) for order, candidate in models.items()}
            cases.append({"material": material, "spec": spec, "data": data, "indices": indices,
                          "models": models, "selected": 1,
                          "initial": {f"谐波{order}阶": record(candidate) for order, candidate in models.items()},
                          "selection_note": "预设一阶保守初始模型；尚未使用外测试或校准反射率选型"})
            checkpoint(cases, upstream, sensitivities, started, "先形成真实附件答案")
            experiment(f"以共享相位的一阶和三阶表示拟合{material}第{spec['折号']}组观测",
                       f"一阶厚度{models[1]['parameters'][0]:.9g}微米，三阶厚度{models[3]['parameters'][0]:.9g}微米",
                       "先保留两种阶数的条件值，再仅依据内部连续留段选择阶数和正则",
                       "求解结果:案例.初始模型")
    for case in cases:
        if time.monotonic() >= started + WORK_SECONDS - 240.0:
            break
        sampled = datasets[case["material"]]["抽样"]
        selection, inner = select_inside(sampled, case["spec"], started)
        case["inner"] = inner
        if selection is None:
            continue
        indices = training_indices(sampled, case["spec"]["训练块"])
        models = {order: fit(sampled, indices, order, selection["ridge"], started, 10.0,
                             grid_count=513) for order in ORDERS}
        if case["spec"]["折号"] == 0:
            models = {order: fit(case["data"], case["indices"], order, selection["ridge"], started, 5.0,
                                 [candidate["parameters"], *[np.array([branch[name] for name in PARAMETER_NAMES])
                                  for branch in candidate["branches"][:3]]], scan=False)
                      for order, candidate in models.items()}
        case["models"], case["selected"] = models, selection["order"]
        case["selection_note"] = "两内层连续留段共同选择同一正则；在该正则下选择阶数；外校准只估区间，外测试只计分"
        checkpoint(cases, upstream, sensitivities, started, "训练内升级")
        experiment(f"在{case['material']}第{case['spec']['折号']}组内部连续留段比较一阶和三阶相位表示",
                   f"完成{len(inner)}组误差比较，选择{case['selected']}阶，条件厚度{models[case['selected']]['parameters'][0]:.9g}微米",
                   "按训练内部误差选型，保留另一阶结果供高阶增量比较",
                   "求解结果:案例.内层选型；求解结果:案例.模型")
    scores, summaries = evaluate_cases(cases, started)
    for summary in summaries:
        experiment(f"将{summary['材料']}第{summary['折号']}组所选相位模型与两个朴素预测比较",
                   f"相对均值下降{summary['所选相对均值下降_%']:.9g}%，相对二次趋势下降{summary['所选相对二次趋势下降_%']:.9g}%",
                   "保留负改善与逐块方向，不以验证结果删去条件厚度",
                   "求解结果:逐折汇总；求解结果:逐块评分")
    sensitivities = sensitivity(cases, started, upstream)
    cross_records = cross_angle(cases, started)
    algebra_and_judgement(cases, scores, sensitivities, cross_records, started)
    checkpoint(cases, upstream, sensitivities, started, "计算结束")
    fold_better = all(row["所选相对均值下降_%"] > 0 and row["所选相对二次趋势下降_%"] > 0 for row in summaries)
    selected_rows = [row for row in scores if row["模型"] == "所选模型"]
    all_better = bool(selected_rows) and all(
        row["标准化均方根误差"] < other["标准化均方根误差"]
        for row in selected_rows for other in scores
        if other["模型"] in ("训练均值", "二次趋势")
        and all(row[key] == other[key] for key in ("材料", "折号", "入射角_度", "测试块")))
    save("执行状态.json", {"正常结束": True, "自设基线目标全部满足": all_better,
         "逐折平均基线目标满足": fold_better, "严格逐块基线目标满足": all_better,
         "实际完成灵敏度数": len(sensitivities), "留角度诊断数": len(cross_records),
         "全局软截止秒": SOFT_SECONDS, "搜索停止秒": WORK_SECONDS,
         "提前收敛": time.monotonic() - started >= WORK_SECONDS,
         "未达目标处置": "保留全部真实数值和反向项；此字段不是G2审查结论",
         "时间声明": "单次数值调用规模受限，每个网格及优化目标检查截止；预留60秒写结果；不是操作系统硬实时保证"}, started)
    core = core_metrics(cases, upstream, sensitivities)[0]
    print(json.dumps(plain({"问题": 3, "变体": 1, "核心指标": core,
                           "基线目标全部满足": all_better, "实际用时秒": time.monotonic() - started}), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
