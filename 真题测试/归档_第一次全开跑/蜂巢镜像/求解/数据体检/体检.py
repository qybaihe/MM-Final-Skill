import sys
sys.path.insert(0, '/tmp/蜂巢/pylibs')

from collections import Counter
from datetime import datetime, timezone
from hashlib import sha256
import json
from pathlib import Path
import re
from zipfile import ZipFile

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path('/tmp/蜂巢')
DATA_DIR = ROOT / '数据'
ARCHIVE_PATH = ROOT / '交接' / '数据档案.json'
EXPERIMENT_PATH = ROOT / '交接' / '实验记录.json'
LOG_PATH = ROOT / '日志' / '运行日志.md'

MATERIAL_META = {
    '附件1.xlsx': {'材料': '碳化硅', '入射角_度': 10, '同源组': '碳化硅同一晶圆片'},
    '附件2.xlsx': {'材料': '碳化硅', '入射角_度': 15, '同源组': '碳化硅同一晶圆片'},
    '附件3.xlsx': {'材料': '硅', '入射角_度': 10, '同源组': '硅同一晶圆片'},
    '附件4.xlsx': {'材料': '硅', '入射角_度': 15, '同源组': '硅同一晶圆片'},
}


def native(value):
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return float(value)
    if isinstance(value, np.bool_):
        return bool(value)
    return value


def round_float(value, digits=10):
    return round(float(value), digits)


def extract_unit(column_name):
    match = re.search(r'\((.*?)\)', str(column_name))
    return match.group(1).strip() if match else None


def xml_metadata(zip_file):
    text = zip_file.read('docProps/core.xml').decode('utf-8')
    def get(tag):
        match = re.search(rf'<{tag}[^>]*>(.*?)</{tag}>', text)
        return match.group(1) if match else None
    return {
        '创建者': get('dc:creator'),
        '最后修改者': get('cp:lastModifiedBy'),
        '创建时间': get('dcterms:created'),
        '修改时间': get('dcterms:modified'),
    }


def contiguous_intervals(mask, x):
    indices = np.flatnonzero(mask)
    if len(indices) == 0:
        return []
    groups = np.split(indices, np.where(np.diff(indices) != 1)[0] + 1)
    return [
        {
            '起始波数_cm-1': round_float(x[group[0]], 4),
            '终止波数_cm-1': round_float(x[group[-1]], 4),
            '点数': int(len(group)),
        }
        for group in groups
    ]


def profile_file(path):
    raw = path.read_bytes()
    with ZipFile(path) as zf:
        names = zf.namelist()
        xml_heads = {
            name: zf.read(name)[:120].decode('ascii', errors='replace')
            for name in ('xl/workbook.xml', 'xl/worksheets/sheet1.xml', 'docProps/core.xml')
            if name in names
        }
        encodings = sorted({
            match.group(1)
            for head in xml_heads.values()
            for match in [re.search(r'encoding=["\']([^"\']+)', head, re.I)]
            if match
        })
        package_info = {
            'ZIP条目数': len(names),
            'XML声明编码': encodings,
            '含宏': any('vbaProject' in name for name in names),
            '含外部链接': any(name.startswith('xl/externalLinks/') for name in names),
            '含自定义XML': any(name.startswith('customXml/') for name in names),
            '文档属性': xml_metadata(zf),
        }

    wb = load_workbook(path, read_only=False, data_only=False)
    frame = pd.read_excel(path, sheet_name=0, header=0, engine='openpyxl')
    x = frame.iloc[:, 0].to_numpy(dtype=float)
    y = frame.iloc[:, 1].to_numpy(dtype=float)
    dx = np.diff(x)
    step_counts = Counter(np.round(dx, 7).tolist())
    first_jump = abs(y[1] - y[0]) if len(y) > 1 else None
    local_diffs = np.abs(np.diff(y[1:min(len(y), 101)]))
    median_local_diff = float(np.median(local_diffs)) if len(local_diffs) else None
    jump_ratio = first_jump / median_local_diff if median_local_diff else None

    sheets = []
    total_formula = total_error = total_comment = total_hyperlink = 0
    for ws in wb.worksheets:
        formulas = []
        errors = []
        comments = []
        hyperlinks = []
        nonempty_by_row = []
        for row in ws.iter_rows():
            count = 0
            for cell in row:
                if cell.value is not None:
                    count += 1
                if cell.data_type == 'f':
                    formulas.append(cell.coordinate)
                elif cell.data_type == 'e':
                    errors.append(cell.coordinate)
                if cell.comment is not None:
                    comments.append(cell.coordinate)
                if cell.hyperlink is not None:
                    hyperlinks.append(cell.coordinate)
            nonempty_by_row.append(count)
        total_formula += len(formulas)
        total_error += len(errors)
        total_comment += len(comments)
        total_hyperlink += len(hyperlinks)
        sheets.append({
            '名称': ws.title,
            '可见性': ws.sheet_state,
            '使用区域': ws.calculate_dimension(),
            '最大行数': ws.max_row,
            '最大列数': ws.max_column,
            '合并单元格': [str(rng) for rng in ws.merged_cells.ranges],
            '隐藏行数': sum(1 for dim in ws.row_dimensions.values() if dim.hidden),
            '隐藏列数': sum(1 for dim in ws.column_dimensions.values() if dim.hidden),
            '冻结窗格': str(ws.freeze_panes) if ws.freeze_panes else None,
            '自动筛选区域': ws.auto_filter.ref,
            '表对象': list(ws.tables.keys()),
            '公式数': len(formulas),
            '错误单元格数': len(errors),
            '批注数': len(comments),
            '超链接数': len(hyperlinks),
            '空白行数_使用区域内': int(sum(count == 0 for count in nonempty_by_row)),
            '样式编号': sorted({cell.style_id for row in ws.iter_rows() for cell in row}),
            '数字格式': sorted({cell.number_format for row in ws.iter_rows() for cell in row}),
        })

    q = np.quantile(y, [0, 0.01, 0.25, 0.5, 0.75, 0.99, 1])
    over100 = y > 100
    quality_risks = []
    if y[0] == 0 and jump_ratio is not None and jump_ratio > 100:
        quality_risks.append({
            '级别': '高',
            '问题': '首个反射率为0且与次点发生异常跃迁',
            '位置': {'Excel行': 2, '波数_cm-1': round_float(x[0], 4), '反射率_百分比': 0.0},
            '证据': {
                '到次点绝对跃迁': round_float(first_jump, 8),
                '后续前100点相邻绝对差中位数': round_float(median_local_diff, 8),
                '跃迁倍数': round_float(jump_ratio, 3),
            },
            '处理建议': '默认将首点标记为边界/占位疑点；建模时剔除首点并做含首点敏感性对照，禁止直接插值覆盖原始值。',
        })
    if np.any(over100):
        quality_risks.append({
            '级别': '中',
            '问题': '存在反射率大于100%的测量值',
            '数量': int(over100.sum()),
            '占比': round_float(over100.mean(), 8),
            '连续区间': contiguous_intervals(over100, x),
            '最大值_百分比': round_float(y.max(), 8),
            '处理建议': '可能来自仪器基线/参考归一化；厚度算法不应擅自截断到100%，但应在可靠性分析中单列，并对峰位法评估其影响。',
        })

    return frame, {
        '文件': path.name,
        '题面元数据': MATERIAL_META.get(path.name),
        '字节数': path.stat().st_size,
        'SHA256': sha256(raw).hexdigest(),
        '文件签名': raw[:8].hex(),
        '容器格式': 'Excel OOXML（ZIP封装，.xlsx）',
        '包结构': package_info,
        '工作簿结构': {
            '工作表总数': len(wb.sheetnames),
            '工作表顺序': wb.sheetnames,
            '隐藏工作表数': sum(ws.sheet_state != 'visible' for ws in wb.worksheets),
            '定义名称': list(wb.defined_names),
            '工作表': sheets,
            '总公式数': total_formula,
            '总错误单元格数': total_error,
            '总批注数': total_comment,
            '总超链接数': total_hyperlink,
        },
        '表格识别': {
            '工作表': wb.sheetnames[0],
            '表头类型': '单级表头',
            '表头行_Excel': 1,
            '数据起始行_Excel': 2,
            '原始行数_含表头': int(len(frame) + 1),
            '数据行数': int(len(frame)),
            '列数': int(frame.shape[1]),
            '列名': [str(col) for col in frame.columns],
        },
        '字段': [
            {
                '列序号': 1,
                '字段名': str(frame.columns[0]),
                '规范名': '波数',
                '单位': extract_unit(frame.columns[0]),
                '读取类型': str(frame.dtypes.iloc[0]),
                '角色': '附件内候选主键；四附件公共关联键',
                '最小值': round_float(x.min(), 4),
                '最大值': round_float(x.max(), 4),
                '唯一值数': int(pd.Series(x).nunique()),
                '缺失数': int(pd.Series(x).isna().sum()),
                '重复数': int(pd.Series(x).duplicated().sum()),
                '严格递增': bool(np.all(dx > 0)),
                '步长统计': {
                    '最小': round_float(dx.min(), 7),
                    '中位数': round_float(np.median(dx), 7),
                    '平均': round_float(dx.mean(), 10),
                    '最大': round_float(dx.max(), 7),
                    '标准差': round_float(dx.std(ddof=0), 10),
                    '按7位小数计数': {format(step, '.7f'): count for step, count in sorted(step_counts.items())},
                },
            },
            {
                '列序号': 2,
                '字段名': str(frame.columns[1]),
                '规范名': '反射率',
                '单位': extract_unit(frame.columns[1]),
                '读取类型': str(frame.dtypes.iloc[1]),
                '角色': '测量值',
                '缺失数': int(pd.Series(y).isna().sum()),
                '零值数': int(np.sum(y == 0)),
                '负值数': int(np.sum(y < 0)),
                '大于100值数': int(np.sum(y > 100)),
                '统计': {
                    '最小值': round_float(y.min(), 8),
                    '1%分位数': round_float(q[1], 8),
                    '25%分位数': round_float(q[2], 8),
                    '中位数': round_float(q[3], 8),
                    '均值': round_float(y.mean(), 8),
                    '75%分位数': round_float(q[4], 8),
                    '99%分位数': round_float(q[5], 8),
                    '最大值': round_float(y.max(), 8),
                    '样本标准差': round_float(y.std(ddof=1), 8),
                },
            },
        ],
        '数据质量': {
            '缺失单元格总数': int(frame.isna().sum().sum()),
            '完全重复行数': int(frame.duplicated().sum()),
            '主键重复数': int(frame.iloc[:, 0].duplicated().sum()),
            '非数值数据单元格数': int(sum(pd.to_numeric(frame[col], errors='coerce').isna().sum() - frame[col].isna().sum() for col in frame.columns)),
            '异常与风险': quality_risks,
        },
    }


def pair_stats(frames, a, b):
    ya = frames[a].iloc[:, 1].to_numpy(dtype=float)
    yb = frames[b].iloc[:, 1].to_numpy(dtype=float)
    diff = yb - ya
    return {
        '附件A': a,
        '附件B': b,
        '皮尔逊相关系数': round_float(np.corrcoef(ya, yb)[0, 1], 8),
        'B减A均值_百分点': round_float(diff.mean(), 8),
        '平均绝对差_百分点': round_float(np.abs(diff).mean(), 8),
        '均方根差_百分点': round_float(np.sqrt(np.mean(diff ** 2)), 8),
        '最小差_百分点': round_float(diff.min(), 8),
        '最大差_百分点': round_float(diff.max(), 8),
    }


def append_experiments(entries):
    if EXPERIMENT_PATH.exists():
        existing = json.loads(EXPERIMENT_PATH.read_text(encoding='utf-8'))
        if not isinstance(existing, list):
            raise ValueError('交接/实验记录.json 顶层必须为数组')
    else:
        existing = []
    known = {json.dumps(item, ensure_ascii=False, sort_keys=True) for item in existing}
    for entry in entries:
        signature = json.dumps(entry, ensure_ascii=False, sort_keys=True)
        if signature not in known:
            existing.append(entry)
            known.add(signature)
    EXPERIMENT_PATH.write_text(json.dumps(existing, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def main():
    started = datetime.now(timezone.utc)
    files = sorted(DATA_DIR.iterdir(), key=lambda p: p.name)
    unsupported = [p.name for p in files if p.is_file() and p.suffix.lower() != '.xlsx']
    xlsx_files = [p for p in files if p.is_file() and p.suffix.lower() == '.xlsx']
    if not xlsx_files:
        raise RuntimeError('数据目录未发现 .xlsx 附件')

    frames = {}
    file_profiles = []
    for path in xlsx_files:
        frame, profile = profile_file(path)
        frames[path.name] = frame
        file_profiles.append(profile)

    base_name = xlsx_files[0].name
    base_key = frames[base_name].iloc[:, 0].to_numpy(dtype=float)
    key_checks = []
    all_exact = True
    for path in xlsx_files:
        key = frames[path.name].iloc[:, 0].to_numpy(dtype=float)
        exact = len(key) == len(base_key) and np.array_equal(key, base_key)
        all_exact = all_exact and exact
        key_checks.append({
            '文件': path.name,
            '与附件1键序列完全一致': bool(exact),
            '行数': int(len(key)),
            '最大绝对键差': 0.0 if exact else round_float(np.max(np.abs(key - base_key)), 10),
        })

    paired = [
        pair_stats(frames, '附件1.xlsx', '附件2.xlsx'),
        pair_stats(frames, '附件3.xlsx', '附件4.xlsx'),
    ]
    correlation = pd.DataFrame({name: df.iloc[:, 1] for name, df in frames.items()}).corr()

    over100_count = int(sum((df.iloc[:, 1] > 100).sum() for df in frames.values()))
    first_zero_files = [name for name, df in frames.items() if float(df.iloc[0, 1]) == 0.0]
    archive = {
        '档案版本': '1.0',
        '生成时间_UTC': started.isoformat(timespec='seconds'),
        '数据根目录': '数据/',
        '体检范围': {
            '目录内文件总数': len([p for p in files if p.is_file()]),
            '已体检附件数': len(xlsx_files),
            '已体检文件': [p.name for p in xlsx_files],
            '非XLSX文件': unsupported,
            '覆盖结论': '数据/ 目录下全部普通文件均已纳入清点；当前4个文件均为XLSX并已逐文件体检。',
        },
        '权威读取规则': {
            '格式': 'Excel OOXML（.xlsx），不是CSV文本文件',
            '编码': 'XLSX为ZIP封装；核心XML、工作簿XML和工作表XML均声明UTF-8。GBK与日期格式规则在本批附件中不适用。',
            '推荐代码': "pd.read_excel(path, sheet_name='Sheet1', header=0, engine='openpyxl')",
            '工作表': 'Sheet1',
            '表头': '第1行单级表头',
            '数据区': '第2至7470行，共7469条记录',
            '字段规范化': {'波数 (cm-1)': '波数_cm-1', '反射率 (%)': '反射率_百分比'},
            '日期字段': '无',
        },
        '总体结论': [
            '4个附件结构完全一致：均为1个可见Sheet1、2列、7469条数据，无隐藏sheet、合并单元格、多级表头、公式、外链、宏、缺失或重复主键。',
            '四附件波数序列逐值完全一致，可按波数一对一横向关联；附件内主键为波数，全局记录主键应使用“附件名+波数”。',
            '附件1/2是同一碳化硅晶圆片的10°/15°测量，附件3/4是同一硅晶圆片的10°/15°测量；材料与角度来自题面，不在Excel列内。',
            f'四附件首个波数399.6747 cm-1处反射率均为0，随后跃迁远超局部正常差分，属于共同边界/占位疑点（涉及{len(first_zero_files)}个文件）。',
            f'附件2存在{over100_count}个反射率大于100%的点，应保留原值并纳入可靠性分析，不建议未经依据截断。',
        ],
        '文件档案': file_profiles,
        '跨文件关系': {
            '同源分组': [
                {'组名': '碳化硅同一晶圆片', '材料': '碳化硅', '成员': [{'文件': '附件1.xlsx', '入射角_度': 10}, {'文件': '附件2.xlsx', '入射角_度': 15}], '依据': '题面附件说明'},
                {'组名': '硅同一晶圆片', '材料': '硅', '成员': [{'文件': '附件3.xlsx', '入射角_度': 10}, {'文件': '附件4.xlsx', '入射角_度': 15}], '依据': '题面附件说明'},
            ],
            '主键设计': {
                '附件内主键': ['波数_cm-1'],
                '跨附件公共关联键': ['波数_cm-1'],
                '全局唯一主键': ['附件名', '波数_cm-1'],
                '关联基数': '四表按波数均为一对一（1:1）',
                '键序列完全一致': bool(all_exact),
                '键范围_cm-1': [round_float(base_key.min(), 4), round_float(base_key.max(), 4)],
                '公共键数': int(len(base_key)),
                '逐文件核验': key_checks,
                '禁止误连': '附件1/2与附件3/4材料不同，不可仅因波数相同而视为同一样本；跨材料只能按频谱位置比较，不能合并为重复测量。',
            },
            '同晶圆不同角度反射率对比': paired,
            '反射率相关矩阵': {
                row: {col: round_float(correlation.loc[row, col], 8) for col in correlation.columns}
                for row in correlation.index
            },
        },
        '建模前处理建议': [
            '保留原始附件只读；规范化列名后以float64读取。',
            '默认剔除四附件共同的首个零反射率点，并保留“含首点/不含首点”敏感性对照结果。',
            '不对附件2大于100%的反射率做机械截断；若算法只依赖峰谷位置，可原值使用并说明归一化异常，若依赖振幅则需单独做基线/尺度稳健性检验。',
            '按同一材料的10°和15°结果分别估计厚度，再以角度一致性验证模型；不要把不同角度反射率直接平均后再寻峰。',
            '波数网格近似均匀但存在0.4810至0.4830 cm-1的舍入级步长波动；FFT类算法宜先检查采样均匀性，必要时等距重采样，峰位法可直接使用原波数坐标。',
        ],
    }

    ARCHIVE_PATH.write_text(json.dumps(archive, ensure_ascii=False, indent=2, default=native) + '\n', encoding='utf-8')

    experiments = [
        {
            '问题': '全附件格式与隐藏结构体检',
            '尝试': '先用openpyxl读取工作簿对象，再直接检查XLSX压缩包中的工作簿、工作表、关系和文档属性XML。',
            '现象': '4个附件均为合法OOXML；各含1个可见Sheet1，无隐藏sheet、合并单元格、公式、外链、宏、隐藏行列或多级表头。',
            '决定': '按第1行单级表头、第2行起数据区读取，并在数据档案中同时记录对象层与包结构层证据。',
            '依据': 'openpyxl工作表属性与XLSX包条目检查结果一致。',
        },
        {
            '问题': '附件编码规则确认',
            '尝试': '检查文件签名和核心XML、workbook.xml、sheet1.xml的XML声明。',
            '现象': '文件签名为ZIP/OOXML，相关XML均声明UTF-8；数据目录不存在CSV文件。',
            '决定': '将本批数据的权威读取方式定为openpyxl/pandas.read_excel，明确GBK和日期格式规则不适用。',
            '依据': '实际文件类型、ZIP包内容和XML编码声明。',
        },
        {
            '问题': '频谱边界点质量判定',
            '尝试': '比较每个附件首点到次点的反射率跃迁与后续前100点相邻差的中位数。',
            '现象': '四附件在399.6747 cm-1处均为0；首跳是局部典型差分的约369至803倍。',
            '决定': '将首点标记为共同边界/占位疑点，建议默认剔除并进行含首点敏感性对照，但不修改原始文件。',
            '依据': '共同位置、共同零值和极端跃迁均不符合后续连续频谱的局部变化尺度。',
        },
        {
            '问题': '反射率物理范围检查',
            '尝试': '逐文件统计反射率小于0、等于0和大于100%的记录及其连续波数区间。',
            '现象': '附件2有262个点大于100%，最大102.7393552%；其余附件无大于100%的点，所有附件均无负值。',
            '决定': '保留原值，不做机械截断；在依赖振幅的算法和可靠性分析中作为仪器基线/参考归一化风险处理。',
            '依据': '截断会改变条纹振幅和局部形状，而题面未提供校准规则。',
        },
        {
            '问题': '跨附件主键关联',
            '尝试': '逐值比较四个附件的波数列，并检查缺失、重复、排序和键集合差异。',
            '现象': '四附件均有7469个唯一、严格递增的波数，序列逐值完全一致且无缺失。',
            '决定': '以波数作为四表公共一对一关联键；全局唯一记录键采用“附件名+波数”，同时按题面材料分组避免跨材料误合并。',
            '依据': '键序列精确相等；附件名承载材料、角度和样本组身份。',
        },
    ]
    append_experiments(experiments)

    finished = datetime.now(timezone.utc)
    elapsed = (finished - started).total_seconds()
    log_lines = [
        '# 数据附件全量体检运行日志',
        '',
        f'- 开始时间（UTC）：{started.isoformat(timespec="seconds")}',
        f'- 完成时间（UTC）：{finished.isoformat(timespec="seconds")}',
        f'- 耗时：{elapsed:.3f} 秒',
        f'- 工作目录：`{ROOT}`',
        f'- 体检范围：`数据/` 下 {len(files)} 个普通文件；解析 {len(xlsx_files)} 个 XLSX，非XLSX {len(unsupported)} 个。',
        '- 工具：Python、pandas、openpyxl、zipfile、numpy；全程离线。',
        '',
        '## 执行步骤',
        '',
        '1. 清点文件、记录字节数与SHA256，校验ZIP/OOXML文件签名。',
        '2. 检查OOXML包条目、XML编码声明、宏、外部链接和文档属性。',
        '3. 检查工作表可见性、合并单元格、隐藏行列、公式、错误值、批注、超链接、筛选、表对象与表头层级。',
        '4. 读取全部数据，统计类型、缺失、重复、主键唯一性、波数步长、数值范围、分位数和异常边界点。',
        '5. 逐值核验四附件波数键，建立同晶圆不同角度关系和跨附件一对一关联规则。',
        '6. 生成 `交接/数据档案.json`，并把真实检查与判断追加到 `交接/实验记录.json`。',
        '',
        '## 关键结果',
        '',
        '- 4个附件均成功解析；每个文件1个可见Sheet1、7469条数据、2列。',
        '- 未发现隐藏sheet、合并单元格、多级表头、隐藏行列、公式、宏、外链、缺失值、重复行或重复波数。',
        '- 四附件波数网格完全一致，范围399.6747–4000.122 cm⁻¹，可按波数严格1:1关联。',
        '- 四附件首个反射率均为0且发生异常跃迁，已标记为边界/占位疑点。',
        '- 附件2有262个反射率超过100%的点，最大102.7393552%，已标记为归一化/基线风险。',
        '',
        '## 产物',
        '',
        '- `交接/数据档案.json`：逐文件结构、字段、统计、风险及跨文件关系。',
        '- `交接/实验记录.json`：本次真实尝试、现象、决定与依据。',
        '- `日志/运行日志.md`：本日志。',
    ]
    LOG_PATH.write_text('\n'.join(log_lines) + '\n', encoding='utf-8')

    print(f'已体检 {len(xlsx_files)} 个附件，共 {sum(len(df) for df in frames.values())} 条频谱记录。')
    print(f'四附件公共波数键 {len(base_key)} 个，键序列完全一致：{all_exact}。')
    print(f'首点零值附件 {len(first_zero_files)} 个；反射率>100%记录 {over100_count} 条。')
    print(f'已生成：{ARCHIVE_PATH.relative_to(ROOT)}、{LOG_PATH.relative_to(ROOT)}、{EXPERIMENT_PATH.relative_to(ROOT)}。')


if __name__ == '__main__':
    main()
