import sys
sys.path.insert(0, '/tmp/蜂巢/pylibs')

import json
import math
import time
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import minimize_scalar
from scipy.signal import find_peaks, savgol_filter


根目录 = Path('/tmp/蜂巢')
输出目录 = 根目录 / '求解' / '问题2' / '红队结果'
输出文件 = 输出目录 / '独立复算_问题2.json'
时间预算秒 = 18 * 60
开始时刻 = time.monotonic()


def 检查时间预算():
    if time.monotonic() - 开始时刻 > 时间预算秒:
        raise TimeoutError('复算接近18分钟时间预算，已主动终止。')


def 读光谱(path):
    """只用 openpyxl 读取题面指定的 Sheet1 两列原始数值。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb['Sheet1']
        rows = list(ws.iter_rows(min_row=2, max_col=2, values_only=True))
    finally:
        wb.close()
    arr = np.asarray(rows, dtype=float)
    if arr.ndim != 2 or arr.shape[1] != 2:
        raise ValueError(f'{path.name} 数据结构不是两列数值。')
    if not np.all(np.isfinite(arr)):
        raise ValueError(f'{path.name} 存在缺失或非有限数值。')
    波数, 反射率 = arr[:, 0], arr[:, 1]
    if not np.all(np.diff(波数) > 0):
        raise ValueError(f'{path.name} 波数不严格递增。')
    return 波数, 反射率


def 奇数窗口(n, preferred, minimum=7):
    w = min(int(preferred), n if n % 2 == 1 else n - 1)
    w = max(minimum, w)
    if w % 2 == 0:
        w -= 1
    return w


def 非偏振菲涅耳反射率(n, 入射角弧度):
    """空气到透明介质的非偏振强度反射率。"""
    s = math.sin(入射角弧度)
    c = math.cos(入射角弧度)
    ct = np.sqrt(1.0 - (s / n) ** 2)
    rs = ((c - n * ct) / (c + n * ct)) ** 2
    rp = ((n * c - ct) / (n * c + ct)) ** 2
    return 0.5 * (rs + rp)


def 由基线反演折射率(基线百分比, 入射角度):
    """
    用高波数透明区的慢变反射率基线逐点反演有效折射率。
    采用查表插值而非逐点迭代；仅为独立复算的有效折射率口径。
    """
    theta = math.radians(入射角度)
    n_grid = np.linspace(1.35, 4.50, 30001)
    r_grid = 非偏振菲涅耳反射率(n_grid, theta)
    r = np.clip(np.asarray(基线百分比) / 100.0,
                r_grid[0], r_grid[-1])
    return np.interp(r, r_grid, n_grid)


def 谐波残差平方(u, y, f):
    """给定光学频率，以常数项、线性漂移及正余弦项作线性回归。"""
    uc = (u - np.mean(u)) / max(np.ptp(u), 1e-12)
    phase = 2.0 * np.pi * f * u
    X = np.column_stack((np.ones_like(u), uc,
                         np.cos(phase), np.sin(phase)))
    coef, _, _, _ = np.linalg.lstsq(X, y, rcond=None)
    err = y - X @ coef
    return float(err @ err)


def 估计主频(u, y):
    """
    在 1--25 微米对应的频率范围内先粗网格、再局部连续优化。
    相位为 2*pi*(2d)*u，因此 d_cm=f/2。
    """
    order = np.argsort(u)
    u = np.asarray(u)[order]
    y = np.asarray(y)[order]
    keep = np.r_[True, np.diff(u) > 1e-10]
    u, y = u[keep], y[keep]
    grid = np.linspace(2.0e-4, 5.0e-3, 1200)
    scores = np.array([谐波残差平方(u, y, f) for f in grid])
    j = int(np.argmin(scores))
    lo = grid[max(0, j - 2)]
    hi = grid[min(len(grid) - 1, j + 2)]
    opt = minimize_scalar(lambda f: 谐波残差平方(u, y, f),
                          bounds=(lo, hi), method='bounded',
                          options={'xatol': 1e-11, 'maxiter': 200})
    f = float(opt.x)
    sse = 谐波残差平方(u, y, f)
    sst = float(np.sum((y - np.mean(y)) ** 2))
    r2 = 1.0 - sse / max(sst, 1e-15)
    return f, r2


def 单窗口复算(波数, 反射率, 入射角度, 下限, 上限):
    mask = (波数 >= 下限) & (波数 <= 上限)
    x = 波数[mask]
    r = 反射率[mask]
    if len(x) < 800:
        raise ValueError('有效窗口数据点不足。')

    # 小窗口抑制仪器高频噪声，大窗口仅提取慢变基线；均不覆盖原始值。
    smooth_w = 奇数窗口(len(r), 31)
    baseline_w = 奇数窗口(len(r), 1601)
    smooth = savgol_filter(r, smooth_w, 3, mode='interp')
    baseline = savgol_filter(smooth, baseline_w, 3, mode='interp')
    residual = smooth - baseline

    n_eff = 由基线反演折射率(baseline, 入射角度)
    theta = math.radians(入射角度)
    u = x * np.sqrt(n_eff ** 2 - math.sin(theta) ** 2)

    # 避开大窗口滤波的端部影响。
    edge = min(450, max(80, len(x) // 12))
    core = slice(edge, len(x) - edge)
    f, r2 = 估计主频(u[core], residual[core])
    thickness_harmonic = f * 5000.0  # f/2 cm -> f*5000 微米

    # 完全不同的交叉校验：在光学坐标上以相邻波峰间距求厚度。
    y_core = residual[core]
    u_core = u[core]
    du_med = float(np.median(np.diff(u_core)))
    expected_period_u = 1.0 / max(f, 1e-12)
    min_distance = max(5, int(0.55 * expected_period_u / max(du_med, 1e-12)))
    prominence = max(0.12 * float(np.std(y_core)), 1e-7)
    peaks, _ = find_peaks(y_core, distance=min_distance,
                          prominence=prominence)
    if len(peaks) >= 4:
        spacings = np.diff(u_core[peaks])
        med = float(np.median(spacings))
        mad = float(np.median(np.abs(spacings - med)))
        good = np.abs(spacings - med) <= max(3.0 * 1.4826 * mad,
                                             0.20 * med)
        spacing = float(np.median(spacings[good])) if np.any(good) else med
        thickness_peak = 5000.0 / spacing
        peak_count = int(len(peaks))
    else:
        thickness_peak = None
        peak_count = int(len(peaks))

    return {
        '窗口下限_cm-1': float(下限),
        '窗口上限_cm-1': float(上限),
        '样本数': int(len(x)),
        '剔除首个占位疑点': bool(波数[0] < 下限),
        '有效折射率中位数': float(np.median(n_eff[core])),
        '有效折射率范围': [float(np.min(n_eff[core])),
                           float(np.max(n_eff[core]))],
        '谐波频率_光学坐标倒数': f,
        '谐波拟合优度_R2': float(r2),
        '谐波厚度_微米': float(thickness_harmonic),
        '峰数': peak_count,
        '峰间距厚度_微米': (None if thickness_peak is None
                           else float(thickness_peak))
    }


def 单附件复算(path, 入射角度):
    波数, 反射率 = 读光谱(path)
    # 主窗口覆盖高波数透明区；四个子窗口用于评估频段稳定性。
    主结果 = 单窗口复算(波数, 反射率, 入射角度, 1450.0, 4000.0)
    子窗口 = [(1450.0, 2600.0), (1800.0, 3000.0),
              (2200.0, 3500.0), (2600.0, 4000.0)]
    子结果 = []
    for lo, hi in 子窗口:
        检查时间预算()
        try:
            子结果.append(单窗口复算(波数, 反射率, 入射角度,
                                      lo, hi))
        except Exception as exc:
            子结果.append({'窗口下限_cm-1': lo, '窗口上限_cm-1': hi,
                           '失败原因': str(exc)})

    thicknesses = [主结果['谐波厚度_微米']]
    thicknesses += [z['谐波厚度_微米'] for z in 子结果
                    if '谐波厚度_微米' in z]
    center = float(np.median(thicknesses))
    mad = float(np.median(np.abs(np.asarray(thicknesses) - center)))

    # 主数值取全透明区谐波回归；分段中位数/MAD只作为可靠性证据。
    return {
        '文件': path.name,
        '入射角_度': float(入射角度),
        '厚度_微米': float(主结果['谐波厚度_微米']),
        '分段厚度中位数_微米': center,
        '分段厚度MAD_微米': mad,
        '主窗口结果': 主结果,
        '子窗口结果': 子结果,
        '原始样本数': int(len(波数)),
        '原始波数范围_cm-1': [float(波数[0]), float(波数[-1])],
        '异常处理': '保留原始数据；计算窗口自然排除399.6747 cm-1首点，不截断大于100%的反射率。'
    }


def 写失败结果(exc):
    输出目录.mkdir(parents=True, exist_ok=True)
    payload = {
        '问题': 2,
        '复算方式': '独立实现，未读建模师代码',
        '状态': '失败',
        '失败原因': str(exc),
        '复算指标': {
            '附件1厚度_微米': None,
            '附件2厚度_微米': None,
            '双角平均厚度_微米': None,
            '双角相对差': None
        }
    }
    输出文件.write_text(json.dumps(payload, ensure_ascii=False, indent=2),
                         encoding='utf-8')


def main():
    输出目录.mkdir(parents=True, exist_ok=True)
    a1 = 单附件复算(根目录 / '数据' / '附件1.xlsx', 10.0)
    检查时间预算()
    a2 = 单附件复算(根目录 / '数据' / '附件2.xlsx', 15.0)
    d1, d2 = a1['厚度_微米'], a2['厚度_微米']
    mean_d = 0.5 * (d1 + d2)
    relative_gap = abs(d1 - d2) / max(abs(mean_d), 1e-12)

    payload = {
        '问题': 2,
        '复算方式': '独立实现，未读建模师代码；慢变基线菲涅耳反演有效折射率，光学坐标谐波回归求厚度，峰间距与分段窗口交叉校验。',
        '状态': '成功',
        '复算指标': {
            '附件1厚度_微米': round(d1, 6),
            '附件2厚度_微米': round(d2, 6),
            '双角平均厚度_微米': round(mean_d, 6),
            '双角相对差': round(relative_gap, 8)
        },
        '口径说明': {
            '附件1厚度_微米': '附件1（同一碳化硅晶圆片，10°）1450–4000 cm⁻¹透明区；原始反射率慢变基线逐点按非偏振菲涅耳公式反演有效折射率，变换至σ·sqrt(n(σ)^2-sin²θ)光学坐标后作单主频谐波回归；厚度单位微米。',
            '附件2厚度_微米': '附件2（同一碳化硅晶圆片，15°）采用与附件1完全相同的窗口、异常处理、折射率反演和聚合方式；不截断大于100%的原始反射率点，厚度单位微米。',
            '双角平均厚度_微米': '附件1与附件2两个独立厚度的等权算术平均，单位微米。',
            '双角相对差': '|附件1厚度-附件2厚度|/双角平均厚度，无量纲。',
            '样本范围与异常点': '每附件原始7469点；厚度主计算仅用1450–4000 cm⁻¹，故399.6747 cm⁻¹的0值首点被排除；窗口内不做人为删点、不截断反射率。',
            '聚合方式': '头条厚度取全主窗口谐波回归值；四个重叠子窗口的中位数和MAD、以及波峰间距法仅用于独立稳定性核验。'
        },
        '附件明细': [a1, a2],
        '运行秒数': round(time.monotonic() - 开始时刻, 3),
        '时间预算秒': 时间预算秒
    }
    输出文件.write_text(json.dumps(payload, ensure_ascii=False, indent=2),
                         encoding='utf-8')
    print(json.dumps({
        '输出文件': str(输出文件),
        '附件1厚度_微米': payload['复算指标']['附件1厚度_微米'],
        '附件2厚度_微米': payload['复算指标']['附件2厚度_微米'],
        '双角相对差': payload['复算指标']['双角相对差'],
        '运行秒数': payload['运行秒数']
    }, ensure_ascii=False))


if __name__ == '__main__':
    try:
        main()
    except Exception as exc:
        写失败结果(exc)
        print(json.dumps({'状态': '失败', '原因': str(exc)}, ensure_ascii=False))
        raise
