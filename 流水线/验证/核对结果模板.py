"""核对结果模板交付物（2026 国赛 A 题起）：题目要求把完整结果按 附件3 的模板写进 result1.xlsx–result4.xlsx。
用法：python3 流水线/验证/核对结果模板.py <成品或镜像目录> <模板目录(含 result*.xlsx)> [--只=resultN.xlsx]
逐个模板：在目录树里找同名文件（排除 数据/ 与 模板目录本身）；核对工作表名、A 列首格、首行距离网格（0…2，步长 0.1 → 21 列，
result4 末列为「药材表面」）、A 列时间步（result1/2 每 1 s，result3/4 每 60 s）、数值非空且为数、四位小数。退出码 1 = 有不合格。"""
import pathlib
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("需要 openpyxl（用 流水线/运行时/venv/bin/python3 跑）")

根 = pathlib.Path(sys.argv[1]).resolve()
模板目录 = pathlib.Path(sys.argv[2]).resolve()
只 = {a.split("=", 1)[1] for a in sys.argv[3:] if a.startswith("--只=")}      # 可选：只核对指定模板（驱动逐问核对用）
坏 = 0


def 报(ok, 文, 说):
    global 坏
    print(f"  {'✓' if ok else '✗'} {文}：{说}")
    if not ok:
        坏 += 1


def 找(名):
    候选 = [p for p in 根.rglob(名) if 模板目录 not in p.parents and "数据" not in p.relative_to(根).parts]
    return sorted(候选, key=lambda p: p.stat().st_mtime)[-1] if 候选 else None


def 网格(ws):
    首行 = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    return 首行


for 模板 in sorted(模板目录.glob("result*.xlsx")):
    名 = 模板.name
    if 只 and 名 not in 只:
        continue
    f = 找(名)
    if not f:
        报(False, 名, "未找到交付文件（应在 求解/ 或 交接/ 下，与模板同名）"); continue
    try:
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        wt = openpyxl.load_workbook(模板, read_only=True)
    except Exception as e:
        报(False, 名, f"打不开：{e}"); continue
    if [w.title for w in wb.worksheets] != [w.title for w in wt.worksheets]:
        报(False, 名, f"工作表名 {[w.title for w in wb.worksheets]} ≠ 模板 {[w.title for w in wt.worksheets]}"); continue
    for ws in wb.worksheets:
        首 = 网格(ws)
        期望列 = 21 + 1 + (1 if 名 == "result4.xlsx" else 0)     # A 列 + 0…2 步长 0.1 共 21 列（+ 药材表面）
        距离 = [x for x in 首[1:] if x is not None]
        try:
            数距 = [float(x) for x in 距离 if not (isinstance(x, str) and "表面" in x)]
            网格对 = len(数距) == 21 and abs(数距[0]) < 1e-9 and abs(数距[-1] - 2) < 1e-9 and all(abs(数距[i+1] - 数距[i] - 0.1) < 1e-6 for i in range(20))
        except Exception:
            网格对 = False
        表面对 = (名 != "result4.xlsx") or (isinstance(距离[-1], str) and "表面" in 距离[-1])
        行 = list(ws.iter_rows(min_row=2, values_only=True))
        时间 = [r[0] for r in 行 if r and r[0] is not None]
        步 = 1 if 名 in ("result1.xlsx", "result2.xlsx") else 60
        时间对 = len(时间) >= 2 and all(isinstance(t, (int, float)) for t in 时间) and all(abs((时间[i+1] - 时间[i]) - 步) < 1e-6 for i in range(min(len(时间) - 1, 50)))
        值 = [c for r in 行 for c in r[1:len(距离) + 1]]
        非空 = sum(1 for v in 值 if isinstance(v, (int, float)))
        四位 = all((isinstance(v, int) or round(v, 4) == v) for v in 值 if isinstance(v, (int, float)))
        报(网格对 and 表面对, f"{名}[{ws.title}]", f"首行距离网格 0…2/0.1（21 列）{'对' if 网格对 else '不对'}；{'药材表面列 ' + ('有' if 表面对 else '无') if 名 == 'result4.xlsx' else ''}")
        报(时间对, f"{名}[{ws.title}]", f"A 列时间 {len(时间)} 行，步长应为 {步} s：{'对' if 时间对 else '不对'}（首格 {首[0]!r}）")
        if 名 == "result4.xlsx":
            # 2026 A 题契约裁定（歧义「收缩后固定距离与表面列以及越界格的处理」）：半径收缩后，超出当时半径的固定距离格
            # 不属于药材，留空而不填 0；「药材表面」列每行必填。故 result4 的非空判据 = 表面列全非空 + 每行的空格只许是尾部连续越界段
            # （中心侧一旦有值、外侧再出现值后又出现空格即为漏填）。其余三个文件仍要求全部非空。
            坏行 = 0
            表面空 = 0
            for r in 行:
                if not r or r[0] is None:
                    continue
                内 = list(r[1:len(距离)])          # 0…2 的 21 个固定距离格
                表 = r[len(距离)] if len(r) > len(距离) else None
                if not isinstance(表, (int, float)):
                    表面空 += 1
                有值 = [isinstance(v, (int, float)) for v in 内]
                if not 有值 or not 有值[0]:
                    坏行 += 1
                    continue
                首空 = 有值.index(False) if False in 有值 else len(有值)
                if any(有值[首空:]):                # 空格之后又出现数值 → 不是尾部越界段
                    坏行 += 1
            报(坏行 == 0 and 表面空 == 0 and 非空 > 0, f"{名}[{ws.title}]",
              f"数值 {非空}/{len(值)} 非空；表面列空格 {表面空} 行；空格非尾部越界段的行 {坏行}（契约裁定：越界格留空、表面列必填）")
        else:
            报(非空 == len(值) and len(值) > 0, f"{名}[{ws.title}]", f"数值 {非空}/{len(值)} 非空")
        报(四位, f"{名}[{ws.title}]", "四位小数" if 四位 else "存在超过四位小数的值")
        # 角格（A1）只作提示不计不合格：契约只规定 A 列时间/首行距离/工作表名，模板角格「时间\\到药材中心的距离」是否沿用属呈现细节。
        模板首 = 网格(wt[ws.title]) if ws.title in wt.sheetnames else None
        if 模板首 and 模板首[0] is not None and str(首[0]).strip() != str(模板首[0]).strip():
            print(f"  △ {名}[{ws.title}]：角格 A1={首[0]!r} 与模板 {模板首[0]!r} 不同（提示，不计不合格）")
print("结果模板核对：", "全部合格" if 坏 == 0 else f"{坏} 项不合格")
sys.exit(1 if 坏 else 0)
